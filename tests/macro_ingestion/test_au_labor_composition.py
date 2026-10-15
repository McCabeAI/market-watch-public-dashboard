"""ABS labour composition, QA, cache, and workbook isolation tests."""

from __future__ import annotations

import copy
import io
import json
import unittest
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
import sys

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.macro_ingestion.adapters.au import (
    clear_workbook_cache,
    fetch_series,
    parse_abs_time_series_workbook,
)
from scripts.macro_ingestion.contract import load_catalog

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "au"
COMPOSITION = FIXTURES / "labour_composition_fixture.xlsx"
HOURS = FIXTURES / "labour_hours_fixture.xlsx"
UNDER = FIXTURES / "labour_underutilisation_fixture.xlsx"
SMOKE_PATH = FIXTURES / "abs_labour_live_smoke.json"

DUE_AUG_LABOUR = datetime(2026, 9, 24, 2, 0, tzinfo=timezone.utc)

AUG_2026 = {
    "A84423050A": 4.6462961,
    "A84423043C": 14836.5771018,
    "A84423041X": 10193.9298744,
    "A84423042A": 4642.6472274,
    "A84423054K": 63.9424839,
    "A84423051C": 67.0582068,
    "A84423046K": 722.9412952,
    "A84423047L": 15559.518397,
    "A84426277X": 2009106.80803,
    "A85255719L": 962.9224142,
    "A85255725J": 6.1886389,
    "A85255726K": 10.834935,
}


def _build_abs_workbook(series_id: str, observations: list[tuple[datetime, float]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data1"
    for _ in range(9):
        ws.append([None])
    ws.append([None, series_id])
    for dt, val in observations:
        ws.append([dt, val])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestAuLaborComposition(unittest.TestCase):
    def setUp(self) -> None:
        clear_workbook_cache()
        self.catalog = load_catalog()
        self.composition_bytes = COMPOSITION.read_bytes()

    def tearDown(self) -> None:
        clear_workbook_cache()

    def _spec(self, series_id_catalog: str) -> dict:
        return copy.deepcopy(next(r for r in self.catalog["series"] if r["id"] == series_id_catalog))

    def test_composition_fixture_august_values(self) -> None:
        body = self.composition_bytes
        for sid, expected in AUG_2026.items():
            if sid.startswith("A852") or sid == "A84426277X":
                continue
            series = parse_abs_time_series_workbook(body, sid, cadence="monthly")
            period, value = series[-1]
            self.assertEqual(period, "2026-08")
            self.assertAlmostEqual(value, expected)

    def test_fetch_unemployment_passes_identity_qa(self) -> None:
        spec = self._spec("AU.Labor.unemployment")
        spec["endpoint"] = "https://fixture.test/aug-2026/62020001.xlsx"

        def opener(url: str, *, timeout: float = 20) -> dict:
            return {
                "ok": True,
                "url": url,
                "http_status": 200,
                "body": self.composition_bytes,
                "error": None,
            }

        payload = fetch_series(spec, opener=opener, now=DUE_AUG_LABOUR)
        self.assertTrue(payload.get("ok"))
        self.assertAlmostEqual(payload["points"][0]["value"], AUG_2026["A84423050A"])

    def test_corrupted_unemployment_fails_identity_mismatch(self) -> None:
        import openpyxl

        spec = self._spec("AU.Labor.unemployment")
        spec["endpoint"] = "https://fixture.test/aug-2026/62020001.xlsx"
        wb = openpyxl.load_workbook(io.BytesIO(self.composition_bytes))
        ws = wb["Data1"]
        for row in ws.iter_rows(min_row=11):
            if row[0].value and str(row[0].value).startswith("2026-08"):
                for cell in row:
                    if cell.value == AUG_2026["A84423050A"]:
                        cell.value = 9.9999999
        buf = io.BytesIO()
        wb.save(buf)
        bad = buf.getvalue()

        def opener(url: str, *, timeout: float = 20) -> dict:
            return {
                "ok": True,
                "url": url,
                "http_status": 200,
                "body": bad,
                "error": None,
            }

        payload = fetch_series(spec, opener=opener, now=DUE_AUG_LABOUR)
        self.assertFalse(payload.get("ok"))
        self.assertEqual(payload.get("status"), "source_failed")
        self.assertTrue(str(payload.get("error") or "").startswith("identity_mismatch:"))

    def test_hours_not_in_table1_workbook(self) -> None:
        spec = {
            "id": "AU.Labor.hours_worked",
            "series_id": "A84426277X",
            "endpoint": "fixture://labour_composition_fixture.xlsx",
            "retrieval_method": "abs_time_series_workbook",
            "transform": "thousand_hours_sa",
            "cadence": "monthly",
        }

        def opener(url: str, *, timeout: float = 20) -> dict:
            return {
                "ok": True,
                "url": url,
                "http_status": 200,
                "body": self.composition_bytes,
                "error": None,
            }

        payload = fetch_series(spec, opener=opener, now=DUE_AUG_LABOUR)
        self.assertFalse(payload.get("ok"))
        self.assertIn("series_id_not_in_workbook", str(payload.get("error")))

    def test_workbook_cache_single_fetch(self) -> None:
        spec_u = self._spec("AU.Labor.unemployment")
        spec_ft = self._spec("AU.Labor.employment")
        spec_ft["series_id"] = "A84423041X"
        spec_ft["transform"] = "thousands_sa"
        url = "https://example.test/aug-2026/62020001.xlsx"
        spec_u["endpoint"] = url
        spec_ft["endpoint"] = url
        calls = 0

        def opener(u: str, *, timeout: float = 20) -> dict:
            nonlocal calls
            calls += 1
            return {
                "ok": True,
                "url": u,
                "http_status": 200,
                "body": self.composition_bytes,
                "error": None,
            }

        p1 = fetch_series(spec_u, opener=opener, now=DUE_AUG_LABOUR)
        p2 = fetch_series(spec_ft, opener=opener, now=DUE_AUG_LABOUR)
        self.assertTrue(p1.get("ok") and p2.get("ok"))
        self.assertEqual(calls, 1)

    def test_revision_parsing_yields_distinct_points(self) -> None:
        import openpyxl

        body_v1 = self.composition_bytes
        wb = openpyxl.load_workbook(io.BytesIO(body_v1))
        ws = wb["Data1"]
        for row in ws.iter_rows(min_row=11):
            if row[0].value and str(row[0].value).startswith("2026-08"):
                for cell in row:
                    if cell.value == AUG_2026["A84423050A"]:
                        cell.value = 4.7
        buf = io.BytesIO()
        wb.save(buf)
        body_v2 = buf.getvalue()
        s1 = parse_abs_time_series_workbook(body_v1, "A84423050A", cadence="monthly")
        s2 = parse_abs_time_series_workbook(body_v2, "A84423050A", cadence="monthly")
        aug1 = next(v for p, v in s1 if p == "2026-08")
        aug2 = next(v for p, v in s2 if p == "2026-08")
        self.assertNotAlmostEqual(aug1, aug2)
        self.assertEqual(len(s1), len(s2))

    def test_404_body_not_cached_or_success(self) -> None:
        spec = self._spec("AU.Labor.unemployment")
        calls = 0

        def opener(url: str, *, timeout: float = 20) -> dict:
            nonlocal calls
            calls += 1
            return {
                "ok": False,
                "url": url,
                "http_status": 404,
                "body": b"not found",
                "error": "HTTP 404",
            }

        payload = fetch_series(spec, opener=opener, now=DUE_AUG_LABOUR)
        self.assertFalse(payload.get("ok"))
        self.assertEqual(payload.get("http_status"), 404)
        self.assertIn("404", str(payload.get("error")))
        fetch_series(spec, opener=opener, now=DUE_AUG_LABOUR)
        self.assertEqual(calls, 2)

    def test_hours_fixture_endpoint(self) -> None:
        spec = {
            "id": "AU.Labor.hours_worked",
            "series_id": "A84426277X",
            "endpoint": "fixture://labour_hours_fixture.xlsx",
            "retrieval_method": "abs_time_series_workbook",
            "transform": "thousand_hours_sa",
            "cadence": "monthly",
        }
        hours_bytes = HOURS.read_bytes()

        def opener(url: str, *, timeout: float = 20) -> dict:
            return {
                "ok": True,
                "url": url,
                "http_status": 200,
                "body": hours_bytes,
                "error": None,
            }

        payload = fetch_series(spec, opener=opener, now=DUE_AUG_LABOUR)
        self.assertTrue(payload.get("ok"))
        self.assertAlmostEqual(payload["points"][0]["value"], AUG_2026["A84426277X"])

    def test_underutilisation_fixture(self) -> None:
        under_bytes = UNDER.read_bytes()
        for sid in ("A85255719L", "A85255725J", "A85255726K"):
            series = parse_abs_time_series_workbook(under_bytes, sid, cadence="monthly")
            self.assertAlmostEqual(series[-1][1], AUG_2026[sid])

    def test_live_smoke_report_present(self) -> None:
        if not SMOKE_PATH.is_file():
            self.skipTest("abs_labour_live_smoke.json not written yet")
        report = json.loads(SMOKE_PATH.read_text(encoding="utf-8"))
        self.assertIn("fetched_at", report)
        probes = report.get("probes") or []
        table1 = next((p for p in probes if "62020001" in str(p.get("url"))), None)
        self.assertIsNotNone(table1)
        if table1.get("ok"):
            self.assertEqual(
                table1.get("sha256"),
                "ef60b5c0bcdc44c9f0bec3ef537a1e0f6965d67a38f7a4ed1b038d6d2007712f",
            )
            if table1.get("series_id") == "A84423050A":
                self.assertAlmostEqual(float(table1.get("value")), 4.6462961)


if __name__ == "__main__":
    unittest.main()
