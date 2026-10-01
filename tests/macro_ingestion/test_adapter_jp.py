"""Japan macro ingestion adapter tests."""

from __future__ import annotations

import copy
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[2]
import sys

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.japan_macro_data import (
    METI_RETAIL_INDEX_URL_EN,
    METI_RETAIL_INDEX_URL_JA,
    extract_meti_retail_workbook_links,
    meti_retail_period_from_url,
)
from scripts.macro_ingestion.adapters.jp import fetch_series
from scripts.macro_ingestion.calendar import release_due
from scripts.macro_ingestion.contract import load_catalog
from scripts.macro_ingestion.runner import run_ingestion
from scripts.macro_ingestion.windows import cutoff_class

FIXTURES = ROOT / "tests/macro_ingestion/fixtures/jp"
CPI_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032103932&fileKind=1"
TANKAN_URL = "https://www.boj.or.jp/en/statistics/tk/gaiyo/2026/tka2606.zip"
PMI_LISTING = "https://www.pmi.spglobal.com/Public/Release/PressReleases"
METI_RETAIL_JULY_URL = (
    "https://www.meti.go.jp/statistics/tyo/syoudou/result/excel/DB_202607S.xlsx"
)
METI_RETAIL_AUGUST_URL = (
    "https://www.meti.go.jp/statistics/tyo/syoudou/result/excel/DB_202608S.xlsx"
)
METI_RETAIL_SOKUHO_URL = "https://www.meti.go.jp/statistics/tyo/syoudou/result/sokuho_2.html"


class TestJapanMacroAdapter(unittest.TestCase):
    def setUp(self) -> None:
        self.catalog = load_catalog()
        self.tmp = tempfile.TemporaryDirectory()
        self.obs_dir = Path(self.tmp.name) / "observations"
        self.health_dir = Path(self.tmp.name) / "health"

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _mini_catalog(self, spec: dict) -> dict:
        cat = copy.deepcopy(self.catalog)
        cat["series"] = [spec]
        return cat

    def _fixture_opener(self, mapping: dict[str, Path]):
        def opener(url: str, *, timeout: float = 20):
            path = mapping.get(url)
            if path is None:
                return {
                    "ok": False,
                    "url": url,
                    "http_status": None,
                    "body": b"",
                    "error": "unexpected_url",
                }
            body = path.read_bytes()
            return {
                "ok": True,
                "url": url,
                "http_status": 200,
                "body": body,
                "error": None,
            }

        return opener

    def _tracking_opener(self, mapping: dict[str, Path]):
        calls: list[str] = []

        def opener(url: str, *, timeout: float = 20):
            calls.append(url)
            return self._fixture_opener(mapping)(url, timeout=timeout)

        return opener, calls

    @staticmethod
    def _mini_meti_retail_xlsx() -> bytes:
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DB_MainC_Graph"
        ws.cell(row=5, column=2, value="2026年8月")
        ws.cell(row=5, column=6, value=3.2)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_cpi_fixture_yields_new_observation_via_runner(self) -> None:
        spec = copy.deepcopy(next(r for r in self.catalog["series"] if r["id"] == "JP.Inflation.headline"))
        cat = self._mini_catalog(spec)
        opener = self._fixture_opener({CPI_URL: FIXTURES / "estat_cpi_2025base_yoy.csv"})
        now = datetime(2026, 9, 23, 12, 0, tzinfo=timezone.utc)
        result = run_ingestion(
            mode="offline",
            countries=["JP"],
            catalog=cat,
            now=now,
            opener=opener,
            observations_dir=self.obs_dir,
            health_dir=self.health_dir,
        )
        self.assertEqual(result["rows"][0]["status"], "new_observation")
        store = json.loads((self.obs_dir / "jp.json").read_text(encoding="utf-8"))
        self.assertTrue(store["observations"])

    def test_tankan_non_zip_source_failed_no_points(self) -> None:
        spec = copy.deepcopy(
            next(r for r in self.catalog["series"] if r["id"] == "JP.Activity.tankan_large_manufacturing")
        )
        opener = self._fixture_opener({TANKAN_URL: FIXTURES / "tankan_not_a_zip.html"})
        payload = fetch_series(spec, opener=opener, now=datetime(2026, 9, 23, tzinfo=timezone.utc))
        self.assertEqual(payload.get("status"), "source_failed")
        self.assertFalse(payload.get("points"))

    def test_pmi_listing_challenge_license_gap(self) -> None:
        spec = copy.deepcopy(next(r for r in self.catalog["series"] if r["id"] == "JP.Activity.business_surveys"))
        opener = self._fixture_opener({PMI_LISTING: FIXTURES / "pmi_listing_challenge.html"})
        payload = fetch_series(spec, opener=opener, now=datetime(2026, 9, 23, tzinfo=timezone.utc))
        self.assertIn(payload.get("status"), {"license_gap", "source_failed"})
        self.assertFalse(payload.get("points"))

    def test_flash_not_due_before_release_instant(self) -> None:
        flash = copy.deepcopy(
            next(r for r in self.catalog["series"] if r["id"] == "JP.Activity.flash_composite_pmi")
        )
        before = datetime(2026, 9, 24, 0, 29, tzinfo=timezone.utc)
        self.assertFalse(release_due(flash, before))

    def test_flash_due_at_release_instant(self) -> None:
        flash = copy.deepcopy(
            next(r for r in self.catalog["series"] if r["id"] == "JP.Activity.flash_composite_pmi")
        )
        at_release = datetime(2026, 9, 24, 0, 30, tzinfo=timezone.utc)
        self.assertTrue(release_due(flash, at_release))

    def test_flash_due_at_0030_utc_is_post_freeze_at_2030_et_prior_day(self) -> None:
        ny = ZoneInfo("America/New_York")
        when = datetime(2026, 9, 23, 20, 30, tzinfo=ny)
        self.assertEqual(when.astimezone(timezone.utc), datetime(2026, 9, 24, 0, 30, tzinfo=timezone.utc))
        self.assertEqual(cutoff_class(when), "post_freeze")

    def test_flash_adapter_does_not_invent_value_without_pdf(self) -> None:
        flash = copy.deepcopy(
            next(r for r in self.catalog["series"] if r["id"] == "JP.Activity.flash_composite_pmi")
        )
        opener = self._fixture_opener({PMI_LISTING: FIXTURES / "pmi_listing_empty.html"})
        payload = fetch_series(flash, opener=opener, now=datetime(2026, 9, 24, 0, 30, tzinfo=timezone.utc))
        self.assertTrue(payload.get("ok"))
        self.assertEqual(payload.get("points"), [])

    def test_live_smoke_report_present(self) -> None:
        report_path = FIXTURES / "live_smoke_report.json"
        self.assertTrue(report_path.is_file(), "live_smoke_report.json missing")
        report = json.loads(report_path.read_text(encoding="utf-8"))
        labels = {c["label"] for c in report.get("checks", [])}
        self.assertTrue({"estat_cpi", "boj_cgpi_page", "sp_japan_pmi_listing"}.issubset(labels))

    def test_meti_retail_discover_latest_workbook_via_listings(self) -> None:
        retail = copy.deepcopy(next(r for r in self.catalog["series"] if r["id"] == "JP.Consumer.retail"))
        ja_index = (
            '<html><body><a href="result/sokuho_2.html">速報 (2026年8月分)</a></body></html>'
        )
        sokuho = (
            "<html><body>"
            '<a href="excel/DB_202607S.xlsx">prelim July</a>'
            '<a href="excel/DB_202608S.xlsx">prelim August</a>'
            "</body></html>"
        )
        en_index = (
            f'<html><body><a href="{METI_RETAIL_AUGUST_URL}">English index</a></body></html>'
        )
        xlsx = self._mini_meti_retail_xlsx()
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            paths = {
                METI_RETAIL_INDEX_URL_JA: root / "ja_index.html",
                METI_RETAIL_SOKUHO_URL: root / "sokuho.html",
                METI_RETAIL_INDEX_URL_EN: root / "en_index.html",
                METI_RETAIL_AUGUST_URL: root / "august.xlsx",
            }
            paths[METI_RETAIL_INDEX_URL_JA].write_text(ja_index, encoding="utf-8")
            paths[METI_RETAIL_SOKUHO_URL].write_text(sokuho, encoding="utf-8")
            paths[METI_RETAIL_INDEX_URL_EN].write_text(en_index, encoding="utf-8")
            paths[METI_RETAIL_AUGUST_URL].write_bytes(xlsx)
            opener, calls = self._tracking_opener(paths)
            payload = fetch_series(
                retail,
                opener=opener,
                now=datetime(2026, 10, 1, tzinfo=timezone.utc),
            )
        self.assertTrue(payload.get("ok"))
        self.assertTrue(str(payload.get("source_url", "")).endswith("DB_202608S.xlsx"))
        self.assertEqual(payload.get("artifact_release_period"), "2026-08")
        self.assertEqual(payload.get("official_latest_period"), "2026-08")
        periods = {p["period"] for p in payload.get("points") or []}
        self.assertIn("2026-08", periods)
        self.assertNotIn(METI_RETAIL_JULY_URL, calls)

    def test_meti_retail_discovery_failure_ignores_pinned_endpoint(self) -> None:
        retail = copy.deepcopy(next(r for r in self.catalog["series"] if r["id"] == "JP.Consumer.retail"))
        empty = "<html><body><p>No workbook links here.</p></body></html>"
        xlsx = self._mini_meti_retail_xlsx()
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            paths = {
                METI_RETAIL_INDEX_URL_JA: root / "ja.html",
                METI_RETAIL_INDEX_URL_EN: root / "en.html",
                METI_RETAIL_JULY_URL: root / "july.xlsx",
            }
            paths[METI_RETAIL_INDEX_URL_JA].write_text(empty, encoding="utf-8")
            paths[METI_RETAIL_INDEX_URL_EN].write_text(empty, encoding="utf-8")
            paths[METI_RETAIL_JULY_URL].write_bytes(xlsx)
            opener, calls = self._tracking_opener(paths)
            payload = fetch_series(
                retail,
                opener=opener,
                now=datetime(2026, 10, 1, tzinfo=timezone.utc),
            )
        self.assertFalse(payload.get("ok"))
        self.assertEqual(payload.get("status"), "source_failed")
        self.assertIn("meti_retail_release_undiscovered", str(payload.get("error") or ""))
        self.assertFalse(payload.get("points"))
        self.assertNotIn(METI_RETAIL_JULY_URL, calls)

    def test_meti_retail_discovery_august_fetch_failure_no_july_fallback(self) -> None:
        retail = copy.deepcopy(next(r for r in self.catalog["series"] if r["id"] == "JP.Consumer.retail"))
        ja_index = '<html><body><a href="result/sokuho_2.html">速報</a></body></html>'
        sokuho = '<html><body><a href="excel/DB_202608S.xlsx">August</a></body></html>'
        en_index = "<html><body></body></html>"
        xlsx = self._mini_meti_retail_xlsx()
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            paths = {
                METI_RETAIL_INDEX_URL_JA: root / "ja.html",
                METI_RETAIL_SOKUHO_URL: root / "sokuho.html",
                METI_RETAIL_INDEX_URL_EN: root / "en.html",
                METI_RETAIL_JULY_URL: root / "july.xlsx",
            }
            paths[METI_RETAIL_INDEX_URL_JA].write_text(ja_index, encoding="utf-8")
            paths[METI_RETAIL_SOKUHO_URL].write_text(sokuho, encoding="utf-8")
            paths[METI_RETAIL_INDEX_URL_EN].write_text(en_index, encoding="utf-8")
            paths[METI_RETAIL_JULY_URL].write_bytes(xlsx)

            def opener(url: str, *, timeout: float = 20):
                if url == METI_RETAIL_AUGUST_URL:
                    return {
                        "ok": False,
                        "url": url,
                        "http_status": 503,
                        "body": b"",
                        "error": "august_unavailable",
                    }
                return self._fixture_opener(paths)(url, timeout=timeout)

            payload = fetch_series(
                retail,
                opener=opener,
                now=datetime(2026, 10, 1, tzinfo=timezone.utc),
            )
        self.assertFalse(payload.get("ok"))
        self.assertEqual(payload.get("status"), "source_failed")
        self.assertFalse(payload.get("points"))

    def test_meti_retail_workbook_link_extraction(self) -> None:
        sokuho_base = METI_RETAIL_SOKUHO_URL
        html = (
            '<a href="excel/DB_202608S.xlsx">August</a>'
            '<a href="excel/DB_202607K-1.xlsx">Final</a>'
        )
        links = extract_meti_retail_workbook_links(html, sokuho_base)
        self.assertEqual(links, [("2026-08", METI_RETAIL_AUGUST_URL)])
        self.assertEqual(meti_retail_period_from_url(METI_RETAIL_AUGUST_URL), "2026-08")
        self.assertIsNone(meti_retail_period_from_url("excel/DB_202607K-1.xlsx"))


if __name__ == "__main__":
    unittest.main()
