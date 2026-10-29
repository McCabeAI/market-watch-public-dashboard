"""New Zealand macro ingestion adapter (Stats NZ, BNZ/BusinessNZ, ANZ)."""

from __future__ import annotations

import csv
import hashlib
import html as html_lib
import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urljoin

from openpyxl import load_workbook
from pypdf import PdfReader

from scripts.macro_ingestion.semantic_guards import qoq_change_thousands_sa_plausible
from scripts.harvest_nz_pci import (
    BNZ_PUBLICATIONS,
    PSI_PDF_RE,
    canonical_pdf_url,
    parse_pci_table,
    parse_pdf_release_date,
    parse_headline_reference_period,
)

COUNTRY = "NZ"
STATS_BASE = "https://www.stats.govt.nz"

_MONTH_NAME_TO_NUM = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}

_RELEASE_HREF_RE = re.compile(
    r'href="(?P<path>/information-releases/(?P<slug>[a-z0-9-]+)-'
    r"(?P<month>january|february|march|april|may|june|july|august|september|october|november|december)-"
    r'(?P<year>\d{4})-quarter/?)"',
    re.IGNORECASE,
)
_EXPERIMENTAL_RELEASE_HREF_RE = re.compile(
    r'href="(?P<path>/experimental/national-accounts-income-saving-assets-and-liabilities-'
    r"(?P<month>january|february|march|april|may|june|july|august|september|october|november|december)-"
    r'(?P<year>\d{4})-quarter/?)"',
    re.IGNORECASE,
)

_CHALLENGE_MARKERS = (
    b"cf-browser-verification",
    b"challenge-platform",
    b"Just a moment",
    b"Attention Required! | Cloudflare",
)

_INFOSHARE_SERIES_REF: dict[tuple[str, str], str] = {
    ("CPIQ.SE9A", "yoy_pct"): "CPIQ.SE9AAC",
    ("CPIQ.SE9A", "qoq_pct"): "CPIQ.SE9APC",
}

_HLFS_COLUMN: dict[str, int] = {
    "HLFQ.S1F3S": 16,  # unemployment rate (total, SA)
    "HLFQ.S1A3S": 2,  # employed thousands
    "HLFQ.S1E3S": 12,  # participation rate
}

_FIXTURES = Path(__file__).resolve().parents[3] / "tests/macro_ingestion/fixtures/nz"


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _fail(
    error: str,
    *,
    http_status: int | None = None,
    status: str | None = None,
    challenge_page: bool = False,
) -> dict[str, Any]:
    out: dict[str, Any] = {
        "ok": False,
        "error": error,
        "http_status": http_status,
        "points": [],
        "raw_sha256": _sha256(error.encode()),
    }
    if status:
        out["status"] = status
    if challenge_page:
        out["challenge_page"] = True
    return out


def _ok(
    body: bytes,
    points: list[dict[str, Any]],
    *,
    http_status: int | None = 200,
    vintage: str = "latest_available",
    source_url: str | None = None,
) -> dict[str, Any]:
    if source_url:
        for pt in points:
            pt.setdefault("source_url", source_url)
    return {
        "ok": True,
        "http_status": http_status,
        "points": points,
        "vintage": vintage,
        "raw_sha256": _sha256(body),
        "body": body,
    }


def is_challenge_page(body: bytes) -> bool:
    if not body:
        return False
    sample = body[:8000]
    if sample[:4] == b"%PDF":
        return False
    lower = sample.lower()
    return any(marker.lower() in lower for marker in _CHALLENGE_MARKERS)


def _fetch_url(opener: Callable[..., dict[str, Any]], url: str, *, timeout: float) -> dict[str, Any]:
    return opener(url, timeout=timeout)


def _decode_html(body: bytes) -> str:
    return body.decode("utf-8", "replace")


def _parse_page_view_data(html: str) -> dict[str, Any] | None:
    match = re.search(r'id="pageViewData" data-value="(.+?)"', html)
    if not match:
        return None
    return json.loads(html_lib.unescape(match.group(1)))


def _document_links(page_data: dict[str, Any]) -> list[tuple[str, str]]:
    links: list[tuple[str, str]] = []
    for block in page_data.get("PageBlocks") or []:
        for doc in block.get("BlockDocuments") or []:
            href = doc.get("DocumentLink")
            name = doc.get("Name") or ""
            if href:
                links.append((urljoin(STATS_BASE, href), name))
    return links


def _pick_document(
    links: list[tuple[str, str]],
    *needles: str,
    exclude: tuple[str, ...] = (),
) -> str | None:
    lowered = [n.lower() for n in needles]
    excluded = [x.lower() for x in exclude]

    def _allowed(blob: str) -> bool:
        return not any(x in blob for x in excluded)

    for url, name in links:
        blob = f"{url} {name}".lower()
        if not _allowed(blob):
            continue
        if all(n in blob for n in lowered):
            return url
    for url, name in links:
        blob = f"{url} {name}".lower()
        if not _allowed(blob):
            continue
        if any(n in blob for n in lowered):
            return url
    return None


def _quarter_from_release_slug(month: str, year: str) -> str:
    month_num = _MONTH_NAME_TO_NUM[month.lower()]
    quarter = (month_num - 1) // 3 + 1
    return f"{int(year)}-Q{quarter}"


def _infer_stats_release_family(spec: dict[str, Any]) -> str:
    component_id = str(spec.get("id") or "")
    if component_id == "NZ.Consumer.income":
        return "national-accounts-income-saving-assets-and-liabilities"
    candidates: list[str] = []
    for raw in list(spec.get("registry_urls") or []) + [spec.get("endpoint") or ""]:
        text = str(raw).lower().strip().rstrip("/")
        if not text:
            continue
        match = re.search(r"/topics/([^/?#]+)", text)
        if match:
            candidates.append(match.group(1))
            continue
        match = re.search(r"/information-releases/([^/?#]+)", text)
        if match:
            slug = match.group(1)
            slug = re.sub(
                r"-(?:january|february|march|april|may|june|july|august|september|october|november|december)-\d{4}-quarter$",
                "",
                slug,
                flags=re.I,
            )
            candidates.append(slug)
            continue
        if "national-accounts-income-saving-assets-and-liabilities" in text:
            candidates.append("national-accounts-income-saving-assets-and-liabilities")
    return candidates[0] if candidates else ""


def discover_stats_nz_release_hrefs(
    html: str,
    *,
    family: str,
) -> list[tuple[str, str]]:
    """Return (absolute_url, observation_period) sorted newest period first."""
    family_key = family.lower().strip()
    found: dict[str, str] = {}
    for pattern in (_RELEASE_HREF_RE, _EXPERIMENTAL_RELEASE_HREF_RE):
        for match in pattern.finditer(html):
            if pattern is _RELEASE_HREF_RE:
                slug = match.group("slug").lower()
            else:
                slug = "national-accounts-income-saving-assets-and-liabilities"
            if family_key and not slug.startswith(family_key):
                continue
            period = _quarter_from_release_slug(match.group("month"), match.group("year"))
            url = urljoin(STATS_BASE, match.group("path"))
            found[url] = period
    return sorted(found.items(), key=lambda item: item[1], reverse=True)


def _looks_like_vintage_stats_release_url(url: str) -> bool:
    lower = url.lower()
    return bool(
        re.search(
            r"/(?:information-releases|experimental)/[a-z0-9-]+-"
            r"(?:january|february|march|april|may|june|july|august|september|october|november|december)-"
            r"\d{4}-quarter/?$",
            lower,
        )
    )


def _resolve_stats_release_url(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
) -> tuple[str, str | None]:
    """Resolve the current Stats NZ release page URL and optional listing period hint."""
    endpoint = str(spec.get("endpoint") or "").strip()
    family = _infer_stats_release_family(spec)
    discovered: list[tuple[str, str]] = []

    listing_urls: list[str] = []
    for raw in list(spec.get("registry_urls") or []) + ([endpoint] if endpoint else []):
        url = urljoin(STATS_BASE, raw) if str(raw).startswith("/") else str(raw)
        if url and url not in listing_urls:
            listing_urls.append(url)

    for listing_url in listing_urls:
        if _looks_like_vintage_stats_release_url(listing_url):
            continue
        page = _fetch_url(opener, listing_url, timeout=timeout)
        if not page.get("ok"):
            continue
        html = _decode_html(page.get("body") or b"")
        discovered.extend(discover_stats_nz_release_hrefs(html, family=family))

    if discovered:
        url, period = discovered[0]
        return url, period

    if endpoint:
        absolute = urljoin(STATS_BASE, endpoint) if endpoint.startswith("/") else endpoint
        return absolute, None
    return "", None


def _stats_period_label(period_cell: str) -> str | None:
    """Map Stats NZ quarter labels like Jun-26 to 2026-Q2."""
    text = period_cell.strip()
    m = re.match(r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-(\d{2})$", text, re.I)
    if not m:
        return None
    month = m.group(1).title()[:3]
    yy = int(m.group(2))
    year = 2000 + yy if yy < 70 else 1900 + yy
    quarter = {"Mar": 1, "Jun": 2, "Sep": 3, "Dec": 4}[month]
    return f"{year}-Q{quarter}"


def _quarter_from_yyyy_mm(period: str) -> str | None:
    m = re.match(r"^(\d{4})\.(\d{2})$", period.strip())
    if not m:
        return None
    year, month = int(m.group(1)), int(m.group(2))
    quarter = (month - 1) // 3 + 1
    return f"{year}-Q{quarter}"


def _infoshare_field(row: dict[str, Any], *keys: str) -> str:
    lowered = {str(key).lower(): value for key, value in row.items() if key}
    for key in keys:
        value = lowered.get(key.lower())
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return ""


def _parse_infoshare_csv(
    body: bytes,
    *,
    series_id: str,
    transform: str,
    source_url: str,
) -> list[dict[str, Any]]:
    text = body.decode("utf-8", "replace")
    reader = csv.DictReader(io.StringIO(text))
    target_ref = _INFOSHARE_SERIES_REF.get((series_id, transform), series_id)
    levels: list[tuple[str, float]] = []
    for row in reader:
        ref = _infoshare_field(row, "Series_Reference", "Series_reference")
        if ref != target_ref:
            continue
        period = _quarter_from_yyyy_mm(_infoshare_field(row, "Period"))
        value_raw = _infoshare_field(row, "Data_Value", "Data_value")
        if not period or not value_raw:
            continue
        levels.append((period, float(value_raw)))
    if not levels:
        return []
    levels.sort(key=lambda item: item[0])
    if transform == "qoq_change_thousands_sa" and series_id == "HLFQ.S1A3S":
        if len(levels) < 2:
            return []
        prior_period, prior_level = levels[-2]
        period, latest_level = levels[-1]
        value = latest_level - prior_level
        if not qoq_change_thousands_sa_plausible(
            value, prior_level=prior_level, latest_level=latest_level
        ):
            return []
        return [
            {
                "period": period,
                "value": value,
                "transformation": transform,
                "revision_status": "final",
                "source_url": source_url,
            }
        ]
    if transform == "yoy_pct" and series_id.startswith("LCIQ."):
        if len(levels) < 2:
            return []
        period, latest_level = levels[-1]
        year_s, quarter_s = period.split("-Q")
        prior_period = f"{int(year_s) - 1}-Q{quarter_s}"
        prior_level = dict(levels).get(prior_period)
        if prior_level is None or prior_level == 0:
            return []
        value = (latest_level / prior_level - 1.0) * 100.0
        return [
            {
                "period": period,
                "value": value,
                "transformation": transform,
                "revision_status": "final",
                "source_url": source_url,
            }
        ]
    period, value = levels[-1]
    return [
        {
            "period": period,
            "value": value,
            "transformation": transform,
            "revision_status": "final",
            "source_url": source_url,
        }
    ]


def _parse_cpi_xlsx_table(
    body: bytes,
    *,
    sheet: str,
    series_suffix: str,
    transform: str,
    source_url: str,
) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for idx, row in enumerate(rows):
        if row and str(row[0]).strip() == "Selected grouping":
            header_idx = idx + 1
            break
    if header_idx is None:
        return []
    header = rows[header_idx]
    # columns pairs: label at 2,4,6...
    col_periods: list[tuple[int, str]] = []
    for col in range(2, len(header), 2):
        label = header[col]
        if label is None:
            continue
        period = _stats_period_label(str(label))
        if period:
            col_periods.append((col, period))
    target_row = None
    for row in rows:
        if not row:
            continue
        ref = row[1]
        if ref == series_suffix:
            target_row = row
            break
    if target_row is None:
        return []
    points: list[dict[str, Any]] = []
    best: tuple[str, float] | None = None
    for col, period in col_periods:
        val = target_row[col]
        if val is None:
            continue
        best = (period, float(val))
    if best:
        period, value = best
        points.append(
            {
                "period": period,
                "value": value,
                "transformation": transform,
                "revision_status": "final",
                "source_url": source_url,
            }
        )
    return points


def _hlfs_block_bounds(rows: list[tuple[Any, ...]], suffix: str) -> tuple[int, int, int] | None:
    """Return (column_index, start_row, end_row) for one Table 1 section."""
    col: int | None = _HLFS_COLUMN.get(f"HLFQ.{suffix}")
    start: int | None = None
    end = len(rows)
    for idx, row in enumerate(rows):
        if not row or not str(row[0]).startswith("Series ref: HLFQ"):
            continue
        if start is not None:
            end = idx
            break
        for cell_idx, cell in enumerate(row):
            if cell == suffix:
                col = cell_idx
                start = idx + 1
                break
    if col is None or start is None:
        return None
    return col, start, end


def _hlfs_block_levels(
    rows: list[tuple[Any, ...]], *, col: int, start: int, end: int
) -> list[tuple[str, float]]:
    year: int | None = None
    observations: list[tuple[str, float]] = []
    for row in rows[start:end]:
        if not row:
            continue
        first = row[0]
        if first == "Quarter":
            continue
        if isinstance(first, str) and first.startswith(("1.", "2.", "3.")):
            break
        second = row[1] if len(row) > 1 else None
        if first in (2024, 2025, 2026, "2024", "2025", "2026"):
            year = int(first)
            if second not in ("Mar", "Jun", "Sep", "Dec"):
                continue
            month = second
        elif year is not None and second in ("Mar", "Jun", "Sep", "Dec"):
            month = second
        else:
            continue
        quarter = {"Mar": 1, "Jun": 2, "Sep": 3, "Dec": 4}[month]
        period = f"{year}-Q{quarter}"
        val = row[col] if len(row) > col else None
        if val is None:
            continue
        observations.append((period, float(val)))
    observations.sort(key=lambda item: item[0])
    return observations


def _parse_hlfs_xlsx(
    body: bytes,
    *,
    series_id: str,
    transform: str,
    source_url: str,
) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    ws = wb["Table 1"] if "Table 1" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    suffix = series_id.split(".")[-1] if series_id.startswith("HLFQ.") else ""
    if not suffix:
        return []
    bounds = _hlfs_block_bounds(rows, suffix)
    if bounds is None:
        return []
    col, start, end = bounds
    observations = _hlfs_block_levels(rows, col=col, start=start, end=end)
    if not observations:
        return []
    period, latest_level = observations[-1]
    if series_id == "HLFQ.S1A3S" and transform == "qoq_change_thousands_sa":
        if len(observations) < 2:
            return []
        _, prior_level = observations[-2]
        value = latest_level - prior_level
        if not qoq_change_thousands_sa_plausible(
            value, prior_level=prior_level, latest_level=latest_level
        ):
            return []
    else:
        value = latest_level
    return [
        {
            "period": period,
            "value": value,
            "transformation": transform,
            "revision_status": "final",
            "source_url": source_url,
        }
    ]


def _period_from_year_month(year: int, month_label: str) -> str | None:
    month = month_label.strip()
    if month.lower().startswith("sept"):
        month = "Sep"
    month = month.title()[:3]
    if month not in {"Mar", "Jun", "Sep", "Dec"}:
        return None
    quarter = {"Mar": 1, "Jun": 2, "Sep": 3, "Dec": 4}[month]
    return f"{year}-Q{quarter}"


def _find_series_column(rows: list[tuple[Any, ...]], series_ref: str) -> int | None:
    target = series_ref.split(".")[-1] if "." in series_ref else series_ref
    for row in rows:
        if not row:
            continue
        for idx, cell in enumerate(row):
            if cell in {series_ref, target}:
                return idx
    return None


def _parse_year_quarter_table_levels(
    rows: list[tuple[Any, ...]],
    *,
    series_ref: str,
    month_col: int = 2,
) -> list[tuple[str, float]]:
    col = _find_series_column(rows, series_ref)
    if col is None:
        return []
    year: int | None = None
    observations: list[tuple[str, float]] = []
    for row in rows:
        if not row or len(row) <= col:
            continue
        first = row[0]
        if isinstance(first, (int, float)) and 1900 <= int(first) <= 2100:
            year = int(first)
        elif isinstance(first, str) and first.isdigit() and len(first) == 4:
            year = int(first)
        month_cell = row[month_col] if len(row) > month_col else None
        if year is None or not isinstance(month_cell, str):
            continue
        period = _period_from_year_month(year, month_cell)
        if not period:
            continue
        val = row[col]
        if isinstance(val, (int, float)):
            observations.append((period, float(val)))
    observations.sort(key=lambda item: item[0])
    return observations


def _parse_retail_table_11(
    body: bytes,
    *,
    series_id: str,
    transform: str,
    source_url: str,
) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    sheet = "Table 11" if "Table 11" in wb.sheetnames else wb.sheetnames[-1]
    rows = list(wb[sheet].iter_rows(values_only=True))
    levels = _parse_year_quarter_table_levels(rows, series_ref=series_id)
    if len(levels) < 2:
        return []
    period, latest = levels[-1]
    if transform == "qoq_pct":
        _, prior = levels[-2]
        if prior == 0:
            return []
        value = (latest / prior - 1.0) * 100.0
    else:
        value = latest
    return [
        {
            "period": period,
            "value": value,
            "transformation": transform,
            "revision_status": "final",
            "source_url": source_url,
        }
    ]


def _parse_gdp_supplementary_table_4(
    body: bytes,
    *,
    series_ref: str,
    transform: str,
    source_url: str,
) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if "Table 4" not in wb.sheetnames:
        return []
    rows = list(wb["Table 4"].iter_rows(values_only=True))
    observations = _parse_year_quarter_table_levels(rows, series_ref=series_ref)
    if not observations:
        return []
    period, value = observations[-1]
    return [
        {
            "period": period,
            "value": value,
            "transformation": transform,
            "revision_status": "final",
            "source_url": source_url,
        }
    ]


def _pick_zip_csv_member(names: list[str], *, series_id: str) -> str | None:
    if series_id.startswith("HLFQ."):
        for name in names:
            lower = name.lower()
            if "hlfs" in lower and lower.endswith(".csv"):
                return name
    if series_id.startswith("LCIQ."):
        for name in names:
            lower = name.lower()
            if "lci" in lower and lower.endswith(".csv"):
                return name
    return None


def _parse_experimental_income_html(html: str, *, source_url: str, transform: str) -> list[dict[str, Any]]:
    text = html_lib.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html)))
    period: str | None = None
    slug_match = re.search(
        r"national-accounts-income-saving-assets-and-liabilities-"
        r"(january|february|march|april|may|june|july|august|september|october|november|december)-(\d{4})-quarter",
        source_url,
        re.I,
    )
    if slug_match:
        period = _quarter_from_release_slug(slug_match.group(1), slug_match.group(2))
    if not period:
        page_data = _parse_page_view_data(html)
        if page_data:
            page_date = str(page_data.get("PageDate") or "")
            date_match = re.match(r"(\d{4})-(\d{2})-\d{2}", page_date)
            if date_match:
                year = int(date_match.group(1))
                month = int(date_match.group(2))
                period = f"{year}-Q{(month - 1) // 3 + 1}"
    value_match = re.search(
        r"household net disposable income[^0-9\-]{0,120}"
        r"(-?\d+(?:\.\d+)?)\s*(?:percent|%)",
        text,
        re.I,
    )
    if not value_match or not period:
        return []
    return [
        {
            "period": period,
            "value": float(value_match.group(1)),
            "transformation": transform,
            "revision_status": "preliminary",
            "source_url": source_url,
        }
    ]


def _parse_stats_document(
    raw: bytes,
    *,
    data_url: str,
    series_id: str,
    transform: str,
    retrieval: str,
) -> list[dict[str, Any]]:
    if data_url.lower().endswith(".csv"):
        return _parse_infoshare_csv(
            raw, series_id=series_id, transform=transform, source_url=data_url
        )
    if data_url.lower().endswith(".xlsx"):
        if series_id.startswith("HLFQ."):
            return _parse_hlfs_xlsx(
                raw, series_id=series_id, transform=transform, source_url=data_url
            )
        if series_id == "SF91KS" or retrieval == "information_release_xlsx_table_11":
            return _parse_retail_table_11(
                raw, series_id=series_id, transform=transform, source_url=data_url
            )
        if series_id == "SG02RSC30P30E" or retrieval == "gdp_supplementary_xlsx_table_4":
            return _parse_gdp_supplementary_table_4(
                raw, series_ref=series_id, transform=transform, source_url=data_url
            )
        if series_id == "CPIQ.SE9NS1450":
            return _parse_cpi_xlsx_table(
                raw,
                sheet="3.03",
                series_suffix="SE9NS1450",
                transform=transform,
                source_url=data_url,
            )
        if retrieval == "information_release_xlsx_tables_3.02_3.03" and transform == "yoy_pct":
            return _parse_cpi_xlsx_table(
                raw,
                sheet="3.03",
                series_suffix="SE9A",
                transform=transform,
                source_url=data_url,
            )
        return _parse_cpi_xlsx_table(
            raw,
            sheet="3.02" if transform == "qoq_pct" else "3.03",
            series_suffix=series_id.split(".")[-1] if "." in series_id else series_id,
            transform=transform,
            source_url=data_url,
        )
    if data_url.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            member = _pick_zip_csv_member(zf.namelist(), series_id=series_id)
            if member:
                csv_body = zf.read(member)
                return _parse_infoshare_csv(
                    csv_body, series_id=series_id, transform=transform, source_url=data_url
                )
    return []


def _fetch_stats_release(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
    csv_needles: tuple[str, ...] | None = None,
    xlsx_needles: tuple[str, ...] | None = None,
    zip_needles: tuple[str, ...] | None = None,
    hlfs_xlsx_needles: tuple[str, ...] | None = None,
) -> dict[str, Any]:
    release_url, _listing_period = _resolve_stats_release_url(spec, opener=opener, timeout=timeout)
    if not release_url:
        return _fail("missing_endpoint", status="source_failed")
    page = _fetch_url(opener, release_url, timeout=timeout)
    if not page.get("ok"):
        return _fail(page.get("error") or "endpoint_fetch_failed", http_status=page.get("http_status"))
    body = page.get("body") or b""
    if is_challenge_page(body):
        return _fail("challenge_page", status="license_gap", challenge_page=True, http_status=page.get("http_status"))
    page_data = _parse_page_view_data(_decode_html(body))
    if not page_data:
        return _fail("release_page_unparsed", http_status=page.get("http_status"))
    links = _document_links(page_data)
    series_id = str(spec.get("series_id") or spec["id"])
    transform = str(spec.get("transform") or "")
    retrieval = str(spec.get("retrieval_method") or "")

    data_url: str | None = None
    if csv_needles or retrieval.endswith("infoshare_csv"):
        data_url = _pick_document(
            links,
            "infoshare",
            "csv",
            exclude=("visualisation", "visualization"),
        )
        if not data_url and csv_needles:
            data_url = _pick_document(
                links,
                *csv_needles,
                exclude=("visualisation", "visualization"),
            )
        if not data_url and retrieval.endswith("infoshare_csv"):
            data_url = _pick_document(
                links,
                "gross-domestic-product",
                "quarter.csv",
                exclude=("visualisation", "visualization", "supplementary", "current-price"),
            )
            if not data_url:
                data_url = _pick_document(
                    links,
                    "csv",
                    exclude=("visualisation", "visualization", "supplementary"),
                )
    if not data_url and xlsx_needles:
        data_url = _pick_document(links, *xlsx_needles)
    if not data_url and hlfs_xlsx_needles:
        data_url = _pick_document(links, *hlfs_xlsx_needles)
    if not data_url and zip_needles:
        data_url = _pick_document(links, *zip_needles)

    if not data_url:
        return _fail("primary_document_not_found", http_status=page.get("http_status"))

    data_resp = _fetch_url(opener, data_url, timeout=timeout)
    if not data_resp.get("ok"):
        return _fail(data_resp.get("error") or "data_fetch_failed", http_status=data_resp.get("http_status"))
    raw = data_resp.get("body") or b""
    points = _parse_stats_document(
        raw,
        data_url=data_url,
        series_id=series_id,
        transform=transform,
        retrieval=retrieval,
    )
    if not points:
        return _fail("empty_parse", http_status=data_resp.get("http_status"))
    payload = _ok(raw, points, http_status=data_resp.get("http_status"), source_url=data_url)
    payload["release_page_url"] = release_url
    return payload


def _fetch_pci(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
) -> dict[str, Any]:
    biz_url = "https://businessnz.org.nz/psi/"
    biz = _fetch_url(opener, biz_url, timeout=timeout)
    biz_body = biz.get("body") or b""
    businessnz_challenge = bool(biz_body and is_challenge_page(biz_body))

    pub = _fetch_url(opener, BNZ_PUBLICATIONS, timeout=timeout)
    pdf_urls: list[str] = []
    if pub.get("ok"):
        html = _decode_html(pub.get("body") or b"")
        pdf_urls = sorted({canonical_pdf_url(m.group(1)) for m in PSI_PDF_RE.finditer(html)})
    if not pdf_urls:
        pdf_urls = [
            "https://www.bnz.co.nz/assets/markets/research/BNZ-BusinessNZ-PSI-Release_August-2026_web.pdf",
        ]

    parsed: list[dict[str, Any]] = []
    last_body = b""
    for url in reversed(pdf_urls):
        resp = _fetch_url(opener, url, timeout=timeout)
        if not resp.get("ok"):
            continue
        data = resp.get("body") or b""
        if data[:4] != b"%PDF":
            continue
        last_body = data
        text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
        pci = parse_pci_table(text)
        headline = parse_headline_reference_period(text)
        if not pci or not headline:
            continue
        value = pci.get(headline)
        if value is None:
            continue
        parsed.append(
            {
                "period": headline,
                "value": value,
                "transformation": spec.get("transform"),
                "revision_status": "final",
                "source_url": url,
            }
        )
        break
    if not parsed:
        if businessnz_challenge:
            return _fail(
                "businessnz_cloudflare_challenge",
                status="license_gap",
                challenge_page=True,
                http_status=biz.get("http_status"),
            )
        return _fail("pci_pdf_unparsed", status="source_failed")
    payload = _ok(last_body, parsed, source_url=parsed[0]["source_url"])
    if businessnz_challenge:
        payload["businessnz_challenge_page"] = True
    return payload


def _parse_anz_confidence(html: str) -> list[dict[str, Any]]:
    text = re.sub(r"<[^>]+>", " ", html)
    text = html_lib.unescape(re.sub(r"\s+", " ", text))
    patterns = [
        re.compile(
            r"consumer confidence(?: index)?[^0-9]{0,40}(\d{2,3}\.\d)\b",
            re.IGNORECASE,
        ),
        re.compile(r"confidence (?:index )?(?:at|to|of) (\d{2,3}\.\d)", re.IGNORECASE),
        re.compile(r"ANZ-Roy Morgan Consumer Confidence[^0-9]{0,20}(\d{2,3}\.\d)", re.IGNORECASE),
    ]
    for pat in patterns:
        match = pat.search(text)
        if match:
            value = float(match.group(1))
            month_match = re.search(
                r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{4})\b",
                text,
                re.I,
            )
            period = "2026-08"
            if month_match:
                month_name = month_match.group(1).title()
                year = int(month_match.group(2))
                month_num = datetime.strptime(month_name[:3], "%b").month
                period = f"{year}-{month_num:02d}"
            return [
                {
                    "period": period,
                    "value": value,
                    "transformation": "index_level",
                    "revision_status": "final",
                }
            ]
    return []


def _fetch_anz_confidence(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
) -> dict[str, Any]:
    endpoint = str(spec.get("endpoint") or "")
    resp = _fetch_url(opener, endpoint, timeout=timeout)
    if not resp.get("ok"):
        return _fail(resp.get("error") or "anz_fetch_failed", http_status=resp.get("http_status"))
    body = resp.get("body") or b""
    points = _parse_anz_confidence(_decode_html(body))
    if not points:
        return _fail("anz_index_not_parseable", status="source_failed", http_status=resp.get("http_status"))
    for pt in points:
        pt["source_url"] = endpoint
    return _ok(body, points, http_status=resp.get("http_status"), source_url=endpoint)


def _fetch_experimental_income(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
) -> dict[str, Any]:
    release_url, _ = _resolve_stats_release_url(spec, opener=opener, timeout=timeout)
    if not release_url:
        return _fail("missing_endpoint", status="source_failed")
    resp = _fetch_url(opener, release_url, timeout=timeout)
    if not resp.get("ok"):
        return _fail(resp.get("error") or "income_release_fetch_failed", http_status=resp.get("http_status"))
    body = resp.get("body") or b""
    if is_challenge_page(body):
        return _fail("challenge_page", status="license_gap", challenge_page=True)
    html = _decode_html(body)
    points = _parse_experimental_income_html(
        html,
        source_url=release_url,
        transform=str(spec.get("transform") or "qoq_pct"),
    )
    if not points:
        return _fail("experimental_income_unparsed", status="source_failed", http_status=resp.get("http_status"))
    return _ok(body, points, http_status=resp.get("http_status"), source_url=release_url)


def _fetch_context_unpinned(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
) -> dict[str, Any]:
    release_url, _ = _resolve_stats_release_url(spec, opener=opener, timeout=timeout)
    if not release_url:
        release_url = str(spec.get("endpoint") or "")
    resp = _fetch_url(opener, release_url, timeout=timeout)
    if not resp.get("ok"):
        return _fail(
            resp.get("error") or "context_endpoint_failed",
            status="source_failed",
            http_status=resp.get("http_status"),
        )
    body = resp.get("body") or b""
    if is_challenge_page(body):
        return _fail("challenge_page", status="license_gap", challenge_page=True)
    if spec["id"] == "NZ.Labor.participation":
        page_data = _parse_page_view_data(_decode_html(body))
        if page_data:
            xlsx = _pick_document(_document_links(page_data), "household-labour-force-survey", "xlsx")
            if xlsx:
                data_resp = _fetch_url(opener, xlsx, timeout=timeout)
                if data_resp.get("ok"):
                    points = _parse_hlfs_xlsx(
                        data_resp.get("body") or b"",
                        series_id="HLFQ.S1E3S",
                        transform="percent",
                        source_url=xlsx,
                    )
                    if points:
                        payload = _ok(data_resp.get("body") or b"", points, source_url=xlsx)
                        payload["pinned_series_id"] = "HLFQ.S1E3S"
                        return payload
    return _fail(
        f"series_id_unpinned_for_{spec['id']}",
        status="not_applicable",
        http_status=resp.get("http_status"),
    )


def fetch_series(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    now: datetime,
    timeout: float = 20,
) -> dict[str, Any]:
    """Return a JSON-ready payload for one catalog row."""
    role = str(spec.get("role") or "")
    if role == "explanatory_alias":
        return _fail("explanatory_alias_skip", status="not_applicable")

    series_id = spec.get("series_id")
    retrieval = str(spec.get("retrieval_method") or "")
    component_id = str(spec["id"])

    if component_id == "NZ.Activity.gdp_expenditure":
        return _fail("context_expenditure_gdp_not_primary", status="not_applicable")

    if retrieval == "harvest_nz_pci_publications_psi_pdf" or series_id == "BUSINESSNZ_PCI_GDP_WEIGHTED":
        return _fetch_pci(spec, opener=opener, timeout=timeout)

    if retrieval == "official_web_page_text" and component_id == "NZ.Consumer.confidence":
        return _fetch_anz_confidence(spec, opener=opener, timeout=timeout)

    if role == "context" and series_id is None:
        return _fetch_context_unpinned(spec, opener=opener, timeout=timeout)

    if series_id in {"CPIQ.SE9A", "CPIQ.SE9NS1450"}:
        needles = ("consumers-price-index", "xlsx") if series_id == "CPIQ.SE9NS1450" else ("infoshare", "csv")
        return _fetch_stats_release(
            spec,
            opener=opener,
            timeout=timeout,
            csv_needles=needles if series_id == "CPIQ.SE9A" else None,
            xlsx_needles=needles if series_id == "CPIQ.SE9NS1450" else None,
        )

    if series_id and str(series_id).startswith("HLFQ."):
        return _fetch_stats_release(
            spec,
            opener=opener,
            timeout=timeout,
            hlfs_xlsx_needles=("household-labour-force-survey", "xlsx"),
            zip_needles=("labour-market-statistics", "zip"),
        )

    if series_id == "LCIQ.SD53Z9":
        return _fetch_stats_release(
            spec,
            opener=opener,
            timeout=timeout,
            zip_needles=("labour-market-statistics", "zip"),
        )

    if series_id and str(series_id).startswith("SNEQ."):
        return _fetch_stats_release(
            spec,
            opener=opener,
            timeout=timeout,
            csv_needles=("infoshare", "csv"),
        )

    if series_id == "SF91KS":
        return _fetch_stats_release(
            spec,
            opener=opener,
            timeout=timeout,
            xlsx_needles=("retail-trade", "xlsx"),
        )

    if series_id == "household_net_disposable_income_sa":
        return _fetch_experimental_income(spec, opener=opener, timeout=timeout)

    if series_id == "SG02RSC30P30E":
        return _fetch_stats_release(
            spec,
            opener=opener,
            timeout=timeout,
            xlsx_needles=("supplementary", "xlsx"),
        )

    if series_id == "BUSINESSNZ_PMI":
        return _fail("manufacturing_pmi_not_implemented", status="license_gap", challenge_page=True)

    endpoint = str(spec.get("endpoint") or "")
    if endpoint:
        resp = _fetch_url(opener, endpoint, timeout=timeout)
        return _fail(
            resp.get("error") or "unsupported_series",
            status="source_failed",
            http_status=resp.get("http_status"),
        )
    return _fail("unsupported_series", status="source_failed")
