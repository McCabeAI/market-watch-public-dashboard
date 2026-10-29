"""Australia macro ingestion adapter (ABS workbooks, S&P PMI, housing context)."""

from __future__ import annotations

import hashlib
import io
import re
from calendar import monthrange
from datetime import date, datetime
from typing import Any, Callable

from scripts.australia_housing_data import collect_australia_housing
from scripts.macro_ingestion.calendar import latest_due_release
from scripts.macro_ingestion.au_labor_qa import (
    TABLE1_SIBLING_SERIES,
    format_identity_mismatch,
    labour_siblings_present,
    validate_table1_labour_point,
)
from scripts.harvest_au_pmi import (
    PMI_LISTING_URL,
    SEED_GUIDS,
    choose_observations,
    discover_australia_listing,
    is_flash_release,
    parse_release_artifact,
    press_release_url,
)

COUNTRY = "AU"

_CHALLENGE_MARKERS = (
    b"access denied",
    b"cf-chl",
    b"challenge-platform",
    b"just a moment",
    b"enable javascript",
)
_HTML_PREFIX = re.compile(rb"^\s*<(!DOCTYPE|html\b)", re.I)

RETAIL_SERIES_CEASED_AFTER = "2025-06"

_CONTEXT_SERIES_ENDPOINTS: dict[str, str] = {
    "AU.Inflation.ppi": (
        "https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/"
        "producer-price-indexes-australia/jun-2026/642701.xlsx"
    ),
    "AU.Inflation.import_price_index": (
        "https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/"
        "international-trade-price-indexes-australia/jun-2026/645701.xlsx"
    ),
    "AU.Inflation.export_price_index": (
        "https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/"
        "international-trade-price-indexes-australia/jun-2026/645704.xlsx"
    ),
    "AU.Labor.participation": (
        "https://www.abs.gov.au/statistics/labour/employment-and-unemployment/"
        "labour-force-australia/jul-2026/62020001.xlsx"
    ),
}

_CONTEXT_SERIES_IDS: dict[str, str] = {
    "AU.Inflation.ppi": "A2314865F",
    "AU.Inflation.import_price_index": "A2295765J",
    "AU.Inflation.export_price_index": "A2294886K",
    "AU.Labor.participation": "A84423051C",
}

_LABOUR_WORKBOOK_IDS = {
    "AU.Labor.unemployment",
    "AU.Labor.employment",
    "AU.Labor.participation",
}

AU_LFS_MODERNISATION_METHODOLOGY_NOTE = (
    "August 2026 Labour Force Survey modernisation may affect comparability "
    "with earlier periods; informational context only — scored series unchanged."
)

_LABOUR_TABLE1_WORKBOOK_MARKER = "62020001.xlsx"

# National accounts time-series workbooks (temperature_history/au.json source_urls).
_ABS_NA_WORKBOOK_BY_SERIES: dict[str, str] = {
    "A2304370T": "5206001_Key_Aggregates.xlsx",
    "A2304183K": "5206002_Expenditure_Volume_Measures.xlsx",
}

_workbook_bytes_cache: dict[str, dict[str, Any]] = {}


def clear_workbook_cache() -> None:
    """Drop cached ABS workbook bytes (for test isolation)."""
    _workbook_bytes_cache.clear()

_ENGLISH_MONTH_SLUGS = (
    "jan",
    "feb",
    "mar",
    "apr",
    "may",
    "jun",
    "jul",
    "aug",
    "sep",
    "oct",
    "nov",
    "dec",
)

_DATED_WORKBOOK_SLUG_RE = re.compile(
    r"/(" + "|".join(_ENGLISH_MONTH_SLUGS) + r")-(\d{4})/",
    re.IGNORECASE,
)
_MONTHLY_PERIOD_RE = re.compile(r"(20\d{2})-(\d{2})")


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _base_payload(**kwargs: Any) -> dict[str, Any]:
    return {
        "ok": False,
        "status": None,
        "points": [],
        "vintage": None,
        "raw_sha256": None,
        "error": None,
        "http_status": None,
        "challenge_page": False,
        "body": b"",
        **kwargs,
    }


def _looks_like_challenge(body: bytes, *, expect: str | None = None) -> bool:
    if not body:
        return False
    lower = body[:8000].lower()
    if any(marker in lower for marker in _CHALLENGE_MARKERS):
        return True
    if expect == "xlsx" and _HTML_PREFIX.match(body):
        return True
    if expect == "pdf" and _HTML_PREFIX.match(body):
        return True
    return False


def _period_from_cell(raw: Any, cadence: str) -> str | None:
    if isinstance(raw, datetime):
        dt = raw.date()
    elif isinstance(raw, date):
        dt = raw
    else:
        text = str(raw or "").strip()
        if not text:
            return None
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                dt = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                dt = None
        if dt is None:
            return None
    if cadence == "quarterly":
        quarter = (dt.month - 1) // 3 + 1
        return f"{dt.year:04d}-Q{quarter}"
    return f"{dt.year:04d}-{dt.month:02d}"


def _quarter_end_month(period: str) -> int:
    return int(period.split("-Q")[1]) * 3


def _parse_series_column(
    rows: list[tuple[Any, ...]],
    series_id: str,
    *,
    cadence: str,
) -> list[tuple[str, float]]:
    if len(rows) < 11:
        raise ValueError("workbook_too_short")

    header_ids = rows[9]
    col_index = None
    for idx, cell in enumerate(header_ids):
        if cell == series_id:
            col_index = idx
            break
    if col_index is None:
        raise ValueError(f"series_id_not_in_workbook:{series_id}")

    if header_ids[col_index] != series_id:
        raise ValueError(f"series_id_header_mismatch:{series_id}")

    out: list[tuple[str, float]] = []
    for row in rows[10:]:
        if not row or row[0] is None:
            continue
        period = _period_from_cell(row[0], cadence)
        if period is None:
            continue
        raw_val = row[col_index] if col_index < len(row) else None
        if raw_val is None:
            continue
        try:
            value = float(raw_val)
        except (TypeError, ValueError):
            continue
        out.append((period, value))
    if not out:
        raise ValueError("no_observations_parsed")
    return out


def parse_abs_time_series_workbook(
    data: bytes,
    series_id: str,
    *,
    cadence: str,
) -> list[tuple[str, float]]:
    """Parse ABS 'Time Series Workbook' Data* sheets for one Series ID column."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency in requirements-market-state
        raise RuntimeError("openpyxl required for ABS workbook parse") from exc

    if data[:2] != b"PK":
        raise ValueError("not_xlsx_zip")

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    data_sheet_names = [name for name in wb.sheetnames if name.startswith("Data")]
    if not data_sheet_names:
        raise ValueError("missing_data_sheet")

    last_exc: ValueError | None = None
    for sheet_name in data_sheet_names:
        data_sheet = wb[sheet_name]
        rows = list(data_sheet.iter_rows(values_only=True))
        try:
            return _parse_series_column(rows, series_id, cadence=cadence)
        except ValueError as exc:
            if str(exc) == f"series_id_not_in_workbook:{series_id}":
                last_exc = exc
                continue
            raise
    if last_exc is not None:
        raise last_exc
    raise ValueError(f"series_id_not_in_workbook:{series_id}")


def parse_abs_workbook_columns(
    data: bytes,
    series_ids: set[str] | frozenset[str],
    *,
    cadence: str,
) -> dict[str, list[tuple[str, float]]]:
    """Parse multiple series id columns present anywhere in Data* sheets."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl required for ABS workbook parse") from exc

    if data[:2] != b"PK":
        raise ValueError("not_xlsx_zip")

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    remaining = set(series_ids)
    found: dict[str, list[tuple[str, float]]] = {}
    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("Data") or not remaining:
            continue
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        if len(rows) < 11:
            continue
        header_ids = rows[9]
        for sid in list(remaining):
            if sid not in header_ids:
                continue
            try:
                found[sid] = _parse_series_column(rows, sid, cadence=cadence)
                remaining.discard(sid)
            except ValueError:
                continue
    return found


def _is_labour_table1_workbook_url(url: str) -> bool:
    return _LABOUR_TABLE1_WORKBOOK_MARKER in (url or "")


def _derive_mom_change_thousands(levels: list[tuple[str, float]]) -> list[tuple[str, float]]:
    derived: list[tuple[str, float]] = []
    for i in range(1, len(levels)):
        period, value = levels[i]
        prior = levels[i - 1][1]
        derived.append((period, float(value - prior)))
    return derived


def _derive_qoq_pct(levels: list[tuple[str, float]]) -> list[tuple[str, float]]:
    derived: list[tuple[str, float]] = []
    for i in range(1, len(levels)):
        period, value = levels[i]
        prior = levels[i - 1][1]
        if prior == 0:
            continue
        derived.append((period, (value - prior) / prior * 100.0))
    return derived


def _latest_point(
    series: list[tuple[str, float]],
    *,
    transform: str,
    source_url: str,
    max_period: str | None = None,
) -> dict[str, Any] | None:
    if not series:
        return None
    filtered = series
    if max_period is not None:
        filtered = [(p, v) for p, v in series if p <= max_period]
    if not filtered:
        return None
    period, value = filtered[-1]
    return {
        "period": period,
        "value": value,
        "transformation": transform,
        "revision_status": "final",
        "source_url": source_url,
    }


def _resolve_series_id(spec: dict[str, Any]) -> str | None:
    sid = spec.get("series_id")
    if not sid:
        return _CONTEXT_SERIES_IDS.get(spec["id"])
    text = str(sid)
    if ";" in text:
        return text.split(";")[0].strip()
    return text


def _national_accounts_workbook_url(url: str, series_id: str) -> str:
    """Route GDP vs domestic-final-demand ABS series to the workbook that contains them."""
    target = _ABS_NA_WORKBOOK_BY_SERIES.get(series_id)
    if not target or not url:
        return url
    for filename in _ABS_NA_WORKBOOK_BY_SERIES.values():
        if filename in url:
            return url.replace(filename, target)
    if url.endswith("/"):
        return url + target
    return url


def _month_slug_for_period(period: str) -> str | None:
    match = _MONTHLY_PERIOD_RE.fullmatch(period.strip())
    if not match:
        return None
    year = int(match.group(1))
    month = int(match.group(2))
    if month < 1 or month > 12:
        return None
    return f"{_ENGLISH_MONTH_SLUGS[month - 1]}-{year}"


def _resolve_abs_workbook_endpoint(
    spec: dict[str, Any],
    *,
    now: datetime | None = None,
) -> tuple[str | None, str | None, bool]:
    """Return (url, expected_period, release_due_unresolved).

    When a monthly release is due with a bound YYYY-MM period, resolve the official
    `{mon-yyyy}` workbook URL. Otherwise keep the catalog static endpoint.
    """
    spec_id = str(spec.get("id") or "")
    static = spec.get("endpoint") or _CONTEXT_SERIES_ENDPOINTS.get(spec_id)
    if now is None:
        return static, None, False

    due = latest_due_release(spec, now)
    if not due:
        return static, None, False

    period = str(due.get("period") or "")
    if not _MONTHLY_PERIOD_RE.fullmatch(period):
        return static, None, False

    slug = _month_slug_for_period(period)
    if not slug:
        return None, period, True

    if static and _DATED_WORKBOOK_SLUG_RE.search(static):
        url = _DATED_WORKBOOK_SLUG_RE.sub(f"/{slug}/", static, count=1)
        return url, period, False

    if spec_id in _LABOUR_WORKBOOK_IDS:
        registry = list(spec.get("registry_urls") or [])
        if registry:
            url = registry[0].rstrip("/") + f"/{slug}/62020001.xlsx"
            return url, period, False

    return None, period, True


def _point_for_period(
    series: list[tuple[str, float]],
    period: str,
    *,
    transform: str,
    source_url: str,
) -> dict[str, Any] | None:
    chosen: dict[str, Any] | None = None
    for p, value in series:
        if p != period:
            continue
        chosen = {
            "period": p,
            "value": value,
            "transformation": transform,
            "revision_status": "final",
            "source_url": source_url,
        }
    return chosen


def _fetch_workbook_bytes(
    url: str,
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
) -> dict[str, Any]:
    cached = _workbook_bytes_cache.get(url)
    if cached is not None:
        return dict(cached)

    resp = opener(url, timeout=timeout)
    body = resp.get("body") or b""
    if _looks_like_challenge(body, expect="xlsx"):
        return _base_payload(
            ok=False,
            status="license_gap",
            challenge_page=True,
            http_status=resp.get("http_status"),
            error="challenge_page",
            raw_sha256=_sha256(body) if body else None,
            body=body,
            url=resp.get("url") or url,
        )
    if not resp.get("ok") or not body:
        http_status = resp.get("http_status")
        error = resp.get("error") or "fetch_failed"
        if http_status == 404 and "404" not in str(error):
            error = "HTTP 404"
        return _base_payload(
            ok=False,
            status="source_failed",
            http_status=http_status,
            error=error,
            body=body,
        )
    if body[:2] != b"PK":
        return _base_payload(
            ok=False,
            status="source_failed",
            http_status=resp.get("http_status"),
            error="not_xlsx",
            raw_sha256=_sha256(body),
            body=body,
        )
    result = {
        "ok": True,
        "body": body,
        "http_status": resp.get("http_status"),
        "error": None,
        "raw_sha256": _sha256(body),
        "url": resp.get("url") or url,
    }
    _workbook_bytes_cache[url] = result
    return dict(result)


def _fetch_abs_workbook_series(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    timeout: float,
    now: datetime | None = None,
    max_period: str | None = None,
) -> dict[str, Any]:
    series_id = _resolve_series_id(spec)
    url, expected_period, unresolved = _resolve_abs_workbook_endpoint(spec, now=now)
    if unresolved:
        return _base_payload(
            ok=False,
            status="source_failed",
            error=f"release_due_unresolved:expected_period={expected_period}",
        )
    if not series_id or not url:
        return _base_payload(
            ok=False,
            status="source_failed",
            error="missing_series_or_endpoint",
        )

    url = _national_accounts_workbook_url(url, series_id)

    fetched = _fetch_workbook_bytes(url, opener=opener, timeout=timeout)
    if not fetched.get("ok"):
        return _base_payload(**{k: fetched[k] for k in fetched if k != "url"})

    cadence = str(spec.get("cadence") or "monthly")
    if cadence == "ceased":
        cadence = "monthly"

    try:
        levels = parse_abs_time_series_workbook(fetched["body"], series_id, cadence=cadence)
    except ValueError as exc:
        return _base_payload(
            ok=False,
            status="source_failed",
            error=str(exc),
            http_status=fetched.get("http_status"),
            raw_sha256=fetched.get("raw_sha256"),
            body=fetched.get("body") or b"",
        )

    transform = str(spec.get("transform") or "")
    retrieval = str(spec.get("retrieval_method") or "")

    if transform == "mom_change_thousands_sa":
        series = _derive_mom_change_thousands(levels)
    elif transform == "qoq_pct" and retrieval == "abs_time_series_workbook_derived_qoq":
        series = _derive_qoq_pct(levels)
    else:
        series = levels

    if expected_period is not None:
        point = _point_for_period(
            series,
            expected_period,
            transform=transform,
            source_url=url,
        )
        qa_period = expected_period
        if point is None:
            observed = series[-1][0] if series else None
            return _base_payload(
                ok=False,
                status="source_failed",
                error=(
                    f"release_due_stale_workbook:expected={expected_period}:observed={observed}"
                ),
                raw_sha256=fetched.get("raw_sha256"),
                body=fetched.get("body") or b"",
                http_status=fetched.get("http_status"),
            )
    else:
        point = _latest_point(
            series,
            transform=transform,
            source_url=url,
            max_period=max_period,
        )
        qa_period = point["period"] if point else None
        if point is None:
            return _base_payload(
                ok=False,
                status="source_failed",
                error="no_point_after_filters",
                raw_sha256=fetched.get("raw_sha256"),
                body=fetched.get("body") or b"",
                http_status=fetched.get("http_status"),
            )

    if (
        qa_period
        and _is_labour_table1_workbook_url(url)
        and series_id in TABLE1_SIBLING_SERIES
    ):
        try:
            sibling_cols = parse_abs_workbook_columns(
                fetched["body"],
                TABLE1_SIBLING_SERIES,
                cadence=cadence,
            )
        except ValueError:
            sibling_cols = {}
        if labour_siblings_present(sibling_cols):
            mismatch = validate_table1_labour_point(qa_period, sibling_cols)
            if mismatch:
                return _base_payload(
                    ok=False,
                    status="source_failed",
                    error=format_identity_mismatch(mismatch),
                    raw_sha256=fetched.get("raw_sha256"),
                    body=fetched.get("body") or b"",
                    http_status=fetched.get("http_status"),
                )

    extra: dict[str, Any] = {}
    if _is_labour_table1_workbook_url(url):
        extra["methodology_note"] = AU_LFS_MODERNISATION_METHODOLOGY_NOTE

    return _base_payload(
        ok=True,
        points=[point],
        vintage="latest_available",
        raw_sha256=fetched.get("raw_sha256"),
        http_status=fetched.get("http_status"),
        body=fetched.get("body") or b"",
        **extra,
    )


def _opener_fetch_bytes(
    opener: Callable[..., dict[str, Any]],
    url: str,
    *,
    timeout: float,
) -> bytes:
    resp = opener(url, timeout=timeout)
    if not resp.get("ok"):
        raise RuntimeError(resp.get("error") or "fetch_failed")
    return resp.get("body") or b""


def _fetch_housing_context(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    now: datetime,
    timeout: float,
) -> dict[str, Any]:
    def fetch_bytes(url: str, **_kwargs: Any) -> bytes:
        return _opener_fetch_bytes(opener, url, timeout=timeout)

    block = collect_australia_housing(today=now.date(), fetch_bytes=fetch_bytes)
    status = block.get("status")
    if status == "unavailable":
        return _base_payload(
            ok=False,
            status="source_failed",
            error="housing_feeds_unavailable",
        )
    return _base_payload(
        ok=True,
        vintage="latest_available",
        raw_sha256=_sha256(repr(block).encode()),
        points=[],
    )


def _fetch_pmi(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    now: datetime,
    timeout: float,
    flash_only: bool,
    finals_only: bool,
) -> dict[str, Any]:
    attempted_urls: list[str] = []
    listing = opener(spec.get("endpoint") or PMI_LISTING_URL, timeout=timeout)
    listing_body = listing.get("body") or b""
    if _looks_like_challenge(listing_body, expect="pdf") or listing.get("http_status") in {403, 202}:
        return _base_payload(
            ok=False,
            status="license_gap",
            challenge_page=True,
            http_status=listing.get("http_status"),
            error="pmi_listing_challenge",
            raw_sha256=_sha256(listing_body) if listing_body else None,
            body=listing_body,
        )

    html = listing_body.decode("utf-8", errors="replace") if listing_body else ""
    discovered = discover_australia_listing(html) if html else []
    guids: list[str] = []
    for guid in SEED_GUIDS:
        if guid not in guids:
            guids.append(guid)
    for row in discovered:
        guid = row.get("guid")
        if guid and guid not in guids:
            guids.append(guid)

    parsed_releases: list[dict[str, Any]] = []
    last_body = listing_body
    last_status = listing.get("http_status")

    for guid in guids:
        url = press_release_url(guid)
        attempted_urls.append(url)
        resp = opener(url, timeout=timeout)
        body = resp.get("body") or b""
        last_body = body or last_body
        last_status = resp.get("http_status")
        if _looks_like_challenge(body, expect="pdf"):
            return _base_payload(
                ok=False,
                status="license_gap",
                challenge_page=True,
                http_status=resp.get("http_status"),
                error="pmi_release_challenge",
                raw_sha256=_sha256(body) if body else None,
                body=body,
            )
        if body[:4] != b"%PDF":
            continue
        meta = parse_release_artifact(body)
        flash = bool(meta.get("is_flash"))
        if flash_only and not flash:
            continue
        if finals_only and flash:
            continue
        parsed_releases.append(
            {
                "source_url": url,
                "release_date": meta.get("release_date"),
                "reference_period": meta.get("reference_period"),
                "composite_value": meta.get("composite_value"),
                "is_flash": flash,
                "restated_priors": meta.get("restated_priors"),
                "body": body,
            }
        )

    if not parsed_releases:
        return _base_payload(
            ok=False,
            status="license_gap",
            challenge_page=True,
            http_status=last_status,
            error="pmi_no_primary_pdf",
            raw_sha256=_sha256(last_body) if last_body else None,
            body=last_body,
        )

    retrieved_at = now.astimezone().isoformat()
    chosen = choose_observations(
        [
            {
                "source_url": rel["source_url"],
                "release_date": rel.get("release_date"),
                "reference_period": rel.get("reference_period"),
                "composite_value": rel.get("composite_value"),
                "is_flash": rel.get("is_flash"),
                "restated_priors": rel.get("restated_priors"),
            }
            for rel in parsed_releases
        ],
        retrieved_at=retrieved_at,
    )
    if not chosen:
        return _base_payload(
            ok=False,
            status="source_failed",
            error="pmi_no_observation_in_window",
            body=parsed_releases[-1].get("body") or b"",
            raw_sha256=_sha256(parsed_releases[-1].get("body") or b""),
        )

    latest_period = sorted(chosen.keys())[-1]
    obs = chosen[latest_period]
    revision = obs.get("revision_status") or "final"
    if flash_only:
        revision = "flash"

    point = {
        "period": latest_period,
        "value": float(obs["value"]),
        "transformation": spec.get("transform"),
        "revision_status": revision,
        "source_url": obs.get("source_url"),
    }
    pdf_body = next(
        (rel.get("body") for rel in parsed_releases if rel.get("source_url") == obs.get("source_url")),
        parsed_releases[-1].get("body") or b"",
    )
    return _base_payload(
        ok=True,
        points=[point],
        vintage="latest_available",
        raw_sha256=_sha256(pdf_body),
        http_status=200,
        body=pdf_body,
    )


def _map_leading_point(point: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    transform = str(spec.get("transform") or "")
    return {
        "period": point["period"],
        "value": point["value"],
        "transformation": point.get("transformation") or transform,
        "revision_status": point.get("revision_status") or "latest_available",
        "source_url": point.get("source_url"),
        "release_date": point.get("release_date"),
        "units": point.get("units") or spec.get("units"),
        "raw_sha256": point.get("raw_sha256"),
    }


def _leading_gap_for_spec(gaps: list[dict[str, Any]], spec_id: str) -> dict[str, Any] | None:
    for gap in gaps:
        if str(gap.get("id") or "") == spec_id:
            return gap
    return None


def _fetch_au_leading_series(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    now: datetime,
) -> dict[str, Any] | None:
    method = str(spec.get("retrieval_method") or "")
    spec_id = str(spec.get("id") or "")
    as_of_date = now.date()

    def _wrap_opener(url: str) -> dict[str, Any]:
        return opener(url, timeout=20)

    if method == "jsa_ivi_xlsx":
        from scripts.macro_ingestion.au_leading import jsa_ivi

        collected = jsa_ivi.fetch_ivi(as_of=now, opener=_wrap_opener)
    elif method == "jsa_reos_xlsx":
        from scripts.macro_ingestion.au_leading import jsa_reos

        collected = jsa_reos.fetch_reos(as_of=now, opener=_wrap_opener)
    elif method == "anz_indeed_xlsx":
        from scripts.macro_ingestion.au_leading import anz

        collected = anz.collect_anz(as_of=as_of_date, opener=_wrap_opener)
    elif method == "nab_monthly_business_survey_pdf":
        from scripts.macro_ingestion.au_leading import nab

        collected = nab.collect_nab(as_of=as_of_date, opener=_wrap_opener)
    elif method == "seek_employment_report_html":
        from scripts.macro_ingestion.au_leading import seek

        collected = seek.collect_seek(as_of=as_of_date, opener=_wrap_opener)
    else:
        return None

    raw_points = collected.get("points") or []
    gaps = collected.get("gaps") or []
    points = [
        _map_leading_point(point, spec)
        for point in raw_points
        if str(point.get("id") or spec_id) == spec_id
    ]

    if points:
        digest = next((p.get("raw_sha256") for p in points if p.get("raw_sha256")), None)
        vintage = raw_points[0].get("vintage") if raw_points else None
        return _base_payload(
            ok=True,
            points=points,
            vintage=vintage,
            raw_sha256=digest,
            http_status=200,
        )

    gap = _leading_gap_for_spec(gaps, spec_id)
    if gap is not None:
        status = str(gap.get("status") or "source_failed")
        if status == "not_yet_released":
            return _base_payload(ok=True, points=[], http_status=None)
        if status in {"license_gap", "source_failed"}:
            return _base_payload(
                ok=False,
                status=status,
                error=str(gap.get("reason") or status),
                points=[],
            )

    if method == "jsa_reos_xlsx" and collected.get("status") in {"source_failed", "license_gap"}:
        return _base_payload(
            ok=False,
            status=str(collected.get("status")),
            error=str(
                gap.get("reason") if gap else collected.get("status") or "REOS workbook missing"
            ),
            points=[],
        )

    if collected.get("status") == "source_failed" and not points:
        return _base_payload(
            ok=False,
            status="source_failed",
            error=str(gap.get("reason") if gap else "leading_source_failed"),
            points=[],
        )

    return _base_payload(ok=True, points=[])


def fetch_series(
    spec: dict[str, Any],
    *,
    opener: Callable[..., dict[str, Any]],
    now: datetime,
    timeout: float = 20,
) -> dict[str, Any]:
    role = spec.get("role")
    spec_id = spec.get("id", "")

    if role == "explanatory_alias":
        return _base_payload(
            ok=False,
            status="not_applicable",
            error="explanatory_alias",
        )

    if spec_id == "AU.Consumer.confidence":
        return _base_payload(
            ok=False,
            status="source_failed",
            error="westpac_unverified_no_primary_adapter",
        )

    if spec.get("retrieval_method") == "existing_australia_housing_data":
        return _fetch_housing_context(spec, opener=opener, now=now, timeout=timeout)

    if spec_id == "AU.Activity.business_surveys":
        return _fetch_pmi(spec, opener=opener, now=now, timeout=timeout, flash_only=False, finals_only=True)

    if spec_id == "AU.Activity.flash_composite_pmi":
        return _fetch_pmi(spec, opener=opener, now=now, timeout=timeout, flash_only=True, finals_only=False)

    max_period: str | None = None
    if spec_id in {"AU.Consumer.retail", "AU.Consumer.retail_ceased"}:
        max_period = RETAIL_SERIES_CEASED_AFTER

    if spec.get("retrieval_method", "").startswith("abs_time_series_workbook"):
        return _fetch_abs_workbook_series(
            spec,
            opener=opener,
            timeout=timeout,
            now=now,
            max_period=max_period,
        )

    if spec.get("series_id") is None and spec_id in _CONTEXT_SERIES_ENDPOINTS:
        overlay = dict(spec)
        overlay["series_id"] = _CONTEXT_SERIES_IDS[spec_id]
        overlay["endpoint"] = _CONTEXT_SERIES_ENDPOINTS[spec_id]
        return _fetch_abs_workbook_series(
            overlay,
            opener=opener,
            timeout=timeout,
            now=now,
            max_period=max_period,
        )

    leading_payload = _fetch_au_leading_series(
        spec,
        opener=opener,
        now=now,
    )
    if leading_payload is not None:
        return leading_payload

    return _base_payload(
        ok=False,
        status="source_failed",
        error="unsupported_au_series",
    )
