"""Jobs and Skills Australia Recruitment Experiences and Outlook Survey (REOS)."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from html import unescape
from typing import Any, Callable
from urllib.parse import urljoin

import openpyxl

from scripts.macro_ingestion.au_leading.asof import filter_observations_by_asof

PUBLISHER = "Jobs and Skills Australia"

CATALOG_ID_RECRUITING = "AU.Labor.jsa_reos_recruiting_share"
CATALOG_ID_DIFFICULTY = "AU.Labor.jsa_reos_recruitment_difficulty"
CATALOG_ID_HIRING_INTENT = "AU.Labor.jsa_reos_hiring_intent"

REOS_LANDING_URL = (
    "https://www.jobsandskills.gov.au/data/recruitment-experiences-and-outlook-survey"
)

SERIES_RECRUITING = "Recruitment rate"
SERIES_DIFFICULTY = "Recruitment difficulty"
SERIES_HIRING = "Expect to increase staff"

UNITS_PERCENT = "percent"


@dataclass(frozen=True)
class FetchAttempt:
    url: str
    http_status: int | None
    content_type: str | None
    sha256: str | None
    outcome: str
    parsed: str | None = None


def _sha256(body: bytes) -> str:
    return hashlib.sha256(body).hexdigest()


def _period_from_dt(when: datetime | date) -> str:
    if isinstance(when, datetime):
        when = when.date()
    return f"{when.year:04d}-{when.month:02d}"


def _parse_release_sentence(html: str) -> date | None:
    m = re.search(
        r"([A-Za-z]+)\s+(\d{4})\s+data were released on\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})",
        html,
        flags=re.IGNORECASE,
    )
    if not m:
        return None
    _, _, day_s, rel_month, rel_year = m.groups()
    months = {
        "january": 1,
        "february": 2,
        "march": 3,
        "april": 4,
        "may": 5,
        "june": 6,
        "july": 7,
        "august": 8,
        "september": 9,
        "october": 10,
        "november": 11,
        "december": 12,
    }
    rel_m = months.get(rel_month.lower())
    if rel_m is None:
        return None
    return date(int(rel_year), rel_m, int(day_s))


def _discover_data_xlsx_url(html: str, base_url: str = REOS_LANDING_URL) -> str | None:
    for m in re.finditer(
        r'href="([^"]*recruitment_insights_report_-_[a-z0-9_]+_data_file\.xlsx)"',
        html,
        flags=re.IGNORECASE,
    ):
        href = unescape(m.group(1))
        if href.startswith("http"):
            return href
        if "/web/" in href and "jobsandskills.gov.au/" in href:
            part = href.split("jobsandskills.gov.au/", 1)[-1]
            return urljoin("https://www.jobsandskills.gov.au/", part.lstrip("/"))
        return urljoin(base_url, href)
    return None


def _read_reos_metadata(wb: openpyxl.Workbook) -> dict[str, dict[str, str]]:
    if "Information" not in wb.sheetnames:
        raise ValueError("missing Information sheet")
    ws = wb["Information"]
    rows = list(ws.iter_rows(min_row=15, max_row=27, min_col=2, max_col=5, values_only=True))
    meta: dict[str, dict[str, str]] = {}
    for row in rows:
        if not row:
            continue
        table_raw, series_raw, freq_raw, _extra = (row + (None, None, None, None))[:4]
        if table_raw is None or series_raw is None:
            continue
        table_no = str(table_raw).strip()
        series = str(series_raw).strip()
        frequency = str(freq_raw or "").strip()
        if not table_no or not series:
            continue
        meta[table_no] = {"series": series, "frequency": frequency}
    return meta


def _sheet_for_table(wb: openpyxl.Workbook, table_no: str) -> str | None:
    table_no = str(table_no).strip()
    major = table_no.split(".", 1)[0]
    for name in wb.sheetnames:
        if name.startswith(f"{table_no} ") or name.startswith(f"{table_no}\t"):
            return name
        if name.startswith(f"{major}."):
            return name
    return None


def _national_monthly_latest(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> tuple[str, float]:
    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = None
    for i, row in enumerate(rows):
        if row and str(row[0]).strip().lower() in {"month", "quarter"}:
            header_row_idx = i
            break
    if header_row_idx is None:
        raise ValueError("month/quarter header row not found")
    header = rows[header_row_idx]
    col_idx = None
    for j, name in enumerate(header[1:], start=1):
        if name and str(name).strip().lower() in {"australia", "national", "aust"}:
            col_idx = j
            break
    if col_idx is None:
        for j, name in enumerate(header[1:], start=1):
            if name and "australia" in str(name).lower():
                col_idx = j
                break
    if col_idx is None:
        raise ValueError("national disaggregation column not found")

    last_period: str | None = None
    last_val: float | None = None
    for row in rows[header_row_idx + 1 :]:
        when = row[0]
        if when is None:
            continue
        if isinstance(when, datetime):
            period = _period_from_dt(when)
        elif isinstance(when, (int, float)):
            # Excel serial date
            from datetime import timedelta

            epoch = datetime(1899, 12, 30)
            when_dt = epoch + timedelta(days=float(when))
            period = _period_from_dt(when_dt)
        else:
            continue
        val = row[col_idx]
        if val is None or not isinstance(val, (int, float)):
            continue
        if isinstance(val, float) and val != val:
            continue
        last_period, last_val = period, float(val)
    if last_period is None or last_val is None:
        raise ValueError("no national monthly values")
    return last_period, last_val


def parse_reos_workbook(
    body: bytes,
    *,
    source_url: str,
    release_date: date,
    raw_sha256: str,
) -> tuple[list[dict[str, Any]], list[str]]:
    """Return (points, gaps) where gaps lists catalog ids that could not be bound."""
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    meta = _read_reos_metadata(wb)
    wanted = {
        SERIES_RECRUITING: CATALOG_ID_RECRUITING,
        SERIES_DIFFICULTY: CATALOG_ID_DIFFICULTY,
        SERIES_HIRING: CATALOG_ID_HIRING,
    }
    points: list[dict[str, Any]] = []
    bound: set[str] = set()

    for table_no, info in meta.items():
        series_name = info.get("series") or ""
        if info.get("frequency", "").lower() != "monthly":
            continue
        catalog_id = wanted.get(series_name)
        if not catalog_id:
            continue
        sheet = _sheet_for_table(wb, table_no)
        if sheet is None:
            continue
        try:
            period, value = _national_monthly_latest(wb[sheet])
        except ValueError:
            continue
        points.append(
            {
                "id": catalog_id,
                "series_id": series_name,
                "publisher": PUBLISHER,
                "weight": 0.0,
                "period": period,
                "value": value,
                "units": UNITS_PERCENT,
                "transformation": "level",
                "series_label": "national",
                "coverage": "Australia",
                "frequency": info.get("frequency"),
                "release_date": release_date.isoformat(),
                "raw_sha256": raw_sha256,
                "source_url": source_url,
                "vintage": f"release:{release_date.isoformat()}",
            }
        )
        bound.add(catalog_id)

    gaps = [cid for cid in wanted.values() if cid not in bound]
    return points, gaps


def fetch_reos(
    *,
    as_of: datetime,
    opener: Callable[[str], dict[str, Any]],
) -> dict[str, Any]:
    attempts: list[FetchAttempt] = []
    gap_rows: list[dict[str, str]] = []
    points: list[dict[str, Any]] = []

    landing = opener(REOS_LANDING_URL)
    body = landing.get("body") or b""
    landing_sha = _sha256(body) if body else None
    attempts.append(
        FetchAttempt(
            url=REOS_LANDING_URL,
            http_status=landing.get("http_status"),
            content_type=landing.get("content_type"),
            sha256=landing_sha,
            outcome="parsed" if landing.get("ok") else "not_retrieved",
        )
    )
    if not landing.get("ok") or not body:
        for cid in (
            CATALOG_ID_RECRUITING,
            CATALOG_ID_DIFFICULTY,
            CATALOG_ID_HIRING_INTENT,
        ):
            gap_rows.append(
                {
                    "id": cid,
                    "publisher": "jsa_reos",
                    "status": "source_failed",
                    "reason": str(landing.get("error") or "REOS landing page not retrieved"),
                }
            )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gap_rows,
            "attempts": attempts,
        }

    html = body.decode("utf-8", errors="replace")
    release_date = _parse_release_sentence(html)
    xlsx_url = _discover_data_xlsx_url(html)
    if not xlsx_url:
        for cid in (
            CATALOG_ID_RECRUITING,
            CATALOG_ID_DIFFICULTY,
            CATALOG_ID_HIRING_INTENT,
        ):
            gap_rows.append(
                {
                    "id": cid,
                    "publisher": "jsa_reos",
                    "status": "source_failed",
                    "reason": "REOS data XLSX link not found on landing page",
                }
            )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gap_rows,
            "attempts": attempts,
        }

    xlsx = opener(xlsx_url)
    xlsx_body = xlsx.get("body") or b""
    xlsx_sha = _sha256(xlsx_body) if xlsx_body else None
    attempts.append(
        FetchAttempt(
            url=xlsx_url,
            http_status=xlsx.get("http_status"),
            content_type=xlsx.get("content_type"),
            sha256=xlsx_sha,
            outcome="parsed" if xlsx.get("ok") else "not_retrieved",
        )
    )
    if not xlsx.get("ok") or not xlsx_body:
        for cid in (
            CATALOG_ID_RECRUITING,
            CATALOG_ID_DIFFICULTY,
            CATALOG_ID_HIRING_INTENT,
        ):
            gap_rows.append(
                {
                    "id": cid,
                    "publisher": "jsa_reos",
                    "status": "source_failed",
                    "reason": str(xlsx.get("error") or "REOS workbook download failed"),
                }
            )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gap_rows,
            "attempts": attempts,
        }

    if not xlsx_body.startswith(b"PK"):
        for cid in (
            CATALOG_ID_RECRUITING,
            CATALOG_ID_DIFFICULTY,
            CATALOG_ID_HIRING_INTENT,
        ):
            gap_rows.append(
                {
                    "id": cid,
                    "publisher": "jsa_reos",
                    "status": "source_failed",
                    "reason": "REOS workbook response is not a valid XLSX archive",
                }
            )
        return {
            "status": "source_failed",
            "points": [],
            "gaps": gap_rows,
            "attempts": attempts,
        }

    if release_date is None:
        for cid in (
            CATALOG_ID_RECRUITING,
            CATALOG_ID_DIFFICULTY,
            CATALOG_ID_HIRING_INTENT,
        ):
            gap_rows.append(
                {
                    "id": cid,
                    "publisher": "jsa_reos",
                    "status": "source_failed",
                    "reason": "release provenance missing; observation quarantined",
                }
            )
        return {
            "status": "source_failed",
            "points": [],
            "gaps": gap_rows,
            "attempts": attempts,
        }

    try:
        points, missing_ids = parse_reos_workbook(
            xlsx_body,
            source_url=xlsx_url,
            release_date=release_date,
            raw_sha256=xlsx_sha or "",
        )
    except ValueError as exc:
        for cid in (
            CATALOG_ID_RECRUITING,
            CATALOG_ID_DIFFICULTY,
            CATALOG_ID_HIRING_INTENT,
        ):
            gap_rows.append(
                {
                    "id": cid,
                    "publisher": "jsa_reos",
                    "status": "source_failed",
                    "reason": str(exc),
                }
            )
        return {
            "status": "source_failed",
            "points": [],
            "gaps": gap_rows,
            "attempts": attempts,
        }

    for cid in missing_ids:
        gap_rows.append(
            {
                "id": cid,
                "publisher": "jsa_reos",
                "status": "license_gap",
                "reason": "series not present in downloaded REOS workbook",
            }
        )

    points = filter_observations_by_asof(points, as_of)
    status = "ok" if points else "source_failed"
    if not points and not gap_rows:
        gap_rows.append(
            {
                "id": CATALOG_ID_RECRUITING,
                "publisher": "jsa_reos",
                "status": "source_failed",
                "reason": "no REOS observations eligible at as_of",
            }
        )
    return {
        "status": status,
        "points": points,
        "gaps": gap_rows,
        "attempts": attempts,
    }


def attempts_to_json(attempts: list[FetchAttempt]) -> list[dict[str, Any]]:
    return [
        {
            "url": a.url,
            "http_status": a.http_status,
            "content_type": a.content_type,
            "sha256": a.sha256,
            "outcome": a.outcome,
            "parsed": a.parsed,
        }
        for a in attempts
    ]
