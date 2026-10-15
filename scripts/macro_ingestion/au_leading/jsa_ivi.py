"""Jobs and Skills Australia Internet Vacancy Index (IVI)."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from html import unescape
from typing import Any, Callable
from urllib.parse import urljoin

import openpyxl

from scripts.macro_ingestion.au_leading.asof import filter_observations_by_asof

PUBLISHER = "Jobs and Skills Australia"
CATALOG_ID_SA_LEVEL = "AU.Labor.jsa_ivi_national_sa"
SERIES_ID_NATIONAL_SA = "AUSTRALIAN_TOTAL"

IVI_LANDING_URL = "https://www.jobsandskills.gov.au/data/internet-vacancy-index"
UNITS_LEVEL = "vacancies"


@dataclass(frozen=True)
class FetchAttempt:
    url: str
    http_status: int | None
    content_type: str | None
    sha256: str | None
    outcome: str
    parsed: str | None = None


def _sha256(body: bytes) -> str:
    return hashlib.sha256(body).hexdigest()


def _period_from_dt(when: datetime | date) -> str:
    if isinstance(when, datetime):
        when = when.date()
    return f"{when.year:04d}-{when.month:02d}"


def _parse_release_sentence(html: str) -> tuple[str | None, date | None]:
    """Return (reference month label, release_date) from landing copy."""
    m = re.search(
        r"([A-Za-z]+)\s+(\d{4})\s+data were released on\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})",
        html,
        flags=re.IGNORECASE,
    )
    if not m:
        return None, None
    month_name, year_s, day_s, rel_month, rel_year = m.groups()
    months = {
        "january": 1,
        "february": 2,
        "march": 3,
        "april": 4,
        "may": 5,
        "june": 6,
        "july": 7,
        "august": 8,
        "september": 9,
        "october": 10,
        "november": 11,
        "december": 12,
    }
    ref_month = months.get(month_name.lower())
    rel_m = months.get(rel_month.lower())
    if ref_month is None or rel_m is None:
        return None, None
    ref_label = f"{month_name} {year_s}"
    release_date = date(int(rel_year), rel_m, int(day_s))
    return ref_label, release_date


def _discover_skill_states_xlsx_url(html: str, base_url: str = IVI_LANDING_URL) -> str | None:
    for m in re.finditer(
        r'href="([^"]*internet_vacancies_anzsco_skill_level_states_and_territories[^"]+\.xlsx)"',
        html,
        flags=re.IGNORECASE,
    ):
        href = unescape(m.group(1))
        if href.startswith("http"):
            return href
        if "/web/" in href and "jobsandskills.gov.au/" in href:
            part = href.split("jobsandskills.gov.au/", 1)[-1]
            return urljoin("https://www.jobsandskills.gov.au/", part.lstrip("/"))
        return urljoin(base_url, href)
    return None


def parse_ivi_skill_states_workbook(
    body: bytes,
    *,
    source_url: str,
    release_date: date,
    raw_sha256: str,
    target_period: str | None = None,
) -> list[dict[str, Any]]:
    """Parse national seasonally adjusted vacancy level (and m/m change) from IVI workbook."""
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if "Seasonally Adjusted" not in wb.sheetnames:
        raise ValueError("missing Seasonally Adjusted sheet")
    ws = wb["Seasonally Adjusted"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    date_cols: list[tuple[int, datetime]] = []
    for idx, cell in enumerate(header):
        if isinstance(cell, datetime):
            date_cols.append((idx, cell))
    if not date_cols:
        raise ValueError("no monthly date columns in Seasonally Adjusted sheet")

    national_row: tuple[Any, ...] | None = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] == "AUSTRALIAN TOTAL" and row[2] == "AUST" and row[3] == 0:
            national_row = row
            break
    if national_row is None:
        raise ValueError("AUSTRALIAN TOTAL row not found")

    series: list[tuple[str, float]] = []
    for idx, when in date_cols:
        val = national_row[idx]
        if val is None or not isinstance(val, (int, float)):
            continue
        if isinstance(val, float) and val != val:
            continue
        series.append((_period_from_dt(when), float(val)))
    if not series:
        raise ValueError("no numeric SA national levels")

    series.sort(key=lambda x: x[0])
    if target_period:
        filtered = [p for p in series if p[0] == target_period]
        if not filtered:
            raise ValueError(f"period {target_period} not in workbook")
        latest_period, latest_val = filtered[-1]
        prior = [p for p in series if p[0] < target_period]
    else:
        latest_period, latest_val = series[-1]
        prior = series[:-1]

    points: list[dict[str, Any]] = [
        {
            "id": CATALOG_ID_SA_LEVEL,
            "series_id": SERIES_ID_NATIONAL_SA,
            "publisher": PUBLISHER,
            "weight": 0.0,
            "period": latest_period,
            "value": latest_val,
            "units": UNITS_LEVEL,
            "transformation": "sa_level",
            "series_label": "seasonally_adjusted",
            "release_date": release_date.isoformat(),
            "raw_sha256": raw_sha256,
            "source_url": source_url,
            "vintage": f"release:{release_date.isoformat()}",
        }
    ]

    if prior:
        prev_period, prev_val = prior[-1]
        if prev_period < latest_period:
            mom_pct = (latest_val - prev_val) / prev_val * 100.0
            points.append(
                {
                    "id": CATALOG_ID_SA_LEVEL,
                    "series_id": SERIES_ID_NATIONAL_SA,
                    "publisher": PUBLISHER,
                    "weight": 0.0,
                    "period": latest_period,
                    "value": mom_pct,
                    "units": "percent",
                    "transformation": "mom_change_from_published_levels",
                    "series_label": "seasonally_adjusted",
                    "release_date": release_date.isoformat(),
                    "raw_sha256": raw_sha256,
                    "source_url": source_url,
                    "vintage": f"release:{release_date.isoformat()}",
                    "derivation": {
                        "from_periods": [prev_period, latest_period],
                        "from_values": [prev_val, latest_val],
                    },
                }
            )
    return points


def fetch_ivi(
    *,
    as_of: datetime,
    opener: Callable[[str], dict[str, Any]],
) -> dict[str, Any]:
    attempts: list[FetchAttempt] = []
    gaps: list[dict[str, str]] = []
    points: list[dict[str, Any]] = []

    landing = opener(IVI_LANDING_URL)
    body = landing.get("body") or b""
    landing_sha = _sha256(body) if body else None
    attempts.append(
        FetchAttempt(
            url=IVI_LANDING_URL,
            http_status=landing.get("http_status"),
            content_type=landing.get("content_type"),
            sha256=landing_sha,
            outcome="parsed" if landing.get("ok") else "not_retrieved",
            parsed=None if not landing.get("ok") else "landing_html",
        )
    )
    if not landing.get("ok") or not body:
        gaps.append(
            {
                "id": CATALOG_ID_SA_LEVEL,
                "publisher": "jsa_ivi",
                "status": "source_failed",
                "reason": str(landing.get("error") or "IVI landing page not retrieved"),
            }
        )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gaps,
            "attempts": attempts,
        }

    html = body.decode("utf-8", errors="replace")
    _, release_date = _parse_release_sentence(html)
    if release_date is None:
        gaps.append(
            {
                "id": CATALOG_ID_SA_LEVEL,
                "publisher": "jsa_ivi",
                "status": "source_failed",
                "reason": "could not parse IVI release sentence from landing page",
            }
        )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gaps,
            "attempts": attempts,
        }

    xlsx_url = _discover_skill_states_xlsx_url(html)
    if not xlsx_url:
        gaps.append(
            {
                "id": CATALOG_ID_SA_LEVEL,
                "publisher": "jsa_ivi",
                "status": "source_failed",
                "reason": "skill-level states IVI XLSX link not found on landing page",
            }
        )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gaps,
            "attempts": attempts,
        }

    xlsx = opener(xlsx_url)
    xlsx_body = xlsx.get("body") or b""
    xlsx_sha = _sha256(xlsx_body) if xlsx_body else None
    attempts.append(
        FetchAttempt(
            url=xlsx_url,
            http_status=xlsx.get("http_status"),
            content_type=xlsx.get("content_type"),
            sha256=xlsx_sha,
            outcome="parsed" if xlsx.get("ok") else "not_retrieved",
            parsed=None,
        )
    )
    if not xlsx.get("ok") or not xlsx_body:
        gaps.append(
            {
                "id": CATALOG_ID_SA_LEVEL,
                "publisher": "jsa_ivi",
                "status": "source_failed",
                "reason": str(xlsx.get("error") or "IVI workbook download failed"),
            }
        )
        return {
            "status": "source_failed",
            "points": points,
            "gaps": gaps,
            "attempts": attempts,
        }

    try:
        points = parse_ivi_skill_states_workbook(
            xlsx_body,
            source_url=xlsx_url,
            release_date=release_date,
            raw_sha256=xlsx_sha or "",
        )
    except ValueError as exc:
        gaps.append(
            {
                "id": CATALOG_ID_SA_LEVEL,
                "publisher": "jsa_ivi",
                "status": "source_failed",
                "reason": str(exc),
            }
        )
        return {
            "status": "source_failed",
            "points": [],
            "gaps": gaps,
            "attempts": attempts,
        }

    points = filter_observations_by_asof(points, as_of)

    status = "ok" if points else "source_failed"
    if not points and not gaps:
        gaps.append(
            {
                "id": CATALOG_ID_SA_LEVEL,
                "publisher": "jsa_ivi",
                "status": "source_failed",
                "reason": "no IVI observations eligible at as_of",
            }
        )
    return {
        "status": status,
        "points": points,
        "gaps": gaps,
        "attempts": attempts,
    }


def attempts_to_json(attempts: list[FetchAttempt]) -> list[dict[str, Any]]:
    return [
        {
            "url": a.url,
            "http_status": a.http_status,
            "content_type": a.content_type,
            "sha256": a.sha256,
            "outcome": a.outcome,
            "parsed": a.parsed,
        }
        for a in attempts
    ]
