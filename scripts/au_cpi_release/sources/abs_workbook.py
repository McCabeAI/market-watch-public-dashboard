"""Archived ABS CPI workbooks.

In-repo July 2026 workbooks under data/temperature_history/raw/au/ are genuine
ABS bytes already in the repo. They are hashed in place and not rewritten.
Table 6 (640106) is parsed with scripts.macro_ingestion.adapters.au for the
headline, trimmed-mean and weighted-median series. Optional Wayback captures
are labeled mirror:web.archive.org and may fail without failing the snapshot.
"""

from __future__ import annotations

import gzip
import hashlib
import io
import json
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from scripts.au_cpi_release.calendar import MIRROR, entry_for
from scripts.au_cpi_release.contract import REGIME_COMPLETE_MONTHLY, Observation
from scripts.au_cpi_release.sources.http import USER_AGENT, SourceError, looks_like_html_or_challenge
from scripts.macro_ingestion.adapters.au import (
    parse_abs_time_series_workbook,
    parse_abs_workbook_columns,
)

REPO_ROOT = Path(__file__).resolve().parents[3]
WORKBOOK_DIR = REPO_ROOT / "data" / "temperature_history" / "raw" / "au"
VINTAGE_JULY = "abs_release_2026-07"
JULY_REFERENCE = "2026-07"

# Series IDs were read from the July 2026 640106 header row (description + Series ID).
TABLE6_SERIES: dict[str, str] = {
    "A130393720C": "headline_nsa_index",
    "A130393722J": "headline_nsa_mom",
    "A130393721F": "headline_nsa_yoy",
    "A130607789R": "headline_sa_index",
    "A130607784C": "headline_sa_mom",
    "A130607783A": "headline_sa_yoy",
    "A130400381L": "trimmed_mean_monthly_sa_index",
    "A130400383T": "trimmed_mean_monthly_sa_mom",
    "A130400382R": "trimmed_mean_monthly_sa_yoy",
    "A130400612K": "weighted_median_monthly_sa_index",
    "A130400614R": "weighted_median_monthly_sa_mom",
    "A130400613L": "weighted_median_monthly_sa_yoy",
}

HEADLINE_INDEX_ID = "A130393720C"

# Original ABS URLs. Wayback is a mirror, not the publisher.
_WAYBACK_FILES: tuple[tuple[str, str], ...] = (
    ("aug-2026", "640101.xlsx"),
    ("aug-2026", "640106.xlsx"),
    ("aug-2026", "640107.xlsx"),
    ("aug-2026", "640108.xlsx"),
    ("aug-2026", "64010Appendix1a.xlsx"),
    ("aug-2026", "64010Appendix2.xlsx"),
    ("jun-2026", "64010Appendix1a.xlsx"),
)


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _sha256(body: bytes) -> str:
    return hashlib.sha256(body).hexdigest()


def _original_workbook_url(slug: str, filename: str) -> str:
    return (
        "https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/"
        f"consumer-price-index-australia/{slug}/{filename}"
    )


def july_workbook_paths() -> list[Path]:
    paths = sorted(WORKBOOK_DIR.glob("cpi_64010*_jul2026.xlsx"))
    return [path for path in paths if path.is_file()]


def _workbook_record(path: Path, body: bytes, *, retrieved_at: str) -> dict[str, object]:
    stem = path.name.replace(".xlsx", "")
    return {
        "accept": None,
        "bytes": len(body),
        "dataflow": f"ABS_WORKBOOK,{stem}",
        "http_status": None,
        "key": None,
        "path": None,
        "repo_path": str(path.relative_to(REPO_ROOT)),
        "retrieved_at": retrieved_at,
        "sha256": _sha256(body),
        "source_id": f"abs_workbook_{stem}",
        "url": f"repo:{path.relative_to(REPO_ROOT).as_posix()}",
        "vintage": VINTAGE_JULY,
    }


def _observations_from_columns(
    columns: dict[str, list[tuple[str, float]]],
    mapping: dict[str, str],
    *,
    source_id: str,
    vintage: str,
    published_at: str,
    regime: str,
) -> list[Observation]:
    observations: list[Observation] = []
    for series_id, canonical_id in mapping.items():
        for period, value in columns.get(series_id) or []:
            observations.append(
                Observation(
                    canonical_id=canonical_id,
                    period=period,
                    value=float(value),
                    published_at=published_at,
                    source_id=source_id,
                    regime=regime,
                    vintage=vintage,
                )
            )
    return observations


def load_july_workbooks(
    *,
    retrieved_at: str | None = None,
) -> tuple[list[Observation], list[dict[str, object]], list[str]]:
    """Hash every in-repo July 2026 CPI workbook and parse table 6."""
    when = retrieved_at or _utc_now()
    paths = july_workbook_paths()
    if not paths:
        raise SourceError("workbook_missing", str(WORKBOOK_DIR))
    notes = [
        "In-repo cpi_64010*_jul2026.xlsx files are the July 2026 ABS vintage. "
        "They are hashed at their repo path and are not copied. "
        f"Observations parsed from 640106 use published_at of the {JULY_REFERENCE} "
        "reference release for every period, because that is when this vintage's bytes "
        "were published (later revisions are inside the file)."
    ]
    records: list[dict[str, object]] = []
    bodies: dict[str, bytes] = {}
    for path in paths:
        body = path.read_bytes()
        if body[:2] != b"PK":
            raise SourceError("workbook_not_xlsx", path.name)
        records.append(_workbook_record(path, body, retrieved_at=when))
        bodies[path.name] = body

    table6_name = "cpi_640106_jul2026.xlsx"
    if table6_name not in bodies:
        raise SourceError("workbook_missing", table6_name)
    parsed = parse_abs_workbook_columns(bodies[table6_name], set(TABLE6_SERIES), cadence="monthly")
    missing = sorted(set(TABLE6_SERIES) - set(parsed))
    if missing:
        notes.append("640106 series ids not parsed: " + ",".join(missing))
    if HEADLINE_INDEX_ID not in parsed:
        raise SourceError("workbook_parse", "640106 missing all-groups index")
    published_at = entry_for(JULY_REFERENCE).release_at
    source_id = "abs_workbook_cpi_640106_jul2026"
    observations = _observations_from_columns(
        parsed,
        TABLE6_SERIES,
        source_id=source_id,
        vintage=VINTAGE_JULY,
        published_at=published_at,
        regime=REGIME_COMPLETE_MONTHLY,
    )

    table1_name = "cpi_640101_jul2026.xlsx"
    if table1_name in bodies:
        table1 = parse_abs_workbook_columns(bodies[table1_name], {HEADLINE_INDEX_ID}, cadence="monthly")
        left = {period: value for period, value in parsed.get(HEADLINE_INDEX_ID, [])}
        right = {period: value for period, value in table1.get(HEADLINE_INDEX_ID, [])}
        overlap = sorted(set(left) & set(right))
        mismatches = [period for period in overlap if left[period] != right[period]]
        if not overlap:
            notes.append("640101 did not yield the all-groups index for a cross-check")
        elif mismatches:
            notes.append(
                "640101 and 640106 all-groups index differ at " + ",".join(mismatches[:8])
            )
        else:
            notes.append(
                f"640101 and 640106 all-groups index agree on {len(overlap)} overlapping months"
            )
    return observations, records, notes


def _maybe_gunzip(body: bytes) -> bytes:
    if body[:2] == b"\x1f\x8b":
        return gzip.decompress(body)
    return body


def _cdx_timestamps(original_url: str, *, timeout: float) -> list[str]:
    from urllib.parse import quote

    bare = original_url.replace("https://", "").replace("http://", "")
    url = (
        "https://web.archive.org/cdx/search/cdx?url="
        + quote(bare, safe="/")
        + "&output=json&fl=timestamp,statuscode,length&filter=statuscode:200&limit=8"
    )
    request = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(request, timeout=timeout) as response:
        payload = json.loads(response.read().decode("utf-8"))
    stamps: list[str] = []
    for row in payload[1:]:
        if row and row[0]:
            stamps.append(str(row[0]))
    return stamps


def _download_capture(original_url: str, timestamp: str, *, timeout: float) -> tuple[bytes, int, str, str | None]:
    capture_url = f"https://web.archive.org/web/{timestamp}id_/{original_url}"
    request = Request(capture_url, headers={"User-Agent": USER_AGENT, "Accept": "*/*"})
    with urlopen(request, timeout=timeout) as response:
        body = _maybe_gunzip(response.read())
        status = int(getattr(response, "status", None) or response.getcode() or 0)
        content_type = response.headers.get("Content-Type")
    return body, status, capture_url, content_type


def _slug_reference(slug: str) -> str:
    months = {
        "jan": "01",
        "feb": "02",
        "mar": "03",
        "apr": "04",
        "may": "05",
        "jun": "06",
        "jul": "07",
        "aug": "08",
        "sep": "09",
        "oct": "10",
        "nov": "11",
        "dec": "12",
    }
    month_name, year = slug.split("-")
    return f"{year}-{months[month_name[:3]]}"


def _appendix_series_map(body: bytes) -> dict[str, tuple[str, str]]:
    """Map series id -> (canonical_id, cadence) for a trimmed-mean appendix, if the layout matches."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise SourceError("openpyxl_missing", "openpyxl required") from exc
    if body[:2] != b"PK":
        return {}
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    found: dict[str, tuple[str, str]] = {}
    try:
        for sheet_name in workbook.sheetnames:
            if not sheet_name.startswith("Data"):
                continue
            rows = []
            for index, row in enumerate(workbook[sheet_name].iter_rows(max_row=10, values_only=True)):
                rows.append(row)
                if index >= 9:
                    break
            if len(rows) < 10:
                continue
            for col, series_id in enumerate(rows[9]):
                if not series_id or series_id == "Series ID":
                    continue
                title = str(rows[0][col] or "")
                lowered = title.lower()
                if "trimmed mean" not in lowered or "australia" not in lowered:
                    continue
                if "sydney" in lowered:
                    continue
                freq = str(rows[4][col] or "").lower()
                cadence = "quarterly" if "quarter" in freq else "monthly"
                if "index number" in lowered:
                    canonical = "appendix_1a_trimmed_mean_index"
                elif "previous year" in lowered or "corresponding" in lowered:
                    canonical = "appendix_1a_trimmed_mean_yoy"
                elif "previous period" in lowered or "quarterly" in lowered:
                    canonical = "appendix_1a_trimmed_mean_qoq"
                else:
                    continue
                found[str(series_id)] = (canonical, cadence)
    finally:
        workbook.close()
    return found


def _parse_wayback_workbook(
    body: bytes,
    filename: str,
    *,
    source_id: str,
    vintage: str,
    published_at: str,
) -> tuple[list[Observation], str]:
    if "Appendix1a" in filename:
        mapping = _appendix_series_map(body)
        if not mapping:
            return [], "appendix_1a_layout_not_recognised"
        observations: list[Observation] = []
        for series_id, (canonical_id, cadence) in mapping.items():
            points = parse_abs_time_series_workbook(body, series_id, cadence=cadence)
            for period, value in points:
                observations.append(
                    Observation(
                        canonical_id=canonical_id,
                        period=period,
                        value=float(value),
                        published_at=published_at,
                        source_id=source_id,
                        regime="quarterly_legacy_tm",
                        vintage=vintage,
                    )
                )
        return observations, f"appendix_1a_parsed_{len(mapping)}_series"
    parsed = parse_abs_workbook_columns(body, set(TABLE6_SERIES), cadence="monthly")
    if not parsed:
        return [], "no_table6_series_ids"
    observations = _observations_from_columns(
        parsed,
        {series_id: canonical for series_id, canonical in TABLE6_SERIES.items() if series_id in parsed},
        source_id=source_id,
        vintage=vintage,
        published_at=published_at,
        regime=REGIME_COMPLETE_MONTHLY,
    )
    return observations, f"parsed_{len(parsed)}_table6_series"


def attempt_wayback_workbooks(
    *,
    budget_s: float = 180.0,
    timeout: float = 25.0,
) -> tuple[list[Observation], list[dict[str, object]], dict[str, bytes], list[str]]:
    """Time-boxed Wayback fetches. Failures are notes, not exceptions."""
    deadline = time.monotonic() + budget_s
    notes: list[str] = []
    observations: list[Observation] = []
    records: list[dict[str, object]] = []
    raw: dict[str, bytes] = {}
    for slug, filename in _WAYBACK_FILES:
        original = _original_workbook_url(slug, filename)
        label = f"{slug}/{filename}"
        if time.monotonic() > deadline:
            notes.append(f"wayback skipped {label}: time box exhausted")
            continue
        try:
            stamps = _cdx_timestamps(original, timeout=min(timeout, 20.0))
        except (HTTPError, URLError, TimeoutError, OSError, json.JSONDecodeError, ValueError) as exc:
            notes.append(f"wayback cdx failed {label}: {type(exc).__name__}: {exc}")
            continue
        if not stamps:
            notes.append(f"wayback cdx empty {label}")
            continue
        body = b""
        status = 0
        capture_url = ""
        used_stamp = ""
        downloaded = False
        for stamp in reversed(stamps):
            if time.monotonic() > deadline:
                notes.append(f"wayback skipped download {label}: time box exhausted")
                break
            try:
                body, status, capture_url, content_type = _download_capture(original, stamp, timeout=timeout)
            except (HTTPError, URLError, TimeoutError, OSError) as exc:
                notes.append(f"wayback download failed {label} {stamp}: {type(exc).__name__}: {exc}")
                continue
            if looks_like_html_or_challenge(body, content_type) or body[:2] != b"PK":
                notes.append(f"wayback {label} {stamp} was not an xlsx ({len(body)} bytes)")
                continue
            used_stamp = stamp
            downloaded = True
            break
        if not downloaded:
            notes.append(f"wayback not obtained {label}")
            continue
        reference = _slug_reference(slug)
        try:
            published_at = entry_for(reference).release_at
        except Exception as exc:  # noqa: BLE001 - record and keep the bytes
            published_at = ""
            notes.append(f"wayback {label} has no calendar entry: {exc}")
        source_id = f"wayback_{slug}_{filename.replace('.xlsx', '')}"
        vintage = f"abs_release_{reference}"
        parsed_obs: list[Observation] = []
        parse_note = "stored_unparsed"
        if published_at:
            try:
                parsed_obs, parse_note = _parse_wayback_workbook(
                    body,
                    filename,
                    source_id=source_id,
                    vintage=vintage,
                    published_at=published_at,
                )
            except (ValueError, SourceError, OSError) as exc:
                parse_note = f"parse_failed:{type(exc).__name__}:{exc}"
        notes.append(
            f"wayback obtained {label} capture {used_stamp} bytes {len(body)} "
            f"source {MIRROR} original {original}; {parse_note}"
        )
        safe_name = f"wayback_{slug}_{filename}"
        records.append(
            {
                "accept": "*/*",
                "bytes": len(body),
                "capture_timestamp": used_stamp,
                "dataflow": f"ABS_WORKBOOK,{filename}",
                "http_status": status,
                "key": None,
                "mirror": MIRROR,
                "original_url": original,
                "path": f"raw/{safe_name}",
                "retrieved_at": _utc_now(),
                "sha256": _sha256(body),
                "source_id": source_id,
                "url": capture_url,
            }
        )
        raw[f"raw/{safe_name}"] = body
        observations.extend(parsed_obs)
    if not records:
        notes.append("wayback: no ABS workbook captures were stored")
    return observations, records, raw, notes
