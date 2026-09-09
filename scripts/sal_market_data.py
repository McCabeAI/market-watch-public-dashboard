#!/usr/bin/env python3
"""Deterministic Market Watch market-state acquisition and analytics (Sal V1).

No model calls. Live authorities are the source of truth. The script fetches
five years of sovereign yield history and G10 FX reference history, computes a
compact derived snapshot, and optionally renders a dashboard fragment.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import statistics
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Mapping, Sequence

USER_AGENT = "MarketWatch-Sal/1.0 (+https://github.com/kevinnmccabe-PC/market-watch-public-dashboard)"
DEFAULT_TIMEOUT = 30
RETRIES = 3

G10 = ("EUR", "GBP", "AUD", "NZD", "USD", "CAD", "CHF", "NOK", "SEK", "JPY")
FX_ORDER = G10
RATE_TENORS = ("2Y", "5Y", "10Y")
RATE_COUNTRIES = ("US", "CA", "AU", "NZ")
RV_PAIRS = (("CA", "US"), ("AU", "US"), ("NZ", "US"), ("AU", "NZ"), ("CA", "AU"))

TREASURY_URL = (
    "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/"
    "daily-treasury-rates.csv/{year}/all?_format=csv&field_tdr_date_value={year}"
    "&page=&type=daily_treasury_yield_curve"
)
BOC_URL = "https://www.bankofcanada.ca/valet/observations/V39051,V39053,V39055,V39056/json"
RBA_URL = "https://www.rba.gov.au/statistics/tables/csv/f2-data.csv"
RBNZ_URL = "https://rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx"
ECB_FX_URL = "https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.csv"


class SalError(RuntimeError):
    pass


def fetch_bytes(url: str, *, timeout: int = DEFAULT_TIMEOUT, retries: int = RETRIES) -> bytes:
    last: Exception | None = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "*/*"})
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                status = getattr(resp, "status", 200)
                if status != 200:
                    raise SalError(f"HTTP {status} from {url}")
                data = resp.read()
                if not data:
                    raise SalError(f"empty response from {url}")
                return data
        except (urllib.error.URLError, TimeoutError, SalError) as exc:
            last = exc
            if attempt + 1 < retries:
                time.sleep(1.5 * (attempt + 1))
    raise SalError(f"failed to fetch {url}: {last}")


def parse_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y", "%d-%b-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    s = str(value).strip().replace("%", "").replace(",", "")
    if not s or s in {"..", "-", "—", "n/a", "NA", "na"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def filter_since(series: Mapping[date, float], start: date) -> dict[date, float]:
    return {d: v for d, v in series.items() if d >= start}


def percentile_rank(values: Sequence[float], x: float) -> float | None:
    if not values:
        return None
    lt = sum(v < x for v in values)
    eq = sum(v == x for v in values)
    return 100.0 * (lt + 0.5 * eq) / len(values)


def zscore(values: Sequence[float], x: float) -> float | None:
    if len(values) < 2:
        return None
    sd = statistics.pstdev(values)
    if sd == 0:
        return 0.0
    return (x - statistics.fmean(values)) / sd


def lookback_value(items: Sequence[tuple[date, float]], n_obs: int) -> float | None:
    if len(items) <= n_obs:
        return None
    return items[-1 - n_obs][1]


def rate_metrics(series: Mapping[date, float]) -> dict:
    items = sorted(series.items())
    if not items:
        raise SalError("empty rate series")
    latest_date, latest = items[-1]
    vals_1y = [v for _, v in items[-252:]]
    vals_5y = [v for _, v in items[-1260:]]

    def bp(n: int) -> float | None:
        old = lookback_value(items, n)
        return None if old is None else (latest - old) * 100.0

    return {
        "as_of": latest_date.isoformat(),
        "value": latest,
        "bp_1d": bp(1),
        "bp_5d": bp(5),
        "bp_1m": bp(21),
        "bp_3m": bp(63),
        "pctile_1y": percentile_rank(vals_1y, latest),
        "pctile_5y": percentile_rank(vals_5y, latest),
        "z_1y": zscore(vals_1y, latest),
        "z_5y": zscore(vals_5y, latest),
    }


def fx_metrics(series: Mapping[date, float]) -> dict:
    items = sorted(series.items())
    if not items:
        raise SalError("empty FX series")
    latest_date, latest = items[-1]
    vals = [v for _, v in items]
    vals_1y = vals[-252:]
    vals_5y = vals[-1260:]

    def ret(n: int) -> float | None:
        old = lookback_value(items, n)
        return None if old in (None, 0) else 100.0 * (latest / old - 1.0)

    def sma(n: int) -> float | None:
        if len(vals) < n:
            return None
        return statistics.fmean(vals[-n:])

    def ma_dist(n: int) -> float | None:
        m = sma(n)
        return None if not m else 100.0 * (latest / m - 1.0)

    logrets = [math.log(vals[i] / vals[i - 1]) for i in range(1, len(vals)) if vals[i] > 0 and vals[i - 1] > 0]

    def rv(n: int) -> float | None:
        if len(logrets) < n:
            return None
        return statistics.pstdev(logrets[-n:]) * math.sqrt(252.0) * 100.0

    return {
        "as_of": latest_date.isoformat(),
        "spot": latest,
        "ret_1d": ret(1),
        "ret_5d": ret(5),
        "ret_1m": ret(21),
        "ret_3m": ret(63),
        "ret_6m": ret(126),
        "ret_1y": ret(252),
        "rv_20d": rv(20),
        "rv_60d": rv(60),
        "pctile_1y": percentile_rank(vals_1y, latest),
        "pctile_5y": percentile_rank(vals_5y, latest),
        "z_1y": zscore(vals_1y, latest),
        "z_5y": zscore(vals_5y, latest),
        "ma20_dist_pct": ma_dist(20),
        "ma50_dist_pct": ma_dist(50),
        "ma200_dist_pct": ma_dist(200),
        "high_52w": max(vals_1y) if vals_1y else None,
        "low_52w": min(vals_1y) if vals_1y else None,
    }


def parse_treasury_csv(text: str) -> dict[str, dict[date, float]]:
    out = {"2Y": {}, "5Y": {}, "10Y": {}, "30Y": {}}
    rows = csv.DictReader(io.StringIO(text))
    keymap = {"2Y": "2 Yr", "5Y": "5 Yr", "10Y": "10 Yr", "30Y": "30 Yr"}
    for row in rows:
        d = parse_date(row.get("Date"))
        if not d:
            continue
        for tenor, col in keymap.items():
            v = to_float(row.get(col))
            if v is not None:
                out[tenor][d] = v
    if any(not out[k] for k in ("2Y", "5Y", "10Y")):
        raise SalError("Treasury CSV missing required tenors")
    return out


def fetch_us_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    merged = {"2Y": {}, "5Y": {}, "10Y": {}, "30Y": {}}
    for year in range(start.year, end.year + 1):
        text = fetch_bytes(TREASURY_URL.format(year=year)).decode("utf-8-sig")
        parsed = parse_treasury_csv(text)
        for tenor, series in parsed.items():
            merged[tenor].update(series)
    return {k: filter_since(v, start) for k, v in merged.items()}


def parse_boc_json(payload: Mapping) -> dict[str, dict[date, float]]:
    series_map = {"V39051": "2Y", "V39053": "5Y", "V39055": "10Y", "V39056": "LONG"}
    out = {v: {} for v in series_map.values()}
    for obs in payload.get("observations", []):
        d = parse_date(obs.get("d"))
        if not d:
            continue
        for sid, tenor in series_map.items():
            cell = obs.get(sid)
            value = cell.get("v") if isinstance(cell, Mapping) else cell
            v = to_float(value)
            if v is not None:
                out[tenor][d] = v
    if any(not out[k] for k in RATE_TENORS):
        raise SalError("BoC response missing required benchmark tenors")
    return out


def fetch_ca_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    qs = urllib.parse.urlencode({"start_date": start.isoformat(), "end_date": end.isoformat()})
    payload = json.loads(fetch_bytes(f"{BOC_URL}?{qs}").decode("utf-8"))
    return parse_boc_json(payload)


def parse_rba_csv(text: str) -> dict[str, dict[date, float]]:
    rows = list(csv.reader(io.StringIO(text)))
    metadata: dict[str, list[str]] = {}
    series_row_idx = None
    for i, row in enumerate(rows):
        if not row:
            continue
        first = row[0].strip().lower()
        if first in {"title", "description", "frequency", "type", "units", "source", "publication date", "series id"}:
            metadata[first] = row
        if first == "series id":
            series_row_idx = i
            break
    if series_row_idx is None:
        raise SalError("RBA F2 CSV missing Series ID row")
    labels = metadata.get("description") or metadata.get("title")
    if not labels:
        raise SalError("RBA F2 CSV missing descriptive metadata")
    selected: dict[int, str] = {}
    for idx in range(1, len(labels)):
        label = labels[idx].strip().lower() if idx < len(labels) else ""
        if "australian government" not in label and "government" not in label:
            continue
        for n, tenor in ((2, "2Y"), (5, "5Y"), (10, "10Y")):
            if re.search(rf"\b{n}\s*[- ]?year\b", label):
                selected[idx] = tenor
    if set(selected.values()) != set(RATE_TENORS):
        raise SalError(f"RBA F2 tenor discovery failed: {selected}")
    out = {k: {} for k in RATE_TENORS}
    for row in rows[series_row_idx + 1 :]:
        if not row:
            continue
        d = parse_date(row[0])
        if not d:
            continue
        for idx, tenor in selected.items():
            if idx < len(row):
                v = to_float(row[idx])
                if v is not None:
                    out[tenor][d] = v
    if any(not out[k] for k in RATE_TENORS):
        raise SalError("RBA F2 data rows missing required tenors")
    return out


def fetch_au_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    parsed = parse_rba_csv(fetch_bytes(RBA_URL).decode("utf-8-sig", errors="replace"))
    return {k: {d: v for d, v in s.items() if start <= d <= end} for k, s in parsed.items()}


def parse_rbnz_xlsx(data: bytes) -> dict[str, dict[date, float]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise SalError("openpyxl is required for the RBNZ workbook") from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        preview = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True))
        for ri, row in enumerate(preview):
            if not row:
                continue
            date_cols = [i for i, v in enumerate(row) if str(v).strip().lower() == "date"]
            if not date_cols:
                continue
            header_start = max(0, ri - 5)
            col_headers: dict[int, str] = {}
            max_cols = max(len(r) for r in preview[header_start : ri + 1])
            header_rows = preview[header_start : ri + 1]
            filled_rows = []
            for rr in header_rows:
                filled = []
                carry = None
                for ci in range(max_cols):
                    val = rr[ci] if ci < len(rr) else None
                    if val not in (None, ""):
                        carry = val
                    filled.append(carry)
                filled_rows.append(filled)
            for ci in range(max_cols):
                parts = []
                for rr in filled_rows:
                    if rr[ci] not in (None, ""):
                        parts.append(str(rr[ci]).strip())
                col_headers[ci] = " | ".join(parts).lower()
            selected: dict[int, str] = {}
            for ci, header in col_headers.items():
                if "government" not in header and "bond" not in header:
                    continue
                for n, tenor in ((2, "2Y"), (5, "5Y"), (10, "10Y")):
                    if re.search(rf"\b{n}\s*year\b", header):
                        selected[ci] = tenor
            if set(selected.values()) == set(RATE_TENORS):
                best = (ws, ri + 1, selected)
                break
        if best:
            break
    if not best:
        raise SalError("RBNZ workbook tenor discovery failed")
    ws, header_row, selected = best
    out = {k: {} for k in RATE_TENORS}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row:
            continue
        d = None
        for v in row[:3]:
            d = parse_date(v)
            if d:
                break
        if not d:
            continue
        for ci, tenor in selected.items():
            if ci < len(row):
                v = to_float(row[ci])
                if v is not None:
                    out[tenor][d] = v
    if any(not out[k] for k in RATE_TENORS):
        raise SalError("RBNZ workbook missing required government tenors")
    return out


def fetch_nz_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    parsed = parse_rbnz_xlsx(fetch_bytes(RBNZ_URL))
    return {k: {d: v for d, v in s.items() if start <= d <= end} for k, s in parsed.items()}


def parse_ecb_fx_csv(text: str) -> dict[str, dict[date, float]]:
    reader = csv.DictReader(io.StringIO(text))
    fieldnames = [x.strip() for x in (reader.fieldnames or [])]
    missing = [c for c in G10 if c != "EUR" and c not in fieldnames]
    if missing:
        raise SalError(f"ECB FX CSV missing G10 currencies: {missing}")
    currency = {c: {} for c in G10}
    for row in reader:
        clean = {str(k).strip(): v for k, v in row.items() if k is not None}
        d = parse_date(clean.get("Date"))
        if not d:
            continue
        currency["EUR"][d] = 1.0
        for c in G10:
            if c == "EUR":
                continue
            v = to_float(clean.get(c))
            if v is not None and v > 0:
                currency[c][d] = v
    if any(not currency[c] for c in G10):
        raise SalError("ECB FX history missing one or more G10 currency histories")
    return currency


def build_fx_crosses(currency_per_eur: Mapping[str, Mapping[date, float]], start: date) -> dict[str, dict[date, float]]:
    out: dict[str, dict[date, float]] = {}
    for i, base in enumerate(FX_ORDER):
        for quote in FX_ORDER[i + 1 :]:
            common = sorted(set(currency_per_eur[base]) & set(currency_per_eur[quote]))
            series = {
                d: currency_per_eur[quote][d] / currency_per_eur[base][d]
                for d in common
                if d >= start and currency_per_eur[base][d] > 0
            }
            out[f"{base}{quote}"] = series
    if len(out) != 45 or any(not s for s in out.values()):
        raise SalError(f"expected 45 populated G10 crosses, got {len(out)}")
    return out


def fetch_fx(start: date) -> dict[str, dict[date, float]]:
    text = fetch_bytes(ECB_FX_URL).decode("utf-8-sig")
    return build_fx_crosses(parse_ecb_fx_csv(text), start)


def curve_spread(short: Mapping[date, float], long: Mapping[date, float]) -> dict[date, float]:
    return {d: (long[d] - short[d]) * 100.0 for d in set(short) & set(long)}


def rv_spread(a: Mapping[date, float], b: Mapping[date, float]) -> dict[date, float]:
    return {d: (a[d] - b[d]) * 100.0 for d in set(a) & set(b)}


def spread_metrics_bps(series: Mapping[date, float]) -> dict:
    items = sorted(series.items())
    if not items:
        raise SalError("empty spread series")
    latest_date, latest = items[-1]
    vals_1y = [v for _, v in items[-252:]]
    vals_5y = [v for _, v in items[-1260:]]

    def chg(n: int) -> float | None:
        old = lookback_value(items, n)
        return None if old is None else latest - old

    return {
        "as_of": latest_date.isoformat(),
        "bps": latest,
        "chg_1d_bps": chg(1),
        "chg_5d_bps": chg(5),
        "chg_1m_bps": chg(21),
        "chg_3m_bps": chg(63),
        "pctile_1y": percentile_rank(vals_1y, latest),
        "pctile_5y": percentile_rank(vals_5y, latest),
        "z_1y": zscore(vals_1y, latest),
        "z_5y": zscore(vals_5y, latest),
    }


def build_snapshot(*, today: date | None = None) -> dict:
    today = today or datetime.now(timezone.utc).date()
    start = today - timedelta(days=366 * 5 + 15)
    rates_raw = {
        "US": fetch_us_rates(start, today),
        "CA": fetch_ca_rates(start, today),
        "AU": fetch_au_rates(start, today),
        "NZ": fetch_nz_rates(start, today),
    }
    rates: dict[str, dict] = {}
    for country, tenors in rates_raw.items():
        metrics = {tenor: rate_metrics(series) for tenor, series in tenors.items() if series}
        curves = {
            "2s10s": spread_metrics_bps(curve_spread(tenors["2Y"], tenors["10Y"])),
            "5s10s": spread_metrics_bps(curve_spread(tenors["5Y"], tenors["10Y"])),
        }
        latest_as_of = max(date.fromisoformat(metrics[t]["as_of"]) for t in RATE_TENORS)
        rates[country] = {
            "tenors": metrics,
            "curves": curves,
            "latest_observation": latest_as_of.isoformat(),
            "age_days": max(0, (today - latest_as_of).days),
        }

    rate_rv: dict[str, dict] = {}
    for a, b in RV_PAIRS:
        for tenor in RATE_TENORS:
            key = f"{a}-{b}_{tenor}"
            rate_rv[key] = spread_metrics_bps(rv_spread(rates_raw[a][tenor], rates_raw[b][tenor]))

    fx_raw = fetch_fx(start)
    fx = {pair: fx_metrics(series) for pair, series in fx_raw.items()}
    latest_fx_date = max(date.fromisoformat(v["as_of"]) for v in fx.values())

    return {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "window_start": start.isoformat(),
        "rates": rates,
        "rate_rv": rate_rv,
        "fx": {
            "source_observation": latest_fx_date.isoformat(),
            "age_days": max(0, (today - latest_fx_date).days),
            "pairs": fx,
        },
        "sources": {
            "US_rates": {"name": "U.S. Treasury", "url": "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?type=daily_treasury_yield_curve"},
            "CA_rates": {"name": "Bank of Canada Valet", "url": "https://www.bankofcanada.ca/valet/docs/"},
            "AU_rates": {"name": "Reserve Bank of Australia F2", "url": "https://www.rba.gov.au/statistics/tables/", "note": "RBA assessed closing yields; research context only; not a financial benchmark."},
            "NZ_rates": {"name": "Reserve Bank of New Zealand B2", "url": "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/wholesale-interest-rates"},
            "FX": {"name": "ECB euro foreign exchange reference rates", "url": "https://www.ecb.europa.eu/stats/policy_and_exchange_rates/euro_reference_exchange_rates/html/index.en.html", "note": "Reference rates for information purposes; cross rates are deterministic ratios of same-fixing ECB EUR reference rates."},
        },
        "method": {
            "model_calls": 0,
            "fx_pair_count": len(fx),
            "rate_rv_count": len(rate_rv),
            "note": "Live authorities remain canonical. This artifact contains only the latest derived research snapshot; no historical warehouse is created in GitHub or Supabase.",
        },
    }


def fnum(v: float | None, digits: int = 2, suffix: str = "") -> str:
    return "—" if v is None else f"{v:.{digits}f}{suffix}"


def signnum(v: float | None, digits: int = 1, suffix: str = "") -> str:
    if v is None:
        return "—"
    return f"{v:+.{digits}f}{suffix}"


def render_html(snapshot: Mapping) -> str:
    rates = snapshot["rates"]
    rows = []
    for c in RATE_COUNTRIES:
        t = rates[c]["tenors"]
        extra = t.get("30Y") or t.get("LONG")
        rows.append(
            f'<tr><th>{c}</th><td>{fnum(t["2Y"]["value"],2,"%")}</td>'
            f'<td>{fnum(t["5Y"]["value"],2,"%")}</td><td>{fnum(t["10Y"]["value"],2,"%")}</td>'
            f'<td>{fnum(extra["value"],2,"%") if extra else "—"}</td>'
            f'<td>{signnum(rates[c]["curves"]["2s10s"]["bps"],0,"bp")}</td>'
            f'<td>{rates[c]["latest_observation"]}</td></tr>'
        )

    rv_keys = ["CA-US_2Y", "CA-US_5Y", "AU-US_2Y", "AU-US_5Y", "NZ-US_2Y", "AU-NZ_2Y", "AU-NZ_5Y"]
    rv_rows = []
    for key in rv_keys:
        m = snapshot["rate_rv"][key]
        flag = " stretch" if (m["pctile_5y"] is not None and (m["pctile_5y"] <= 10 or m["pctile_5y"] >= 90)) else ""
        rv_rows.append(
            f'<tr class="{flag.strip()}"><th>{key.replace("_", " ")}</th><td>{signnum(m["bps"],0,"bp")}</td>'
            f'<td>{signnum(m["chg_1m_bps"],0,"bp")}</td><td>{fnum(m["pctile_5y"],0,"th")}</td><td>{m["as_of"]}</td></tr>'
        )

    focus_fx = ["EURUSD", "GBPUSD", "USDJPY", "USDCHF", "USDCAD", "AUDUSD", "NZDUSD", "AUDNZD", "AUDCAD", "NZDCAD", "CADJPY", "AUDJPY", "NZDJPY", "NOKSEK"]
    fx_rows = []
    pairs = snapshot["fx"]["pairs"]
    for pair in focus_fx:
        if pair not in pairs:
            continue
        m = pairs[pair]
        fx_rows.append(
            f'<tr><th>{pair[:3]}/{pair[3:]}</th><td>{fnum(m["spot"],4)}</td><td>{signnum(m["ret_1d"],2,"%")}</td>'
            f'<td>{signnum(m["ret_1m"],2,"%")}</td><td>{signnum(m["ret_3m"],2,"%")}</td>'
            f'<td>{fnum(m["pctile_5y"],0,"th")}</td><td>{fnum(m["rv_20d"],1,"%")}</td></tr>'
        )

    all_fx_rows = []
    for pair in sorted(pairs):
        m = pairs[pair]
        all_fx_rows.append(
            f'<tr><th>{pair[:3]}/{pair[3:]}</th><td>{fnum(m["spot"],5)}</td><td>{signnum(m["ret_5d"],2,"%")}</td>'
            f'<td>{signnum(m["ret_1m"],2,"%")}</td><td>{signnum(m["ret_3m"],2,"%")}</td><td>{signnum(m["ret_1y"],2,"%")}</td>'
            f'<td>{fnum(m["pctile_5y"],0,"th")}</td></tr>'
        )

    return f'''<section class="sal-market" id="sal-market-data">
<div class="stitle">Rates &amp; G10 FX · Deterministic Market Context</div>
<div class="sal-meta">Sal V1 · zero model calls · rates: Treasury / BoC / RBA / RBNZ · FX: ECB daily reference rates · generated {snapshot["generated_at"][:16].replace("T"," ")} UTC</div>
<div class="sal-grid">
<div class="sal-card"><div class="sal-card-title">Sovereign Rates Monitor</div><table><thead><tr><th></th><th>2Y</th><th>5Y</th><th>10Y</th><th>Long</th><th>2s10s</th><th>As of</th></tr></thead><tbody>{''.join(rows)}</tbody></table><div class="sal-note">Australia: RBA assessed closing yields, research context only, not a financial benchmark.</div></div>
<div class="sal-card"><div class="sal-card-title">Rates Relative Value</div><table><thead><tr><th>Spread</th><th>Now</th><th>1M Δ</th><th>5Y %ile</th><th>As of</th></tr></thead><tbody>{''.join(rv_rows)}</tbody></table><div class="sal-note">Statistical stretch only: ≤10th or ≥90th percentile. It is not a trade signal or fair-value estimate.</div></div>
</div>
<div class="sal-card sal-fx"><div class="sal-card-title">G10 FX Focus</div><table><thead><tr><th>Pair</th><th>Spot</th><th>1D</th><th>1M</th><th>3M</th><th>5Y %ile</th><th>20D RV</th></tr></thead><tbody>{''.join(fx_rows)}</tbody></table>
<details><summary>All 45 G10 crosses</summary><div class="sal-scroll"><table><thead><tr><th>Pair</th><th>Spot</th><th>5D</th><th>1M</th><th>3M</th><th>1Y</th><th>5Y %ile</th></tr></thead><tbody>{''.join(all_fx_rows)}</tbody></table></div></details>
<div class="sal-note">FX values are research reference rates, not executable prices. Crosses are deterministic ratios of same-fixing ECB EUR reference rates.</div></div>
</section>'''


def validate_snapshot(s: Mapping) -> None:
    if s.get("method", {}).get("model_calls") != 0:
        raise SalError("Sal must make zero model calls")
    if s.get("method", {}).get("fx_pair_count") != 45:
        raise SalError("Sal must produce exactly 45 G10 FX crosses")
    if s.get("method", {}).get("rate_rv_count") != len(RV_PAIRS) * len(RATE_TENORS):
        raise SalError("unexpected rate RV count")
    for c in RATE_COUNTRIES:
        for t in RATE_TENORS:
            value = s["rates"][c]["tenors"][t]["value"]
            if not (0 < value < 25):
                raise SalError(f"implausible {c} {t} yield: {value}")
    for pair, m in s["fx"]["pairs"].items():
        if not (0 < m["spot"] < 1000):
            raise SalError(f"implausible FX spot {pair}: {m['spot']}")


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", required=True, help="JSON snapshot path")
    ap.add_argument("--html", help="optional rendered dashboard fragment")
    ap.add_argument("--today", help="override current date (YYYY-MM-DD), useful for testing")
    args = ap.parse_args(argv)
    today = date.fromisoformat(args.today) if args.today else None
    snapshot = build_snapshot(today=today)
    validate_snapshot(snapshot)
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(snapshot, indent=2, sort_keys=True) + "\n")
    if args.html:
        hp = Path(args.html)
        hp.parent.mkdir(parents=True, exist_ok=True)
        hp.write_text(render_html(snapshot) + "\n")
    print(json.dumps({
        "status": "ok",
        "fx_pairs": snapshot["method"]["fx_pair_count"],
        "rate_rv": snapshot["method"]["rate_rv_count"],
        "fx_as_of": snapshot["fx"]["source_observation"],
        "rate_as_of": {c: snapshot["rates"][c]["latest_observation"] for c in RATE_COUNTRIES},
    }))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SalError as exc:
        print(f"SAL ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
