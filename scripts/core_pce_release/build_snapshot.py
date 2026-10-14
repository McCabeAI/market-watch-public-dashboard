"""Build the immutable 2026-09-30 BEA snapshot from official workbooks.

Not on the release hot path. Requires openpyxl. The committed JSON is what
prepare/run load.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from decimal import Decimal
from pathlib import Path

from scripts.core_pce_release.classify import classify_source

OFFICIAL_CORE_LINE = 374
TABLE_LAST_LINE = 368
EXCLUSION_LINES = (73, 113, 169)
CORE_MILLIONS = Decimal("19556763")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_expenditures(path: Path) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["U20405-Q"]
    rows: list[dict] = []
    header = None
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if index == 8:
            header = row
        if index < 9 or row[0] is None:
            continue
        try:
            line = int(row[0])
        except (TypeError, ValueError):
            continue
        if line > TABLE_LAST_LINE:
            continue
        name = row[1] or ""
        value = row[272]
        numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
        rows.append(
            {
                "line": line,
                "name": str(name).strip(),
                "indent": len(name) - len(str(name).lstrip(" ")),
                "code": row[2],
                "q2": Decimal(str(value)) if numeric else None,
                "suppressed": not numeric,
            }
        )
    if not rows or header is None or header[272] != "2026Q2":
        raise SystemExit("U20405-Q did not end at 2026Q2 in column 272")
    rows[0]["indent"] = -1
    return rows


def _load_sources(path: Path) -> dict[int, dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook.active
    dated = sheet.cell(1, 4).value
    sources: dict[int, dict] = {}
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if index < 5 or row[0] is None:
            continue
        try:
            item = int(row[0])
        except (TypeError, ValueError):
            continue
        sources[item] = {
            "name": str(row[1]).strip() if row[1] else "",
            "bea_series_code": row[2],
            "source_text": None if row[3] is None else str(row[3]).strip(),
            "seasonal_status": None if row[4] is None else str(row[4]).strip(),
        }
    return {"dated": str(dated.date()) if hasattr(dated, "date") else str(dated), "rows": sources}


def _children_index(rows: list[dict]) -> dict[int, list[int]]:
    children: dict[int, list[int]] = {i: [] for i in range(len(rows))}
    stack: list[tuple[int, int]] = []
    for index, row in enumerate(rows):
        while stack and stack[-1][0] >= row["indent"]:
            stack.pop()
        if stack:
            children[stack[-1][1]].append(index)
        stack.append((row["indent"], index))
    return children


def _is_less(row: dict) -> bool:
    return row["name"].lower().startswith("less:")


def _select(rows: list[dict], children: dict[int, list[int]], index: int, sign: int) -> list[tuple[int, int]]:
    row = rows[index]
    kids = children[index]
    if not kids or row["q2"] is None:
        return [(index, sign)]
    if any(rows[child]["q2"] is None for child in kids):
        return [(index, sign)]
    kid_sum = sum((rows[child]["q2"] for child in kids), Decimal(0))
    if abs(kid_sum - row["q2"]) <= Decimal("1.5"):
        selected: list[tuple[int, int]] = []
        for child in kids:
            selected.extend(_select(rows, children, child, sign))
        return selected
    less = [child for child in kids if _is_less(rows[child])]
    positive = [child for child in kids if not _is_less(rows[child])]
    if less and positive:
        net = sum((rows[child]["q2"] for child in positive), Decimal(0)) - sum(
            (rows[child]["q2"] for child in less), Decimal(0)
        )
        if abs(net - row["q2"]) <= Decimal("1.5"):
            selected = []
            for child in positive:
                selected.extend(_select(rows, children, child, sign))
            for child in less:
                selected.extend(_select(rows, children, child, -sign))
            return selected
    return [(index, sign)]


def _subtree(children: dict[int, list[int]], index: int) -> set[int]:
    found = {index}
    for child in children[index]:
        found |= _subtree(children, child)
    return found


def _ancestors(parent_of: dict[int, int], index: int) -> list[int]:
    found: list[int] = []
    cursor = parent_of.get(index)
    while cursor is not None:
        found.append(cursor)
        cursor = parent_of.get(cursor)
    return found


def _bridge_disposition(bridge_path: Path, components: list[dict]) -> list[dict]:
    by_line = {row["line"]: row for row in components}
    dispositions = []
    with bridge_path.open(newline="") as handle:
        for record in csv.DictReader(handle):
            raw_lines = str(record["pce_line"])
            lines = []
            if "-" in raw_lines:
                start, end = raw_lines.split("-", 1)
                lines = list(range(int(start), int(end) + 1))
            elif raw_lines.strip():
                lines = [int(raw_lines)]
            aug = (record.get("aug_2026_mom_pct") or "").strip()
            basis = record.get("observation_basis") or ""
            for line in lines:
                component = by_line.get(line)
                reason = "line_not_in_core_partition"
                accepted = False
                if "PPI" in basis:
                    basis_family = "PPI"
                elif "CPI" in basis or "Table 2" in basis:
                    basis_family = "CPI"
                else:
                    basis_family = None
                if component is None:
                    reason = "line_not_in_current_core_partition"
                elif not aug:
                    reason = "no_august_mom"
                elif component["source_family"] == "IMPLICIT":
                    reason = "implicit_source_not_a_ppi_carry"
                elif basis_family and basis_family != component["source_family"]:
                    reason = "observation_basis_does_not_match_current_source"
                elif component["policy"] != "bls_sa_direct":
                    reason = "seasonal_or_method_not_direct_sa:" + component["policy"]
                else:
                    accepted = True
                    reason = "accepted_verified_august_sa_mom"
                dispositions.append(
                    {
                        "bridge_component": record["component"],
                        "line": line,
                        "bridge_q2_millions": record["q2_2026_current_dollar_millions"],
                        "bridge_aug_mom_pct": aug or None,
                        "observation_basis": basis,
                        "current_source_family": None if component is None else component["source_family"],
                        "current_policy": None if component is None else component["policy"],
                        "current_source_text": None if component is None else component["source_text"],
                        "accepted_as_september_carry": accepted,
                        "reason": reason,
                    }
                )
    return dispositions


def build(expenditure_xlsx: Path, source_xlsx: Path, fred_csv: Path, bridge_csv: Path, out_dir: Path) -> None:
    rows = _load_expenditures(expenditure_xlsx)
    source_book = _load_sources(source_xlsx)
    sources = source_book["rows"]
    children = _children_index(rows)
    parent_of: dict[int, int] = {}
    for parent, kids in children.items():
        for child in kids:
            parent_of[child] = parent
    selected = _select(rows, children, 0, 1)
    excluded: set[int] = set()
    by_line_index = {row["line"]: index for index, row in enumerate(rows)}
    for line in EXCLUSION_LINES:
        excluded |= _subtree(children, by_line_index[line])
    core_units = [(index, sign) for index, sign in selected if index not in excluded]
    code_to_source_lines: dict[str, list[int]] = {}
    for line, source in sources.items():
        code = rows[by_line_index[line]]["code"] if line in by_line_index else None
        if code:
            code_to_source_lines.setdefault(str(code), []).append(line)

    components = []
    signed = Decimal(0)
    for index, sign in core_units:
        row = rows[index]
        signed += Decimal(sign) * row["q2"]
        source = sources.get(row["line"])
        inherited_from = None
        inheritance = None
        if source is None:
            donors = [
                line
                for line in code_to_source_lines.get(str(row["code"]), [])
                if line != row["line"] and sources[line].get("source_text")
            ]
            usable = []
            for line in donors:
                family, _policy = classify_source(sources[line]["source_text"], sources[line]["seasonal_status"])
                if family not in {"POINTER", "MISSING"}:
                    usable.append(line)
            if len(usable) == 1:
                inherited_from = usable[0]
                source = sources[inherited_from]
                inheritance = "nipa_code_match_under_table_2_4_4_line_356"
        source_text = None if source is None else source.get("source_text")
        seasonal = None if source is None else source.get("seasonal_status")
        family, policy = classify_source(source_text, seasonal)
        weight = (Decimal(sign) * row["q2"]) / CORE_MILLIONS
        ancestor_lines = [rows[item]["line"] for item in _ancestors(parent_of, index)]
        components.append(
            {
                "line": row["line"],
                "name": row["name"],
                "nipa_code": row["code"],
                "bea_series_code": None if source is None else source.get("bea_series_code"),
                "q2_2026_millions": format(row["q2"], "f"),
                "sign": sign,
                "weight_share": format(weight, "f"),
                "source_text": source_text,
                "seasonal_status": seasonal,
                "source_family": family,
                "policy": policy,
                "inherited_from_line": inherited_from,
                "inheritance": inheritance,
                "ancestor_lines": ancestor_lines,
            }
        )

    components.sort(key=lambda item: item["line"])
    out_dir.mkdir(parents=True, exist_ok=True)
    fred_dest = out_dir / "pcepilfe.csv"
    shutil.copyfile(fred_csv, fred_dest)
    component_path = out_dir / "components.json"
    component_path.write_text(json.dumps({"components": components}, indent=2) + "\n", encoding="utf-8")
    dispositions = _bridge_disposition(bridge_csv, components)
    (out_dir / "august_bridge_disposition.json").write_text(
        json.dumps({"rows": dispositions}, indent=2) + "\n", encoding="utf-8"
    )

    def _weight(policy: str | None = None, family: str | None = None) -> Decimal:
        total = Decimal(0)
        for row in components:
            if policy is not None and row["policy"] != policy:
                continue
            if family is not None and row["source_family"] != family:
                continue
            total += Decimal(row["weight_share"])
        return total

    residual = CORE_MILLIONS - signed
    manifest = {
        "schema": "core_pce_snapshot_v1",
        "source_map": {
            "table": "BEA Table 2.4.4 U-D",
            "dated": source_book["dated"],
            "url": "https://www.bea.gov/sites/default/files/2026-10/FAQ521-Sep26-update.xlsx",
            "sha256": _sha256(source_xlsx),
        },
        "expenditures": {
            "table": "BEA Table 2.4.5U",
            "sheet": "U20405-Q",
            "period": "2026Q2",
            "published": "2026-09-30",
            "url": "https://apps.bea.gov/national/Release/XLS/Underlying/Section2All_xls.xlsx",
            "sha256": _sha256(expenditure_xlsx),
            "official_core_pce_line": OFFICIAL_CORE_LINE,
            "official_core_pce_millions": format(CORE_MILLIONS, "f"),
            "signed_partition_millions": format(signed, "f"),
            "rounding_residual_millions": format(residual, "f"),
            "food_energy_exclusion_lines": list(EXCLUSION_LINES),
            "note": "Residual is BEA million-dollar rounding on net identities. Weights are not renormalized.",
        },
        "pcepilfe": {
            "series": "PCEPILFE",
            "url": "https://fred.stlouisfed.org/graph/fredgraph.csv?id=PCEPILFE",
            "retrieved_on": "2026-10-14",
            "sha256": _sha256(fred_dest),
            "fingerprint": {
                "2025-11": "127.363",
                "2025-12": "127.765",
                "2026-02": "128.701",
                "2026-03": "129.030",
                "2026-05": "129.796",
                "2026-06": "129.969",
                "2026-07": "130.133",
                "2026-08": "130.455",
            },
            "not_used": "data/temperature_history/raw/us/PCEPILFE.csv is a different vintage and ends 2026-07",
        },
        "component_count": len(components),
        "weight_share_sum": format(sum((Decimal(row["weight_share"]) for row in components), Decimal(0)), "f"),
        "bls_sa_direct_weight_share": format(_weight(policy="bls_sa_direct"), "f"),
        "accepted_august_bridge_carries": sum(1 for row in dispositions if row["accepted_as_september_carry"]),
        "components_sha256": _sha256(component_path),
    }
    (out_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: manifest[k] for k in ("component_count", "weight_share_sum", "bls_sa_direct_weight_share", "accepted_august_bridge_carries")}, indent=2))
    print("residual", manifest["expenditures"]["rounding_residual_millions"])


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the core PCE release snapshot")
    parser.add_argument("--expenditures", type=Path, required=True)
    parser.add_argument("--source-map", type=Path, required=True)
    parser.add_argument("--fred", type=Path, required=True)
    parser.add_argument("--bridge", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    build(args.expenditures, args.source_map, args.fred, args.bridge, args.out)


if __name__ == "__main__":
    main()
