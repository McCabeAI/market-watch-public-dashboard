#!/usr/bin/env python3
"""Deterministic Market Watch market-state generator.

No model calls. No API keys. Live official/ECB authorities remain the source
of truth. The script fetches sovereign yield history and ECB G10 FX reference
history, computes a compact research snapshot, and exits non-zero if a
required source is missing or fabricated substitution would be required.

This is research/reference state, not executable pricing.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import statistics
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Mapping, Sequence

from scripts.australia_housing_data import collect_australia_housing, validate_australia_housing
from scripts.cross_asset_data import collect_cross_assets
from scripts.market_opportunities import build_opportunities
from scripts.official_curve_data import collect_official_curves, validate_official_curves
from scripts.positioning_data import build_positioning, validate_positioning
from scripts.policy_path_data import (
    build_tradable_rate_curves,
    collect_policy_paths,
    validate_policy_paths,
    validate_tradable_rate_curves,
)

USER_AGENT = (
    "MarketWatch-MarketState/1.0 "
    "(+https://github.com/McCabeAI/market-watch-public-dashboard)"
)
BROWSER_USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
DEFAULT_TIMEOUT = 45
RETRIES = 3
ECB_TIMEOUT = 90
ECB_RETRIES = 4

G10 = ("EUR", "GBP", "AUD", "NZD", "USD", "CAD", "CHF", "NOK", "SEK", "JPY")
FX_ORDER = G10
RATE_TENORS = ("2Y", "5Y", "10Y")
RATE_COUNTRIES = ("US", "CA", "AU", "NZ")
RV_PAIRS = (("CA", "US"), ("AU", "US"), ("NZ", "US"), ("AU", "NZ"), ("CA", "AU"))

# Daily sources may miss weekends/holidays. RBA F2 is weekly with a 2-business-day lag.
STALE_AFTER_DAYS = {"US": 4, "CA": 4, "AU": 12, "NZ": 4, "FX": 4}
FAIL_AFTER_DAYS = 21

TREASURY_URL = (
    "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/"
    "daily-treasury-rates.csv/{year}/all?_format=csv&field_tdr_date_value={year}"
    "&page=&type=daily_treasury_yield_curve"
)
TREASURY_PAGE = (
    "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/"
    "TextView?type=daily_treasury_yield_curve"
)
BOC_SERIES = {
    "BD.CDN.2YR.DQ.YLD": "2Y",
    "BD.CDN.5YR.DQ.YLD": "5Y",
    "BD.CDN.10YR.DQ.YLD": "10Y",
    "BD.CDN.LONG.DQ.YLD": "LONG",
}
BOC_GROUP_URL = "https://www.bankofcanada.ca/valet/observations/group/bond_yields_benchmark/json"
BOC_SERIES_URL = "https://www.bankofcanada.ca/valet/observations/{series}/json"
BOC_DOCS = "https://www.bankofcanada.ca/valet/docs/"
RBA_URL = "https://www.rba.gov.au/statistics/tables/csv/f2-data.csv"
RBA_PAGE = "https://www.rba.gov.au/statistics/tables/"
RBNZ_URL = "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx"
RBNZ_PAGE = (
    "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/"
    "wholesale-interest-rates"
)
ECB_FX_SDMX_URL = (
    "https://data-api.ecb.europa.eu/service/data/EXR/"
    "D.{currencies}.EUR.SP00.A"
)
ECB_FX_PAGE = (
    "https://www.ecb.europa.eu/stats/policy_and_exchange_rates/"
    "euro_reference_exchange_rates/html/index.en.html"
)
# Classic hist CSV is retained only as a parser test fixture. The public
# eurofxref-hist.csv file is not used in production because CDN copies have
# been observed stale or corrupted; SDMX is the live official series.


class MarketStateError(RuntimeError):
    """Hard failure: a required official source or calculation is unusable."""


def fetch_bytes(
    url: str,
    *,
    timeout: int = DEFAULT_TIMEOUT,
    retries: int = RETRIES,
    user_agent: str | None = None,
    referer: str | None = None,
) -> bytes:
    last: Exception | None = None
    headers = {
        "User-Agent": user_agent or USER_AGENT,
        "Accept": "*/*",
        "Accept-Language": "en-US,en;q=0.8",
    }
    if referer:
        headers["Referer"] = referer
    if "rbnz.govt.nz" in url:
        headers["Referer"] = RBNZ_PAGE
        headers["Accept"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*"
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                status = getattr(resp, "status", 200)
                if status != 200:
                    raise MarketStateError(f"HTTP {status} from {url}")
                data = resp.read()
                if not data:
                    raise MarketStateError(f"empty response from {url}")
                return data
        except urllib.error.HTTPError as exc:
            last = exc
            if attempt + 1 < retries:
                # 5xx and timeouts are retried; 403/404 are still retried a bounded number of times
                time.sleep(2.0 * (attempt + 1) if exc.code >= 500 else 1.5 * (attempt + 1))
        except (urllib.error.URLError, TimeoutError, MarketStateError) as exc:
            last = exc
            if attempt + 1 < retries:
                time.sleep(1.5 * (attempt + 1))
    if isinstance(last, urllib.error.HTTPError) and last.code == 403 and "rbnz.govt.nz" in url:
        raise MarketStateError(
            f"RBNZ official B2 workbook returned HTTP 403 from {url}. "
            "Refusing to substitute a third-party or vendor feed."
        ) from last
    raise MarketStateError(f"failed to fetch {url}: {last}")


def parse_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y", "%d-%b-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    s = str(value).strip().replace("%", "").replace(",", "")
    if not s or s in {"..", "-", "—", "n/a", "NA", "na"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def filter_since(series: Mapping[date, float], start: date) -> dict[date, float]:
    return {d: v for d, v in series.items() if d >= start}


def percentile_rank(values: Sequence[float], x: float) -> float | None:
    if not values:
        return None
    lt = sum(v < x for v in values)
    eq = sum(v == x for v in values)
    return 100.0 * (lt + 0.5 * eq) / len(values)


def zscore(values: Sequence[float], x: float) -> float | None:
    if len(values) < 2:
        return None
    sd = statistics.pstdev(values)
    if sd == 0:
        return 0.0
    return (x - statistics.fmean(values)) / sd


def lookback_value(items: Sequence[tuple[date, float]], n_obs: int) -> float | None:
    if len(items) <= n_obs:
        return None
    return items[-1 - n_obs][1]


def _move_analogs(
    series: Mapping[date, float],
    *,
    lookback: int = 21,
    scale: float = 100.0,
    forward_1m: int = 21,
    forward_3m: int = 63,
    count: int = 4,
) -> list[dict]:
    """Closest prior 1-month moves plus subsequent outcomes.

    The dates/math are deterministic. A trader must still explain why an
    episode is economically comparable and which regime differences matter.
    """
    items = sorted(series.items())
    if len(items) <= lookback + forward_3m + 20:
        return []
    current_move = (items[-1][1] - items[-1 - lookback][1]) * scale
    candidates: list[tuple[float, int, dict]] = []
    last_candidate = len(items) - 1 - forward_3m
    for i in range(lookback, last_candidate + 1):
        prior_move = (items[i][1] - items[i - lookback][1]) * scale
        row = {
            "as_of": items[i][0].isoformat(),
            "level": items[i][1],
            "prior_1m_move": round(prior_move, 4),
            "forward_1m_move": round((items[i + forward_1m][1] - items[i][1]) * scale, 4),
            "forward_3m_move": round((items[i + forward_3m][1] - items[i][1]) * scale, 4),
            "distance_from_current_move": round(abs(prior_move - current_move), 4),
        }
        candidates.append((abs(prior_move - current_move), i, row))
    selected: list[dict] = []
    selected_i: list[int] = []
    for _distance, idx, row in sorted(candidates, key=lambda x: x[0]):
        if any(abs(idx - old) < 21 for old in selected_i):
            continue
        selected.append(row)
        selected_i.append(idx)
        if len(selected) >= count:
            break
    return selected


def rate_metrics(series: Mapping[date, float]) -> dict:
    items = sorted(series.items())
    if not items:
        raise MarketStateError("empty rate series")
    latest_date, latest = items[-1]
    vals_1y = [v for _, v in items[-252:]]
    vals_5y = [v for _, v in items[-1260:]]

    def bp(n: int) -> float | None:
        old = lookback_value(items, n)
        return None if old is None else (latest - old) * 100.0

    return {
        "as_of": latest_date.isoformat(),
        "value": latest,
        "bp_1d": bp(1),
        "bp_5d": bp(5),
        "bp_1m": bp(21),
        "bp_3m": bp(63),
        "bp_6m": bp(126),
        "bp_1y": bp(252),
        "historical_move_analogs": _move_analogs(series, scale=100.0),
        "pctile_1y": percentile_rank(vals_1y, latest),
        "pctile_5y": percentile_rank(vals_5y, latest) if len(vals_5y) >= 60 else None,
        "z_1y": zscore(vals_1y, latest),
        "z_5y": zscore(vals_5y, latest) if len(vals_5y) >= 60 else None,
    }


def fx_metrics(series: Mapping[date, float]) -> dict:
    items = sorted(series.items())
    if not items:
        raise MarketStateError("empty FX series")
    latest_date, latest = items[-1]
    vals = [v for _, v in items]
    vals_1y = vals[-252:]
    vals_5y = vals[-1260:]

    def ret(n: int) -> float | None:
        old = lookback_value(items, n)
        return None if old in (None, 0) else 100.0 * (latest / old - 1.0)

    logrets = [math.log(vals[i] / vals[i - 1]) for i in range(1, len(vals)) if vals[i] > 0 and vals[i - 1] > 0]

    def rv(n: int) -> float | None:
        if len(logrets) < n:
            return None
        return statistics.pstdev(logrets[-n:]) * math.sqrt(252.0) * 100.0

    return {
        "as_of": latest_date.isoformat(),
        "spot": latest,
        "ret_1d": ret(1),
        "ret_5d": ret(5),
        "ret_1m": ret(21),
        "ret_3m": ret(63),
        "rv_20d": rv(20),
        "rv_60d": rv(60),
        "pctile_1y": percentile_rank(vals_1y, latest),
        "pctile_5y": percentile_rank(vals_5y, latest) if len(vals_5y) >= 60 else None,
        "z_1y": zscore(vals_1y, latest),
        "z_5y": zscore(vals_5y, latest) if len(vals_5y) >= 60 else None,
    }


def parse_treasury_csv(text: str) -> dict[str, dict[date, float]]:
    out = {"2Y": {}, "5Y": {}, "10Y": {}, "30Y": {}}
    rows = csv.DictReader(io.StringIO(text))
    keymap = {"2Y": "2 Yr", "5Y": "5 Yr", "10Y": "10 Yr", "30Y": "30 Yr"}
    for row in rows:
        d = parse_date(row.get("Date"))
        if not d:
            continue
        for tenor, col in keymap.items():
            v = to_float(row.get(col))
            if v is not None:
                out[tenor][d] = v
    if any(not out[k] for k in ("2Y", "5Y", "10Y", "30Y")):
        raise MarketStateError("Treasury CSV missing required tenors")
    return out


def fetch_us_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    merged = {"2Y": {}, "5Y": {}, "10Y": {}, "30Y": {}}
    for year in range(start.year, end.year + 1):
        text = fetch_bytes(TREASURY_URL.format(year=year)).decode("utf-8-sig")
        parsed = parse_treasury_csv(text)
        for tenor, series in parsed.items():
            merged[tenor].update(series)
    out = {k: filter_since(v, start) for k, v in merged.items()}
    if any(not out[k] for k in ("2Y", "5Y", "10Y", "30Y")):
        raise MarketStateError("Treasury history missing required tenors after date filter")
    return out


def parse_boc_json(payload: Mapping, *, require_all: bool = True) -> dict[str, dict[date, float]]:
    out = {v: {} for v in BOC_SERIES.values()}
    for obs in payload.get("observations", []):
        d = parse_date(obs.get("d"))
        if not d:
            continue
        for sid, tenor in BOC_SERIES.items():
            cell = obs.get(sid)
            value = cell.get("v") if isinstance(cell, Mapping) else cell
            v = to_float(value)
            if v is not None:
                out[tenor][d] = v
    if require_all and (any(not out[k] for k in RATE_TENORS) or not out["LONG"]):
        raise MarketStateError("BoC response missing required benchmark tenors")
    return out


def fetch_ca_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    """Fetch current BoC benchmark bond yields.

    The legacy V39051/V39053/V39055/V39056 series IDs now 404. Use the official
    `bond_yields_benchmark` group, then fail closed if a required tenor is absent.
    """
    qs = urllib.parse.urlencode({"start_date": start.isoformat(), "end_date": end.isoformat()})
    payload = json.loads(fetch_bytes(f"{BOC_GROUP_URL}?{qs}").decode("utf-8"))
    parsed = parse_boc_json(payload)
    missing = [k for k in (*RATE_TENORS, "LONG") if not parsed.get(k)]
    if not missing:
        return parsed
    for series_id, tenor in BOC_SERIES.items():
        if parsed.get(tenor):
            continue
        series_payload = json.loads(fetch_bytes(f"{BOC_SERIES_URL.format(series=series_id)}?{qs}").decode("utf-8"))
        extra = parse_boc_json(series_payload, require_all=False)
        parsed[tenor].update(extra.get(tenor, {}))
    if any(not parsed[k] for k in RATE_TENORS) or not parsed["LONG"]:
        raise MarketStateError("BoC benchmark group missing required tenors")
    return parsed


def parse_rba_csv(text: str) -> dict[str, dict[date, float]]:
    rows = list(csv.reader(io.StringIO(text)))
    metadata: dict[str, list[str]] = {}
    series_row_idx = None
    for i, row in enumerate(rows):
        if not row:
            continue
        first = row[0].strip().lower()
        if first in {"title", "description", "frequency", "type", "units", "source", "publication date", "series id"}:
            metadata[first] = row
        if first == "series id":
            series_row_idx = i
            break
    if series_row_idx is None:
        raise MarketStateError("RBA F2 CSV missing Series ID row")
    selected: dict[int, str] = {}
    for key in ("title", "description"):
        labels = metadata.get(key)
        if not labels:
            continue
        candidate: dict[int, str] = {}
        for idx in range(1, len(labels)):
            label = labels[idx].strip().lower() if idx < len(labels) else ""
            if "indexed" in label:
                continue
            if "australian government" not in label and "government" not in label:
                continue
            for n, tenor in ((2, "2Y"), (5, "5Y"), (10, "10Y")):
                if re.search(rf"\b{n}\s*[- ]?years?\b", label):
                    candidate[idx] = tenor
        if set(candidate.values()) == set(RATE_TENORS):
            selected = candidate
            break
    if set(selected.values()) != set(RATE_TENORS):
        raise MarketStateError(f"RBA F2 tenor discovery failed: {selected}")
    out = {k: {} for k in RATE_TENORS}
    for row in rows[series_row_idx + 1 :]:
        if not row:
            continue
        d = parse_date(row[0])
        if not d:
            continue
        for idx, tenor in selected.items():
            if idx < len(row):
                v = to_float(row[idx])
                if v is not None:
                    out[tenor][d] = v
    if any(not out[k] for k in RATE_TENORS):
        raise MarketStateError("RBA F2 data rows missing required tenors")
    return out


def fetch_au_rates(start: date, end: date) -> dict[str, dict[date, float]]:
    parsed = parse_rba_csv(fetch_bytes(RBA_URL).decode("utf-8-sig", errors="replace"))
    out = {k: {d: v for d, v in s.items() if start <= d <= end} for k, s in parsed.items()}
    if any(not out[k] for k in RATE_TENORS):
        raise MarketStateError("RBA F2 history missing required tenors after date filter")
    return out


def parse_rbnz_xlsx(data: bytes) -> dict[str, dict[date, float]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise MarketStateError("openpyxl is required for the RBNZ workbook") from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        preview = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 40, 40), values_only=True))
        for ri, row in enumerate(preview):
            if not row:
                continue
            date_cols = [i for i, v in enumerate(row) if str(v).strip().lower() == "date"]
            if not date_cols:
                continue
            header_start = max(0, ri - 5)
            max_cols = max(len(r) for r in preview[header_start : ri + 1])
            header_rows = preview[header_start : ri + 1]
            filled_rows = []
            for rr in header_rows:
                filled = []
                carry = None
                for ci in range(max_cols):
                    val = rr[ci] if ci < len(rr) else None
                    if val not in (None, ""):
                        carry = val
                    filled.append(carry)
                filled_rows.append(filled)
            col_headers: dict[int, str] = {}
            for ci in range(max_cols):
                parts = []
                for rr in filled_rows:
                    if rr[ci] not in (None, ""):
                        parts.append(str(rr[ci]).strip())
                col_headers[ci] = " | ".join(parts).lower()
            selected: dict[int, str] = {}
            for ci, header in col_headers.items():
                if "government" not in header and "bond" not in header:
                    continue
                for n, tenor in ((2, "2Y"), (5, "5Y"), (10, "10Y")):
                    if re.search(rf"\b{n}\s*year\b", header):
                        selected[ci] = tenor
            if set(selected.values()) == set(RATE_TENORS):
                best = (ws, ri + 1, selected)
                break
        if best:
            break
    if not best:
        raise MarketStateError("RBNZ workbook tenor discovery failed")
    ws, header_row, selected = best
    out = {k: {} for k in RATE_TENORS}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row:
            continue
        d = None
        for v in row[:3]:
            d = parse_date(v)
            if d:
                break
        if not d:
            continue
        for ci, tenor in selected.items():
            if ci < len(row):
                v = to_float(row[ci])
                if v is not None:
                    out[tenor][d] = v
    if any(not out[k] for k in RATE_TENORS):
        raise MarketStateError("RBNZ workbook missing required government tenors")
    return out


def fetch_nz_rates(
    start: date,
    end: date,
    *,
    workbook_bytes: bytes | None = None,
) -> dict[str, dict[date, float]]:
    if workbook_bytes is None:
        try:
            workbook_bytes = fetch_bytes(RBNZ_URL, user_agent=BROWSER_USER_AGENT)
        except MarketStateError:
            workbook_bytes = fetch_bytes(RBNZ_URL)
    if workbook_bytes[:2] != b"PK":
        raise MarketStateError("RBNZ download was not an xlsx workbook; refusing to parse a substitute page")
    parsed = parse_rbnz_xlsx(workbook_bytes)
    out = {k: {d: v for d, v in s.items() if start <= d <= end} for k, s in parsed.items()}
    if any(not out[k] for k in RATE_TENORS):
        raise MarketStateError("RBNZ history missing required tenors after date filter")
    return out


def parse_ecb_sdmx_csv(text: str) -> dict[str, dict[date, float]]:
    """Parse ECB Data Portal SDMX-CSV daily EUR reference rates.

    Completeness is checked by the caller so a combined query and a
    per-currency fallback can share the same parser.
    """
    reader = csv.DictReader(io.StringIO(text))
    currency = {c: {} for c in G10}
    needed = {c for c in G10 if c != "EUR"}
    for row in reader:
        ccy = str(row.get("CURRENCY") or "").strip()
        if ccy not in needed:
            continue
        d = parse_date(row.get("TIME_PERIOD"))
        v = to_float(row.get("OBS_VALUE"))
        if d and v is not None and v > 0:
            currency[ccy][d] = v
            currency["EUR"][d] = 1.0
    if not any(currency[c] for c in needed):
        raise MarketStateError("ECB SDMX FX contained no G10 observations")
    return currency


def require_g10_fx(currency: Mapping[str, Mapping[date, float]]) -> None:
    missing = [c for c in G10 if not currency.get(c)]
    if missing:
        raise MarketStateError(f"ECB SDMX FX missing G10 currencies: {missing}")


def parse_ecb_fx_csv(text: str) -> dict[str, dict[date, float]]:
    reader = csv.DictReader(io.StringIO(text))
    fieldnames = [x.strip() for x in (reader.fieldnames or [])]
    missing = [c for c in G10 if c != "EUR" and c not in fieldnames]
    if missing:
        raise MarketStateError(f"ECB FX CSV missing G10 currencies: {missing}")
    currency = {c: {} for c in G10}
    for row in reader:
        clean = {str(k).strip(): v for k, v in row.items() if k is not None}
        d = parse_date(clean.get("Date"))
        if not d:
            continue
        currency["EUR"][d] = 1.0
        for c in G10:
            if c == "EUR":
                continue
            v = to_float(clean.get(c))
            if v is not None and v > 0:
                currency[c][d] = v
    if any(not currency[c] for c in G10):
        raise MarketStateError("ECB FX history missing one or more G10 currency histories")
    return currency


def build_fx_crosses(currency_per_eur: Mapping[str, Mapping[date, float]], start: date) -> dict[str, dict[date, float]]:
    """Derive all 45 unique G10 crosses from the same ECB EUR fixing date.

    ECB publishes foreign-currency units per EUR. BASEQUOTE = QUOTE_per_EUR / BASE_per_EUR.
    Legs from different fixing dates are never mixed.
    """
    out: dict[str, dict[date, float]] = {}
    for i, base in enumerate(FX_ORDER):
        for quote in FX_ORDER[i + 1 :]:
            common = sorted(set(currency_per_eur[base]) & set(currency_per_eur[quote]))
            series = {
                d: currency_per_eur[quote][d] / currency_per_eur[base][d]
                for d in common
                if d >= start and currency_per_eur[base][d] > 0
            }
            out[f"{base}{quote}"] = series
    if len(out) != 45 or any(not s for s in out.values()):
        raise MarketStateError(f"expected 45 populated G10 crosses, got {len(out)}")
    return out


def _fetch_ecb_sdmx(
    currencies: str,
    start: date,
    *,
    timeout: int = ECB_TIMEOUT,
    retries: int = ECB_RETRIES,
) -> dict[str, dict[date, float]]:
    qs = urllib.parse.urlencode({"format": "csvdata", "startPeriod": start.isoformat()})
    url = f"{ECB_FX_SDMX_URL.format(currencies=currencies)}?{qs}"
    return parse_ecb_sdmx_csv(fetch_bytes(url, timeout=timeout, retries=retries).decode("utf-8-sig"))


def _merge_fx_legs(
    dest: dict[str, dict[date, float]], extra: Mapping[str, Mapping[date, float]]
) -> None:
    for key, series in extra.items():
        dest.setdefault(key, {}).update(series)


def fetch_fx(start: date) -> dict[str, dict[date, float]]:
    needed = [c for c in G10 if c != "EUR"]
    try:
        # Combined 5-year query is one round-trip but ECB sometimes 504s it.
        parsed = _fetch_ecb_sdmx("+".join(needed), start, timeout=60, retries=2)
        require_g10_fx(parsed)
    except MarketStateError as exc:
        parsed = {c: {} for c in G10}
        errors = [str(exc)]
        for ccy in needed:
            try:
                extra = _fetch_ecb_sdmx(ccy, start, timeout=ECB_TIMEOUT, retries=ECB_RETRIES)
            except MarketStateError as inner:
                errors.append(f"{ccy}: {inner}")
                continue
            _merge_fx_legs(parsed, extra)
        try:
            require_g10_fx(parsed)
        except MarketStateError:
            raise MarketStateError(
                "ECB SDMX combined query failed and per-currency fallback still missing "
                f"{[c for c in G10 if not parsed[c]]}: {errors[-1]}"
            ) from exc
    latest = max(max(series) for series in parsed.values() if series)
    if (datetime.now(timezone.utc).date() - latest).days > FAIL_AFTER_DAYS:
        raise MarketStateError(f"ECB SDMX latest observation {latest.isoformat()} is too old")
    usd = parsed["USD"]
    latest_usd = usd[max(usd)]
    if not (0.5 < latest_usd < 2.5):
        raise MarketStateError(f"implausible ECB EURUSD reference rate: {latest_usd}")
    return build_fx_crosses(parsed, start)


def curve_spread(short: Mapping[date, float], long: Mapping[date, float]) -> dict[date, float]:
    return {d: (long[d] - short[d]) * 100.0 for d in set(short) & set(long)}


def rv_spread(a: Mapping[date, float], b: Mapping[date, float]) -> dict[date, float]:
    """Matching-tenor spread in basis points. Common observation dates only; no forward-fill."""
    return {d: (a[d] - b[d]) * 100.0 for d in set(a) & set(b)}


def spread_metrics_bps(series: Mapping[date, float]) -> dict:
    items = sorted(series.items())
    if not items:
        raise MarketStateError("empty spread series")
    latest_date, latest = items[-1]
    vals_1y = [v for _, v in items[-252:]]
    vals_5y = [v for _, v in items[-1260:]]

    def chg(n: int) -> float | None:
        old = lookback_value(items, n)
        return None if old is None else latest - old

    return {
        "as_of": latest_date.isoformat(),
        "bps": latest,
        "chg_1d_bps": chg(1),
        "chg_5d_bps": chg(5),
        "chg_1m_bps": chg(21),
        "chg_3m_bps": chg(63),
        "chg_6m_bps": chg(126),
        "chg_1y_bps": chg(252),
        "historical_move_analogs": _move_analogs(series, scale=1.0),
        "pctile_1y": percentile_rank(vals_1y, latest),
        "pctile_5y": percentile_rank(vals_5y, latest) if len(vals_5y) >= 60 else None,
        "z_1y": zscore(vals_1y, latest),
        "z_5y": zscore(vals_5y, latest) if len(vals_5y) >= 60 else None,
    }


def _freshness(age_days: int, key: str) -> str:
    if age_days > FAIL_AFTER_DAYS:
        raise MarketStateError(
            f"{key} observation is {age_days} days old; refusing to emit a fabricated or silently substituted substitute"
        )
    return "stale" if age_days > STALE_AFTER_DAYS[key] else "ok"


def _empty_tenor_metrics() -> dict:
    return {
        "as_of": None,
        "value": None,
        "bp_1d": None,
        "bp_5d": None,
        "bp_1m": None,
        "bp_3m": None,
        "pctile_1y": None,
        "pctile_5y": None,
        "z_1y": None,
        "z_5y": None,
    }


def _unavailable_rate_rv(reason: str) -> dict:
    return {
        "as_of": None,
        "bps": None,
        "chg_1d_bps": None,
        "chg_5d_bps": None,
        "chg_1m_bps": None,
        "chg_3m_bps": None,
        "pctile_1y": None,
        "pctile_5y": None,
        "z_1y": None,
        "z_5y": None,
        "status": "unavailable",
        "reason": reason,
    }


def _nz_rates_unavailable(error: str) -> tuple[dict[str, dict[date, float]], dict]:
    """Return empty NZ raw series and a rates block with explicit provenance."""
    empty = {t: {} for t in RATE_TENORS}
    block = {
        "tenors": {t: _empty_tenor_metrics() for t in RATE_TENORS},
        "curves": {},
        "latest_observation": None,
        "age_days": None,
        "status": "unavailable",
        "error": error,
    }
    return empty, block


def build_snapshot(
    *,
    today: date | None = None,
    nz_workbook: Path | None = None,
    include_cross_assets: bool = True,
    include_positioning: bool | None = None,
    include_policy_paths: bool | None = None,
    include_official_curves: bool | None = None,
    include_australia_housing: bool | None = None,
) -> dict:
    today = today or datetime.now(timezone.utc).date()
    start = today - timedelta(days=366 * 5 + 15)
    nz_bytes = nz_workbook.read_bytes() if nz_workbook else None
    nz_error: str | None = None
    try:
        nz_raw = fetch_nz_rates(start, today, workbook_bytes=nz_bytes)
    except MarketStateError as exc:
        nz_error = str(exc)
        nz_raw, nz_block = _nz_rates_unavailable(nz_error)
    else:
        nz_block = None

    rates_raw = {
        "US": fetch_us_rates(start, today),
        "CA": fetch_ca_rates(start, today),
        "AU": fetch_au_rates(start, today),
        "NZ": nz_raw,
    }
    rates: dict[str, dict] = {}
    stale_sources: list[str] = []
    unavailable_sources: list[str] = []
    if nz_error:
        unavailable_sources.append("NZ_rates")
        stale_sources.append("NZ_rates")
    for country, tenors in rates_raw.items():
        if country == "NZ" and nz_block is not None:
            rates[country] = nz_block
            continue
        required = ("2Y", "5Y", "10Y", "30Y") if country == "US" else ("2Y", "5Y", "10Y", "LONG") if country == "CA" else RATE_TENORS
        missing = [t for t in required if t not in tenors or not tenors[t]]
        if missing:
            raise MarketStateError(f"{country} missing required tenors: {missing}")
        metrics = {tenor: rate_metrics(series) for tenor, series in tenors.items() if series}
        curves = {
            "2s10s": spread_metrics_bps(curve_spread(tenors["2Y"], tenors["10Y"])),
            "5s10s": spread_metrics_bps(curve_spread(tenors["5Y"], tenors["10Y"])),
        }
        latest_as_of = max(date.fromisoformat(metrics[t]["as_of"]) for t in RATE_TENORS)
        age_days = max(0, (today - latest_as_of).days)
        status = _freshness(age_days, country)
        if status == "stale":
            stale_sources.append(f"{country}_rates")
        rates[country] = {
            "tenors": metrics,
            "curves": curves,
            "latest_observation": latest_as_of.isoformat(),
            "age_days": age_days,
            "status": status,
        }

    rate_rv: dict[str, dict] = {}
    nz_available = nz_error is None
    for a, b in RV_PAIRS:
        for tenor in RATE_TENORS:
            if not nz_available and ("NZ" in (a, b)):
                rate_rv[f"{a}-{b}_{tenor}"] = _unavailable_rate_rv(nz_error or "NZ rates unavailable")
                continue
            common = rv_spread(rates_raw[a][tenor], rates_raw[b][tenor])
            if not common:
                raise MarketStateError(f"no overlapping {a}/{b} {tenor} observation dates; refusing to forward-fill")
            rate_rv[f"{a}-{b}_{tenor}"] = spread_metrics_bps(common)

    fx_raw = fetch_fx(start)
    fx = {pair: fx_metrics(series) for pair, series in fx_raw.items()}
    latest_fx_date = max(date.fromisoformat(v["as_of"]) for v in fx.values())
    fx_age = max(0, (today - latest_fx_date).days)
    fx_status = _freshness(fx_age, "FX")
    if fx_status == "stale":
        stale_sources.append("FX")

    if include_australia_housing is None:
        include_australia_housing = include_cross_assets
    australia_housing = (
        collect_australia_housing(today=today, fetch_bytes=fetch_bytes)
        if include_australia_housing
        else {
            "country": "AU",
            "status": "unavailable",
            "feeds": {
                name: {
                    "status": "unavailable",
                    "url": "disabled",
                    "as_of": None,
                    "error": "Australian housing collection disabled for this invocation",
                }
                for name in (
                    "prices_and_turnover",
                    "building_approvals",
                    "housing_lending",
                    "housing_credit",
                    "mortgage_rates",
                    "mortgage_cash_flow",
                )
            },
            "method": {
                "model_calls": 0,
                "credentials_required": [],
                "purpose": "Official Australian housing transmission data for Market Watch and Trader Room context.",
                "trader_packet_delivery": "market_state_passthrough",
            },
        }
    )

    cross_raw, cross_meta = collect_cross_assets(start, today, fetch_bytes) if include_cross_assets else ({}, {})
    opportunities = build_opportunities(rates_raw, fx_raw, cross_raw, cross_meta, today)
    if include_policy_paths is None:
        include_policy_paths = include_cross_assets
    policy_paths = (
        collect_policy_paths(today=today, fetch_bytes=fetch_bytes)
        if include_policy_paths
        else {
            "status": "unavailable",
            "countries": {
                c: {"status": "unavailable", "error": "policy path collection disabled for this invocation"}
                for c in ("US", "CA", "AU")
            },
            "sources": {},
            "method": {
                "model_calls": 0,
                "credentials_required": [],
                "purpose": "policy expectations context before sovereign-curve/RV trade construction",
            },
        }
    )
    tradable_rate_curves = build_tradable_rate_curves(policy_paths)
    if include_policy_paths and tradable_rate_curves.get("status") != "ok":
        for curve_id, block in (tradable_rate_curves.get("curves") or {}).items():
            if (block or {}).get("status") != "ok":
                unavailable_sources.append(f"{curve_id}_tradable_curve")

    if include_official_curves is None:
        include_official_curves = include_cross_assets
    official_curves = (
        collect_official_curves(today=today, fetch_bytes=fetch_bytes)
        if include_official_curves
        else {
            "status": "unavailable",
            "countries": {
                c: {"status": "unavailable", "error": "official curve collection disabled for this invocation"}
                for c in ("US", "CA", "AU")
            },
            "errors": {},
            "method": {
                "model_calls": 0,
                "credentials_required": [],
                "purpose": "official government zero/forward curve proxy for deterministic paper fwd-fwd and curve marks",
                "ois_equivalence": "proxy_only",
            },
        }
    )
    if include_official_curves and official_curves.get("status") != "ok":
        for country in ("US", "CA", "AU"):
            if (official_curves.get("countries", {}).get(country) or {}).get("status") != "ok":
                unavailable_sources.append(f"{country}_official_curve")
    if include_positioning is None:
        include_positioning = include_cross_assets
    positioning_start = today - timedelta(days=366 * 3 + 30)
    positioning = (
        build_positioning(today=today, start=positioning_start, fetch_bytes=fetch_bytes)
        if include_positioning
        else {
            "status": "unavailable",
            "error": "positioning collection disabled for this invocation",
            "cftc_tff": {"status": "unavailable", "error": "positioning collection disabled for this invocation", "instruments": {}},
            "cme": {"status": "unavailable", "error": "positioning collection disabled for this invocation", "instruments": {}},
            "method": {"model_calls": 0, "credentials_required": []},
        }
    )
    if include_positioning:
        if positioning["cftc_tff"].get("status") == "unavailable":
            unavailable_sources.append("CFTC_positioning")
        if positioning["cme"].get("status") == "unavailable":
            unavailable_sources.append("CME_positioning")
        if positioning["cftc_tff"].get("status") == "stale":
            stale_sources.append("CFTC_positioning")
        if positioning["cme"].get("status") == "stale":
            stale_sources.append("CME_positioning")
    packet_status = "stale" if stale_sources else "ok"
    nz_source = {
        "name": "Reserve Bank of New Zealand B2 wholesale interest rates",
        "url": RBNZ_PAGE,
        "download_url": RBNZ_URL,
        "observation_date": rates["NZ"]["latest_observation"],
        "status": rates["NZ"]["status"],
        "note": "Indicative closing government-bond yields with a one-day publication lag.",
    }
    if nz_error:
        nz_source["error"] = nz_error
        nz_source["observation_date"] = None
    return {
        "policy_paths": policy_paths,
        "tradable_rate_curves": tradable_rate_curves,
        "official_curves": official_curves,
        "australia_housing": australia_housing,
        "cross_assets": {"series": cross_meta, "status": "partial" if any(m["status"] != "ok" for m in cross_meta.values()) else "ok"},
        "opportunities": opportunities,
        "positioning": positioning,
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "window_start": start.isoformat(),
        "status": packet_status,
        "stale_sources": stale_sources,
        "unavailable_sources": unavailable_sources,
        "rates": rates,
        "rate_rv": rate_rv,
        "fx": {
            "source_observation": latest_fx_date.isoformat(),
            "age_days": fx_age,
            "status": fx_status,
            "pairs": fx,
        },
        "sources": {
            "US_rates": {
                "name": "U.S. Treasury daily par yield curve",
                "url": TREASURY_PAGE,
                "download_url": TREASURY_URL.format(year=today.year),
                "observation_date": rates["US"]["latest_observation"],
                "status": rates["US"]["status"],
            },
            "CA_rates": {
                "name": "Bank of Canada Valet benchmark bonds",
                "url": BOC_DOCS,
                "download_url": BOC_GROUP_URL,
                "observation_date": rates["CA"]["latest_observation"],
                "status": rates["CA"]["status"],
            },
            "AU_rates": {
                "name": "Reserve Bank of Australia F2 government-bond yields",
                "url": RBA_PAGE,
                "download_url": RBA_URL,
                "observation_date": rates["AU"]["latest_observation"],
                "status": rates["AU"]["status"],
                "note": "RBA assessed closing yields; research context only; not a financial benchmark; typically weekly with a two-business-day lag.",
            },
            "NZ_rates": nz_source,
            "CFTC_positioning": {
                "name": "CFTC Traders in Financial Futures - Futures Only",
                "url": positioning["cftc_tff"].get("source_url", "https://publicreporting.cftc.gov/stories/s/TFF-Futures-Only/98ig-3k9y/"),
                "download_url": positioning["cftc_tff"].get("api_url", "https://publicreporting.cftc.gov/resource/gpe5-46if.json"),
                "observation_date": positioning["cftc_tff"].get("report_date"),
                "status": positioning["cftc_tff"].get("status"),
                **({"error": positioning["cftc_tff"].get("error")} if positioning["cftc_tff"].get("status") == "unavailable" else {}),
            },
            "CME_positioning": {
                "name": "CME Group FX futures and options volume/open-interest service",
                "url": positioning["cme"].get("source_url", "https://www.cmegroup.com/market-data/volume-open-interest.html"),
                "download_url": positioning["cme"].get("api_template", "https://www.cmegroup.com/CmeWS/mvc/Volume/LastTotals/{product_id}?days=30"),
                "observation_date": positioning["cme"].get("trade_date"),
                "status": positioning["cme"].get("status"),
                **({"error": positioning["cme"].get("error")} if positioning["cme"].get("status") == "unavailable" else {}),
            },
            "FX": {
                "name": "ECB euro foreign exchange reference rates",
                "url": ECB_FX_PAGE,
                "download_url": ECB_FX_SDMX_URL.format(
                    currencies="+".join(c for c in G10 if c != "EUR")
                ),
                "observation_date": latest_fx_date.isoformat(),
                "status": fx_status,
                "note": "Information/reference rates, not executable prices. All 45 G10 crosses are deterministic ratios of the same ECB EUR fixing date.",
            },
        },
        "method": {
            "model_calls": 0,
            "fx_source": "ecb_euro_reference_crosses",
            "fx_cross_construction": True,
            "fx_pair_count": len(fx),
            "rate_rv_count": len(rate_rv),
            "credentials_required": [],
            "note": (
                "Live authorities remain canonical. This artifact is a compact research snapshot. "
                "US, Canada, Australia and ECB FX are required; a blocked official RBNZ source "
                "marks NZ rates and NZ-dependent RV spreads unavailable without fabricating data. "
                "Cross-country spreads use exact common observation dates only. "
                "Policy-path context uses official overnight benchmarks plus public money-market data. "
                "The tradable paper rates universe is explicitly SOFR via CME SR3, CORRA via MX CRA, and AONIA via ASX IB; a position stays on its entry curve family until close. "
                "Official government zero/forward curves are supplemental bond-curve inputs and are not required to manufacture a swap curve. "
                "Australian housing context is collected from ABS dwelling prices/transfers, building approvals and lending indicators plus RBA housing credit, mortgage rates and household housing-loan payment data. "
                "CFTC TFF supplies trader-class ownership/crowding context and CME's public volume/open-interest service supplies daily FX futures and aggregate options OI history. "
                "No historical warehouse is written to GitHub or Supabase."
            ),
        },
    }


def validate_snapshot(s: Mapping) -> None:
    if s.get("method", {}).get("model_calls") != 0:
        raise MarketStateError("generator must make zero model calls")
    if s.get("method", {}).get("credentials_required") not in ([], None):
        raise MarketStateError("MVP path must not require credentials")
    if s.get("method", {}).get("fx_source") != "ecb_euro_reference_crosses":
        raise MarketStateError("FX source must be same-fixing ECB EUR reference crosses")
    if s.get("method", {}).get("fx_pair_count") != 45:
        raise MarketStateError("generator must produce exactly 45 G10 FX crosses")
    if s.get("method", {}).get("rate_rv_count") != len(RV_PAIRS) * len(RATE_TENORS):
        raise MarketStateError("unexpected rate RV count")
    try:
        validate_positioning(s.get("positioning") or {})
    except Exception as exc:
        raise MarketStateError(f"invalid positioning block: {exc}") from exc
    try:
        validate_policy_paths(s.get("policy_paths") or {})
    except Exception as exc:
        raise MarketStateError(f"invalid policy_paths block: {exc}") from exc
    try:
        validate_tradable_rate_curves(s.get("tradable_rate_curves") or {})
    except Exception as exc:
        raise MarketStateError(f"invalid tradable_rate_curves block: {exc}") from exc
    try:
        validate_official_curves(s.get("official_curves") or {})
    except Exception as exc:
        raise MarketStateError(f"invalid official_curves block: {exc}") from exc
    try:
        validate_australia_housing(s.get("australia_housing") or {})
    except Exception as exc:
        raise MarketStateError(f"invalid australia_housing block: {exc}") from exc
    if set(s.get("rates", {})) != set(RATE_COUNTRIES):
        raise MarketStateError("rates block must contain US, CA, AU and NZ")
    for c in RATE_COUNTRIES:
        block = s["rates"][c]
        if block.get("status") == "unavailable":
            if c != "NZ":
                raise MarketStateError(f"only NZ may be unavailable, not {c}")
            if not block.get("error"):
                raise MarketStateError("NZ unavailable block must include error provenance")
            continue
        for t in RATE_TENORS:
            value = block["tenors"][t]["value"]
            if not (0 < value < 25):
                raise MarketStateError(f"implausible {c} {t} yield: {value}")
            if not block["tenors"][t].get("as_of"):
                raise MarketStateError(f"{c} {t} missing observation date")
        extra = "30Y" if c == "US" else "LONG" if c == "CA" else None
        if extra and extra not in block["tenors"]:
            raise MarketStateError(f"{c} missing {extra} tenor")
    for _key, rv in s.get("rate_rv", {}).items():
        if rv.get("status") == "unavailable":
            if rv.get("bps") is not None:
                raise MarketStateError("unavailable RV spread must not include fabricated bps")
            continue
        if rv.get("bps") is None:
            raise MarketStateError("available RV spread missing bps")
    if len(s["fx"]["pairs"]) != 45:
        raise MarketStateError("FX pairs must be exactly 45")
    for required in ("EURUSD", "USDJPY", "AUDNZD", "NOKSEK", "USDCAD"):
        if required not in s["fx"]["pairs"]:
            raise MarketStateError(f"missing required FX pair {required}")
    for pair, m in s["fx"]["pairs"].items():
        if not (0 < m["spot"] < 1000):
            raise MarketStateError(f"implausible FX spot {pair}: {m['spot']}")
        if not m.get("as_of"):
            raise MarketStateError(f"{pair} missing observation date")
    for key, src in s.get("sources", {}).items():
        if not src.get("url"):
            raise MarketStateError(f"source {key} missing url")
        if src.get("status") == "unavailable":
            if not src.get("error"):
                raise MarketStateError(f"source {key} unavailable without provenance")
            continue
        if not src.get("observation_date"):
            raise MarketStateError(f"source {key} missing observation_date")
    if not s.get("generated_at"):
        raise MarketStateError("generated_at missing")


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Generate the Market Watch daily market-state packet")
    ap.add_argument("--output", required=True, help="JSON snapshot path")
    ap.add_argument("--today", help="override current date (YYYY-MM-DD), useful for testing")
    ap.add_argument(
        "--nz-workbook",
        help="optional local official RBNZ B2 xlsx (same file as the public download URL)",
    )
    ap.add_argument("--public", action="store_true", help="Omit restricted index observations from the public read model")
    args = ap.parse_args(argv)
    today = date.fromisoformat(args.today) if args.today else None
    nz_workbook = Path(args.nz_workbook) if args.nz_workbook else None
    if nz_workbook and not nz_workbook.is_file():
        raise MarketStateError(f"RBNZ workbook not found: {nz_workbook}")
    snapshot = build_snapshot(today=today, nz_workbook=nz_workbook)
    validate_snapshot(snapshot)
    if args.public:
        from scripts.cross_asset_data import public_snapshot
        snapshot = public_snapshot(snapshot)
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(snapshot, indent=2, sort_keys=True) + "\n")
    print(
        json.dumps(
            {
                "status": snapshot["status"],
                "fx_pairs": snapshot["method"]["fx_pair_count"],
                "rate_rv": snapshot["method"]["rate_rv_count"],
                "fx_as_of": snapshot["fx"]["source_observation"],
                "rate_as_of": {c: snapshot["rates"][c]["latest_observation"] for c in RATE_COUNTRIES},
                "stale_sources": snapshot["stale_sources"],
                "positioning_status": snapshot["positioning"]["status"],
                "output": str(out),
            }
        )
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except MarketStateError as exc:
        print(f"MARKET STATE ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
