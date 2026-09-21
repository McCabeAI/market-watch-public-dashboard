"""Official Australian housing data for Market Watch."""

from __future__ import annotations

import csv
import html
import io
import re
from calendar import monthrange
from datetime import date, datetime
from typing import Any, Callable, Mapping
from urllib.parse import urljoin

ABS_PRICES_URL = "https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/total-value-dwellings/latest-release"
ABS_APPROVALS_URL = "https://www.abs.gov.au/statistics/industry/building-and-construction/building-approvals-australia/latest-release"
ABS_LENDING_URL = "https://www.abs.gov.au/statistics/economy/finance/lending-indicators/latest-release"
RBA_D1_URL = "https://www.rba.gov.au/statistics/tables/csv/d1-data.csv"
RBA_F6_URL = "https://www.rba.gov.au/statistics/tables/csv/f6-data.csv"
RBA_E13_URL = "https://www.rba.gov.au/statistics/tables/csv/e13-data.csv"

D1 = {
    "housing_mom_pct": "DGFACHM", "housing_yoy_pct": "DGFACH12",
    "owner_occupier_mom_pct": "DGFACOHM", "owner_occupier_yoy_pct": "DGFACOH12",
    "investor_mom_pct": "DGFACIHM", "investor_yoy_pct": "DGFACIH12",
}
F6 = {
    "owner_occupier_outstanding_rate_pct": "FLRHOOTA",
    "owner_occupier_new_rate_pct": "FLRHOFTA",
    "investor_outstanding_rate_pct": "FLRHIOTA",
    "investor_new_rate_pct": "FLRHIFTA",
}
E13 = {
    "total_interest_charged_mn": "LPHTIC",
    "total_scheduled_repayments_mn": "LPHTSP",
    "total_excess_payments_mn": "LPHTEX",
    "interest_to_income_pct": "LPHTICRI",
    "scheduled_repayments_to_income_pct": "LPHTSPRI",
    "excess_payments_to_income_pct": "LPHTEXRI",
}
REQUIRED_FEEDS = {
    "prices_and_turnover", "building_approvals", "housing_lending",
    "housing_credit", "mortgage_rates", "mortgage_cash_flow",
}


class AustraliaHousingError(RuntimeError):
    pass


def _num(x):
    if x is None:
        return None
    try:
        return float(str(x).strip().replace(",", "").replace("$", "").replace("%", ""))
    except (TypeError, ValueError):
        return None


def _date(x):
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    for fmt in ("%b-%y", "%b %Y", "%B %Y"):
        try:
            d = datetime.strptime(s, fmt)
            return date(d.year, d.month, monthrange(d.year, d.month)[1])
        except ValueError:
            pass
    q = re.fullmatch(r"(March|June|September|December) Quarter (\d{4})", s, re.I)
    if q:
        month = {"march": 3, "june": 6, "september": 9, "december": 12}[q.group(1).lower()]
        year = int(q.group(2))
        return date(year, month, monthrange(year, month)[1])
    return None


def _text(raw):
    raw = re.sub(r"(?is)<script.*?</script>|<style.*?</style>", " ", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    return re.sub(r"\s+", " ", html.unescape(raw)).strip()


def _ref(text):
    m = re.search(r"Reference period\s+(.+?)\s+Released", text, re.I)
    return (m.group(1).strip(), _date(m.group(1))) if m else (None, None)


def _metric(value, unit, **changes):
    return {"value": value, "unit": unit, **{k: v for k, v in changes.items() if v is not None}}


def _triplet(block, label):
    m = re.search(
        rf"{re.escape(label)}\s+([\d,.]+)\s+(-?\d+(?:\.\d+)?)\s+(-?\d+(?:\.\d+)?)",
        block, re.I,
    )
    if not m:
        raise AustraliaHousingError(f"missing ABS row {label}")
    vals = tuple(_num(m.group(i)) for i in range(1, 4))
    if any(v is None for v in vals):
        raise AustraliaHousingError(f"invalid ABS row {label}")
    return vals


def parse_abs_prices(raw):
    text = _text(raw)
    label, as_of = _ref(text)
    total = re.search(
        r"total value of residential dwellings in Australia (fell|rose) by \$([\d,.]+) billion to \$([\d,.]+) billion",
        text, re.I,
    )
    count = re.search(r"number of residential dwellings (rose|fell) by ([\d,]+) to ([\d,]+)", text, re.I)
    mean = re.search(r"mean price of residential dwellings (fell|rose) by \$([\d,]+) to \$([\d,]+)", text, re.I)
    if not (total and count and mean):
        raise AustraliaHousingError("ABS dwelling-price headlines missing")
    xlsx = None
    for m in re.finditer(r'href=["\']([^"\']+\.xlsx[^"\']*)["\']', raw, re.I):
        if "643202" in m.group(1):
            xlsx = urljoin(ABS_PRICES_URL, html.unescape(m.group(1)))
            break
    sign = lambda word: -1 if word.lower() == "fell" else 1
    return {
        "reference_period": label, "as_of": as_of.isoformat() if as_of else None,
        "metrics": {
            "total_dwelling_value_billion_aud": _metric(_num(total.group(3)), "AUD billion", change_amount=sign(total.group(1)) * (_num(total.group(2)) or 0)),
            "dwelling_count": _metric(int(_num(count.group(3)) or 0), "dwellings", change_amount=sign(count.group(1)) * int(_num(count.group(2)) or 0)),
            "mean_dwelling_price_aud": _metric(_num(mean.group(3)), "AUD", change_amount=sign(mean.group(1)) * (_num(mean.group(2)) or 0)),
        },
        "transfer_workbook_url": xlsx,
    }


def parse_abs_transfers(data):
    try:
        import openpyxl
    except ImportError as exc:
        raise AustraliaHousingError("openpyxl required for ABS transfer workbook") from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        head_i, headers = None, None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 12, 12), values_only=True), 1):
            vals = [str(v or "").strip() for v in row]
            if sum("Transfers" in v for v in vals) >= 2:
                head_i, headers = i, vals
                break
        if not headers:
            continue
        house = [i for i, h in enumerate(headers) if "Number of Established House Transfers" in h and "Total Australia" not in h]
        attached = [i for i, h in enumerate(headers) if "Number of Attached Dwelling Transfers" in h and "Total Australia" not in h]
        if not (house or attached):
            continue
        latest = None
        for row in ws.iter_rows(min_row=head_i + 1, values_only=True):
            d = _date(row[0] if row else None)
            if d and (latest is None or d > latest[0]):
                latest = (d, row)
        if not latest:
            continue
        d, row = latest
        def summ(cols):
            vals = [_num(row[i]) for i in cols if i < len(row)]
            vals = [v for v in vals if v is not None]
            return int(round(sum(vals))) if vals else None
        h, a = summ(house), summ(attached)
        return {
            "as_of": d.isoformat(),
            "metrics": {
                "established_house_transfers": _metric(h, "transfers"),
                "attached_dwelling_transfers": _metric(a, "transfers"),
                "residential_transfers_derived": _metric((h or 0) + (a or 0) if h is not None or a is not None else None, "transfers"),
            },
            "note": "National transfer count is deterministically summed from ABS capital-city/rest-of-state house and attached-dwelling transfer series.",
        }
    raise AustraliaHousingError("ABS transfer workbook headers not found")


def parse_abs_approvals(raw):
    text = _text(raw)
    label, as_of = _ref(text)
    total = _triplet(text, "Total dwelling units approved")
    houses = _triplet(text, "Private sector houses")
    other = _triplet(text, "Private sector dwellings excluding houses")
    value = re.search(r"value of total residential building (fell|rose) (\d+(?:\.\d+)?)% to \$([\d,.]+)b", text, re.I)
    if not value:
        raise AustraliaHousingError("ABS residential-building value missing")
    sign = -1 if value.group(1).lower() == "fell" else 1
    return {
        "reference_period": label, "as_of": as_of.isoformat() if as_of else None,
        "metrics": {
            "total_dwellings_approved": _metric(total[0], "dwellings", change_mom_pct=total[1], change_yoy_pct=total[2]),
            "private_houses_approved": _metric(houses[0], "dwellings", change_mom_pct=houses[1], change_yoy_pct=houses[2]),
            "private_other_dwellings_approved": _metric(other[0], "dwellings", change_mom_pct=other[1], change_yoy_pct=other[2]),
            "residential_building_value_billion_aud": _metric(_num(value.group(3)), "AUD billion", change_mom_pct=sign * (_num(value.group(2)) or 0)),
        },
    }


def parse_abs_lending(raw):
    text = _text(raw)
    label, as_of = _ref(text)
    a = text.find("Number of new loan commitments for dwellings")
    b = text.find("Value of new loan commitments for dwellings")
    if a < 0 or b <= a:
        raise AustraliaHousingError("ABS lending tables missing")
    n, v = text[a:b], text[b:]
    nt, no, nf, ni = (_triplet(n, x) for x in ("Total loan commitments", "Owner occupier", "First home buyers", "Investor"))
    vt, vo, vf, vi = (_triplet(v, x) for x in ("Total loan commitments", "Owner occupier", "First home buyers", "Investor"))
    def q(x, unit):
        return _metric(x[0], unit, change_qoq_pct=x[1], change_yoy_pct=x[2])
    return {
        "reference_period": label, "as_of": as_of.isoformat() if as_of else None,
        "metrics": {
            "new_commitments_count": q(nt, "loans"), "owner_occupier_count": q(no, "loans"),
            "first_home_buyer_count": q(nf, "loans"), "investor_count": q(ni, "loans"),
            "new_commitments_value_billion_aud": q(vt, "AUD billion"),
            "owner_occupier_value_billion_aud": q(vo, "AUD billion"),
            "first_home_buyer_value_billion_aud": q(vf, "AUD billion"),
            "investor_value_billion_aud": q(vi, "AUD billion"),
        },
    }


def parse_rba(text, selected: Mapping[str, str]):
    rows = list(csv.reader(io.StringIO(text)))
    meta, ids_i = {}, None
    for i, row in enumerate(rows):
        if not row:
            continue
        key = row[0].strip().lower()
        if key in {"title", "frequency", "units", "publication date", "series id"}:
            meta[key] = row
        if key == "series id":
            ids_i = i
            break
    if ids_i is None:
        raise AustraliaHousingError("RBA Series ID row missing")
    ids = meta["series id"]
    cols = {str(v).strip(): i for i, v in enumerate(ids) if i}
    missing = [sid for sid in selected.values() if sid not in cols]
    if missing:
        raise AustraliaHousingError(f"RBA series missing: {missing}")
    latest = None
    for row in rows[ids_i + 1:]:
        d = _date(row[0] if row else None)
        if d and (latest is None or d > latest[0]):
            latest = (d, row)
    if not latest:
        raise AustraliaHousingError("RBA observations missing")
    d, row = latest
    metrics, series = {}, {}
    for name, sid in selected.items():
        col = cols[sid]
        metrics[name] = _num(row[col] if col < len(row) else None)
        series[name] = {
            "series_id": sid,
            "title": meta.get("title", [None] * (col + 1))[col] if len(meta.get("title", [])) > col else None,
            "unit": meta.get("units", [None] * (col + 1))[col] if len(meta.get("units", [])) > col else None,
        }
    return {"as_of": d.isoformat(), "metrics": metrics, "series": series}


def _one(name, url, cadence, today, fetch, parser):
    try:
        parsed = parser(fetch(url).decode("utf-8-sig", errors="replace"))
        d = _date(parsed.get("as_of"))
        age = (today - d).days if d else None
        stale = age is not None and age > (70 if cadence == "monthly" else 130)
        return {"status": "stale" if stale else "ok", "source_name": name, "url": url, "age_days": age, **parsed}
    except Exception as exc:
        return {"status": "unavailable", "source_name": name, "url": url, "as_of": None, "age_days": None, "error": str(exc)}


def collect_australia_housing(*, today: date, fetch_bytes: Callable[..., bytes]):
    prices = _one("ABS Total Value of Dwellings", ABS_PRICES_URL, "quarterly", today, fetch_bytes, parse_abs_prices)
    if prices.get("status") != "unavailable" and prices.get("transfer_workbook_url"):
        try:
            prices["transfers"] = parse_abs_transfers(fetch_bytes(prices["transfer_workbook_url"]))
        except Exception as exc:
            prices["transfers"] = {"status": "unavailable", "error": str(exc), "as_of": None}
    approvals = _one("ABS Building Approvals", ABS_APPROVALS_URL, "monthly", today, fetch_bytes, parse_abs_approvals)
    lending = _one("ABS Lending Indicators", ABS_LENDING_URL, "quarterly", today, fetch_bytes, parse_abs_lending)
    credit = _one("RBA D1 Growth in Selected Financial Aggregates", RBA_D1_URL, "monthly", today, fetch_bytes, lambda x: parse_rba(x, D1))
    rates = _one("RBA F6 Housing Lending Rates", RBA_F6_URL, "monthly", today, fetch_bytes, lambda x: parse_rba(x, F6))
    cash = _one("RBA E13 Housing Loan Payments", RBA_E13_URL, "quarterly", today, fetch_bytes, lambda x: parse_rba(x, E13))
    feeds = {
        "prices_and_turnover": prices, "building_approvals": approvals, "housing_lending": lending,
        "housing_credit": credit, "mortgage_rates": rates, "mortgage_cash_flow": cash,
    }
    usable = sum(x.get("status") in {"ok", "stale"} for x in feeds.values())
    unavailable = sum(x.get("status") == "unavailable" for x in feeds.values())
    status = "unavailable" if not usable else "partial" if unavailable or any(x.get("status") == "stale" for x in feeds.values()) else "ok"
    return {
        "country": "AU", "status": status, "feeds": feeds,
        "method": {
            "model_calls": 0, "credentials_required": [],
            "purpose": "Official Australian housing transmission data for Market Watch and Trader Room context.",
            "trader_packet_delivery": "market_state_passthrough",
        },
    }


def validate_australia_housing(block: Mapping[str, Any]):
    if block.get("country") != "AU" or block.get("status") not in {"ok", "partial", "unavailable"}:
        raise AustraliaHousingError("invalid Australian housing block")
    feeds = block.get("feeds")
    if not isinstance(feeds, Mapping) or set(feeds) != REQUIRED_FEEDS:
        raise AustraliaHousingError("Australian housing feeds incomplete")
    for name, feed in feeds.items():
        if feed.get("status") not in {"ok", "stale", "unavailable"} or not feed.get("url"):
            raise AustraliaHousingError(f"invalid Australian housing feed {name}")
        if feed.get("status") == "unavailable" and not feed.get("error"):
            raise AustraliaHousingError(f"{name} unavailable without provenance")
        if feed.get("status") != "unavailable" and not feed.get("as_of"):
            raise AustraliaHousingError(f"{name} missing as_of")
